#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
verify_data.py -- INDEPENDENT verifier for the AMS workbook -> docs/data conversion.

    py tools/verify_data.py "<xlsx>" docs/data [--quick] [--samples N]

Reads the xlsx (openpyxl read_only, data_only=False, formulas evaluated by an
own mini Excel engine written here) and compares it with docs/data/*.json.
Writes tools/verify_report.json and tools/verify_report.md next to this script.
Exit code 0 only when there are zero errors.

This file deliberately does NOT import or read tools/extract.py / tools/amsx*:
it is an independent oracle built from the specs + the xlsx only.
"""
import sys, os, re, json, math, time, zipfile, argparse, collections, html
import datetime as dt
import multiprocessing as mp
from decimal import Decimal, ROUND_HALF_UP
from xml.etree import ElementTree as ET

try:
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')
except Exception:
    pass

import openpyxl
from openpyxl.worksheet._reader import WorkSheetParser
from openpyxl.styles.numbers import BUILTIN_FORMATS

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_CHARTS_CONTRACT = (r"C:/Users/bacon/AppData/Local/Temp/claude/C--Users-bacon-AMS/"
                           r"fc90b50f-6cf5-4df3-af01-0acbf9ded8c3/scratchpad/charts_probe/charts_contract.json")

# ----------------------------------------------------------------------------------------------
# EXPECTED values collected from the analyst specs (specs/00.md .. 27.md, charts.md)
# ----------------------------------------------------------------------------------------------
GRID_IDS = {'00', '01', '14', '16', '17', '18', '24'}
CARD_IDS = {'02'}
CHUNKED = {'21', '22'}


def expected_mode(sid):
    return 'grid' if sid in GRID_IDS else ('card' if sid in CARD_IDS else 'table')


EXPECTED_ROWS = {'03': 1928, '04': 1675, '05': 14066, '06': 1928, '07': 1013, '08': 6386, '09': 140,
                 '10': 13682, '11': 5138, '12': 2007, '13': 761, '15': 201, '19': 3484, '20': 9271,
                 '21': 204562, '22': 140381, '23': 890, '25': 2126, '26': 2178, '27': 1370}
EXPECTED_COLS = {'03': 120, '04': 36, '05': 11, '06': 30, '07': 25, '08': 40, '09': 18, '10': 32, '11': 24,
                 '12': 36, '13': 14, '15': 20, '19': 13, '20': 14, '21': 17, '22': 20, '23': 8, '25': 20,
                 '26': 13, '27': 14}
EXPECTED_FREEZE = {'03': 2, '04': 6, '05': 1, '06': 2, '07': 7, '08': 4, '09': 4, '10': 5, '11': 3, '12': 3,
                   '13': 2, '15': 4, '19': 0, '20': 2, '21': 2, '22': 3, '23': 1, '25': 1, '26': 2, '27': 4}
# hidden parity column ("組") that spec 21/22 §5 allows the extractor to drop (zebra via ui.group_zebra)
DROPPABLE_COLS = {'21': {17}, '22': {20}}
EXPECTED_HEADER_KPI = {'07': {'高': 155, '中': 386, '低': 454, '資訊': 18}}
EXPECTED_CELLS = {  # (sheet id) -> {coord: expected evaluated value}
    '24': {'E53': 1928, 'E54': 1675, 'E55': 14066, 'E56': 1928, 'E57': 1928, 'E58': 1928, 'E59': 1013,
           'E62': 6386, 'E63': 4757, 'E66': 140, 'E67': 13682, 'E68': 5138, 'E71': 2007, 'E72': 761,
           'E78': 3484, 'E79': 9271, 'E80': 204562, 'E81': 140381, 'E82': 890, 'E84': 2126, 'E87': 2178,
           'E88': 1370},
    '14': {'F6': 1928, 'F7': 1577, 'F8': 17, 'F9': 0.9899, 'F10': 87.4, 'F11': 0, 'F12': 8116, 'F13': 4757,
           'F14': 435, 'F15': 379, 'F16': 155, 'F17': 429, 'F18': 7, 'F19': 28, 'F20': 26, 'F21': 367,
           'F22': 1053, 'F279': 13682},
    '07': {'S1': 155, 'U1': 386, 'W1': 454, 'Y1': 18},
}
EXPECTED_ALL_CHECK = {'24': ('✓', 63), '14': ('✓', 17 + 9)}   # all IF formulas evaluate to ✓ (count of formula ✓)
EXPECTED_01_KPI = {  # cell: (label, target sheet id, r, number_format)
    'B6': (1928, '03', 0, '#,##0'), 'F6': (1577, '14', 131, '#,##0'), 'J6': (17, '14', 41, '#,##0'),
    'N6': (0.9898507462686568, '15', 0, '0.0%'), 'R6': (87.4, '18', 5, '0.0'), 'V6': ('0 / 1,928', '18', 5, None),
    'B11': (8116, '16', 5, '#,##0'), 'F11': (4757, '08', 0, '#,##0'), 'J11': (435, '08', 0, '#,##0'),
    'N11': (379, '07', 0, '#,##0'), 'R11': (155, '07', 0, '#,##0'), 'V11': (429, '12', 0, '#,##0'),
    'X72': ('→ 18', '18', 1, None), 'X73': ('→ 00', '00', 1, None), 'X74': ('→ 16', '16', 1, None),
    'X75': ('→ 08', '08', None, None), 'X76': ('→ 07', '07', None, None), 'X77': ('→ 08', '08', None, None),
    'X78': ('→ 18', '18', 1, None), 'X79': ('→ 03', '03', None, None), 'X80': ('→ 12', '12', None, None),
    'X81': ('→ 17', '17', 1, None),
}
EXPECTED_SECTIONS = {
    '00': [4, 19, 37, 44, 56, 63, 106, 138], '01': [71, 83],
    '14': [3, 26, 41, 65, 82, 131, 142, 157, 186, 199, 212, 239, 253, 266, 283, 307],
    '16': [3, 31, 43, 73], '17': [3, 26, 54, 68], '18': [3, 16, 28, 70],
    '24': [4, 34, 48, 90, 111, 123, 138, 151, 167]}
EXPECTED_CHART_COUNT = {'01': 6, '16': 2}
CARD_DEFAULT = ('G11HAD60BT001', 'D00759', 229)
CARD_VECTORS = [  # query -> (alias, count, 03 idx0 or None, status startswith)
    ('G11HAD60BT001', 'D00759', 1, 229, '來源：1 目前AMS位號［鍵 G11HAD60BT001］ → 目前位號 G11HAD60BT001'),
    ('=g11had60bt001 ', 'D00759', 1, 229, '來源：1 目前AMS位號［鍵 G11HAD60BT001］ → 目前位號 G11HAD60BT001'),
    ('d01760', 'D01760', 1, 0, '來源：0 設備別名［鍵 D01760］ → 目前位號 G11_90LT-1'),
    ('1-pi-cw014-1', 'D01380', 2, 1268, None),
    ('C10GHC53', '', 0, None, '來源：12 測試定義清單位號(無對應設備)［鍵 C10GHC53］（測試定義清單位號，無對應設備）'),
    ('XYZ123', '', '', None, '查無此字串（萬用字元 * ? 不適用）：請到『05_位號索引』以 Ctrl+F 搜尋部分字串'),
    ('', '', '', None, '請輸入位號…'),
]
# CF rule "true" counts measured by the analysts: sheet -> {priority: ('cells'|'rows', n)}
EXPECTED_CF_COUNTS = {
    '03': {1: ('rows', 253), 2: ('cells', 24), 3: ('cells', 656), 4: ('cells', 26), 5: ('cells', 83),
           6: ('cells', 784), 7: ('cells', 341), 8: ('cells', 23), 9: ('cells', 444), 10: ('cells', 22),
           11: ('cells', 18), 12: ('cells', 17), 13: ('cells', 1), 14: ('cells', 1), 16: ('cells', 65),
           17: ('cells', 136), 18: ('cells', 194), 19: ('cells', 296), 20: ('cells', 0), 22: ('cells', 132),
           23: ('cells', 114), 24: ('cells', 132), 25: ('cells', 0), 26: ('cells', 17), 27: ('cells', 1246),
           28: ('cells', 32), 29: ('cells', 2)},
    '04': {1: ('cells', 1675), 2: ('cells', 1675), 3: ('rows', 17), 4: ('cells', 132), 5: ('cells', 114),
           6: ('cells', 132), 7: ('cells', 0), 8: ('cells', 17), 9: ('cells', 1246), 10: ('cells', 32),
           11: ('cells', 2), 12: ('cells', 40), 13: ('cells', 23), 14: ('cells', 435), 15: ('rows', 298),
           16: ('cells', 656), 17: ('cells', 18)},
    '05': {1: ('cells', 4371)},
    '06': {1: ('cells', 142), 2: ('cells', 9), 3: ('cells', 26), 4: ('cells', 144)},
    '07': {1: ('cells', 155), 2: ('cells', 386), 3: ('cells', 454), 4: ('cells', 18), 5: ('rows', 0),
           6: ('cells', 1013), 7: ('cells', 1013), 8: ('cells', 1013), 9: ('cells', 1013), 10: ('cells', 1013)},
    '08': {1: ('rows', 2259), 2: ('rows', 1629), 3: ('cells', 485), 4: ('cells', 32), 5: ('cells', 168),
           6: ('cells', 573), 7: ('cells', 11),  # spec 08 says 10; direct count of Z containing 浮點捨入 = 11 (incl. multi-line Z3254)
           8: ('cells', 0), 9: ('cells', 0), 10: ('cells', 6386)},
    '09': {1: ('rows', 18), 2: ('rows', 14)},
    '10': {1: ('rows', 6001), 2: ('rows', 140), 3: ('rows', 3), 4: ('rows', 1928), 5: ('rows', 5),
           6: ('rows', 1), 7: ('rows', 146), 8: ('rows', 3)},
    '11': {1: ('cells', 403)},
    '12': {1: ('cells', 2007), 2: ('cells', 2007), 3: ('cells', 132), 4: ('cells', 114), 5: ('cells', 132),
           6: ('cells', 0), 7: ('cells', 17), 8: ('cells', 1578), 9: ('cells', 32), 10: ('cells', 2),
           11: ('rows', 1578), 12: ('cells', 17)},
    '13': {1: ('rows', 332)},
    '15': {1: ('rows', 1), 2: ('rows', 8), 3: ('rows', 18), 4: ('cells', 86), 5: ('cells', 66)},
    '19': {1: ('rows', 24), 2: ('rows', 1473)},
    '20': {1: ('cells', 8681), 2: ('cells', 372)},
    '21': {1: ('rows', 104137)}, '22': {1: ('rows', 71677)},
    '23': {1: ('cells', 80), 2: ('cells', 153)},
    '25': {1: ('rows', 43)}, '26': {1: ('rows', 7)},
    '14': {1: ('cells', 0), 2: ('cells', 17), 3: ('cells', 0), 4: ('cells', 0), 6: ('cells', 7)},
    '16': {1: ('rows', 7)}, '18': {2: ('cells', 1)}, '24': {1: ('cells', 73), 2: ('cells', 0)},
}

# ----------------------------------------------------------------------------------------------
# Report
# ----------------------------------------------------------------------------------------------


def short(v, n=160):
    if v is None:
        return None
    try:
        s = v if isinstance(v, str) else json.dumps(v, ensure_ascii=False, default=str)
    except Exception:
        s = repr(v)
    return s if len(s) <= n else s[:n] + '…(%d chars)' % len(s)


class Report:
    def __init__(self, nsamples=8):
        self.n = nsamples
        self.counts = collections.Counter()
        self.samples = collections.defaultdict(list)
        self.notes = []
        self.unchecked = []
        self.stats = collections.defaultdict(dict)

    def add(self, level, cat, sheet, msg, where=None, exp=None, act=None):
        k = (level, cat, sheet or '-')
        self.counts[k] += 1
        if len(self.samples[k]) < self.n:
            d = {'msg': msg}
            if where is not None:
                d['where'] = where
            if exp is not None or act is not None:
                d['expected'] = short(exp)
                d['actual'] = short(act)
            self.samples[k].append(d)

    def err(self, cat, sheet, msg, where=None, exp=None, act=None):
        self.add('error', cat, sheet, msg, where, exp, act)

    def warn(self, cat, sheet, msg, where=None, exp=None, act=None):
        self.add('warning', cat, sheet, msg, where, exp, act)

    def info(self, sheet, msg):
        self.notes.append('[%s] %s' % (sheet or '-', msg))

    def uncheck(self, sheet, msg):
        self.unchecked.append('[%s] %s' % (sheet or '-', msg))

    def state(self):
        return {'counts': [[list(k), v] for k, v in self.counts.items()],
                'samples': [[list(k), v] for k, v in self.samples.items()],
                'notes': self.notes, 'unchecked': self.unchecked, 'stats': dict(self.stats)}

    def merge(self, st):
        for k, v in st['counts']:
            self.counts[tuple(k)] += v
        for k, v in st['samples']:
            lst = self.samples[tuple(k)]
            for s in v:
                if len(lst) < self.n:
                    lst.append(s)
        self.notes += st['notes']
        self.unchecked += st['unchecked']
        for k, v in st['stats'].items():
            self.stats[k].update(v)

    def total(self, level):
        return sum(v for k, v in self.counts.items() if k[0] == level)


# ----------------------------------------------------------------------------------------------
# small helpers
# ----------------------------------------------------------------------------------------------

def col2num(s):
    n = 0
    for ch in s:
        n = n * 26 + ord(ch) - 64
    return n


def num2col(n):
    s = ''
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def coord(r, c):
    return '%s%d' % (num2col(c), r)


def parse_coord(s):
    m = re.match(r'^\$?([A-Z]{1,3})\$?(\d+)$', s)
    return int(m.group(2)), col2num(m.group(1))


def parse_area(ref):
    """'A4:DP1932' -> (r1,c1,r2,c2); 'C16' -> (16,3,16,3)"""
    a = ref.split(':')
    r1, c1 = parse_coord(a[0])
    r2, c2 = parse_coord(a[1]) if len(a) > 1 else (r1, c1)
    return r1, c1, r2, c2


def half_up(x):
    return int(math.floor(x + 0.5))


def normc(c):
    """normalise colour -> '#RRGGBB' or None"""
    if c is None or c == '' or c is False:
        return None
    if isinstance(c, str):
        s = c.strip().lstrip('#').upper()
        if len(s) == 8:
            s = s[2:]
        if len(s) == 3:
            s = ''.join(ch * 2 for ch in s)
        if re.fullmatch(r'[0-9A-F]{6}', s):
            return '#' + s
        return c
    return None


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def num_eq(a, b):
    if a == b:
        return True
    try:
        return abs(a - b) <= 1e-9 * max(abs(a), abs(b)) or abs(a - b) < 1e-12
    except Exception:
        return False


EPOCH = dt.datetime(1899, 12, 30)


def to_serial(v):
    if isinstance(v, dt.datetime):
        return (v - EPOCH).total_seconds() / 86400.0
    if isinstance(v, dt.date):
        return float((v - EPOCH.date()).days)
    return v


def general_str(x):
    if isinstance(x, bool):
        return 'TRUE' if x else 'FALSE'
    if isinstance(x, int):
        return str(x)
    if isinstance(x, float):
        if x == int(x) and abs(x) < 1e15:
            return str(int(x))
        s = '%.15g' % x
        if 'e' in s:
            m, e = s.split('e')
            s = m + 'E' + ('+' if int(e) >= 0 else '-') + '%02d' % abs(int(e))
        return s
    return str(x)


def fmt_datetime(v, nf):
    """CONTRACT v2 datetime rule (implemented independently):
       nf contains '.000'  -> YYYY-MM-DD HH:MM:SS.fff
       nf == yyyy-mm-dd    -> YYYY-MM-DD
       nf == yyyy-mm-dd hh:mm -> YYYY-MM-DD HH:MM
       other date formats  -> YYYY-MM-DD HH:MM:SS, rounded to the nearest second (Excel display)."""
    if isinstance(v, dt.time):
        v = dt.datetime.combine(dt.date(1899, 12, 30), v)
    elif isinstance(v, dt.date) and not isinstance(v, dt.datetime):
        v = dt.datetime(v.year, v.month, v.day)
    nfl = (nf or '').lower().replace('\\', '').replace('"', '').strip()
    if '.000' in nfl:
        ms = (v.microsecond + 500) // 1000
        base = v.replace(microsecond=0)
        if ms >= 1000:
            base += dt.timedelta(seconds=1)
            ms = 0
        return base.strftime('%Y-%m-%d %H:%M:%S') + '.%03d' % ms
    if nfl == 'yyyy-mm-dd':
        return v.strftime('%Y-%m-%d')
    r = (v + dt.timedelta(microseconds=500000)).replace(microsecond=0)
    if nfl == 'yyyy-mm-dd hh:mm':
        return r.strftime('%Y-%m-%d %H:%M')
    return r.strftime('%Y-%m-%d %H:%M:%S')


def norm_nf(nf):
    if nf is None or nf in ('', 'General', '@'):
        return None
    return nf


def is_date_nf(nf):
    n = (nf or '').lower()
    return any(t in n for t in ('yy', 'dd', 'hh', 'h:mm', 'mm:ss'))


# ----------------------------------------------------------------------------------------------
# Workbook reading (zip/XML for structure, openpyxl WorkSheetParser for cells)
# ----------------------------------------------------------------------------------------------
NS_MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
NS_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS_PKG = 'http://schemas.openxmlformats.org/package/2006/relationships'


class Formula:
    __slots__ = ('text',)

    def __init__(self, t):
        self.text = t

    def __repr__(self):
        return 'Formula(%r)' % self.text[:90]


def xml_headtail(z, part):
    """Return (head, tail) text of a worksheet XML without the (possibly huge) <sheetData>."""
    head, tail = [], []
    state = 0
    carry = b''
    with z.open(part) as f:
        while True:
            chunk = f.read(1 << 22)
            if not chunk:
                break
            s = carry + chunk
            carry = b''
            if state == 0:
                i = s.find(b'<sheetData')
                if i < 0:
                    head.append(s[:-32])
                    carry = s[-32:]
                    continue
                head.append(s[:i])
                state = 1
                s = s[i:]
            if state == 1:
                j = s.find(b'</sheetData>')
                if j < 0:
                    carry = s[-32:]
                    continue
                state = 2
                s = s[j:]
            if state == 2:
                tail.append(s)
    if state == 0:
        head.append(carry)
    return b''.join(head).decode('utf-8'), b''.join(tail).decode('utf-8')


def _attrs(tagtext):
    return dict(re.findall(r'([\w:]+)="([^"]*)"', tagtext))


class SheetXML:
    def __init__(self, head, tail):
        self.head, self.tail = head, tail
        self.merges = [parse_area(m) for m in re.findall(r'<mergeCell ref="([^"]+)"', tail)]
        self.cols = []
        for m in re.finditer(r'<col [^>]*/?>', head):
            a = _attrs(m.group(0))
            self.cols.append({'min': int(a['min']), 'max': int(a['max']), 'width': float(a.get('width', 0) or 0),
                              'hidden': a.get('hidden') in ('1', 'true'), 'custom': a.get('customWidth') in ('1', 'true')})
        m = re.search(r'<pane [^>]*/?>', head)
        pa = _attrs(m.group(0)) if m else {}
        self.xsplit = int(float(pa.get('xSplit', 0) or 0))
        self.ysplit = int(float(pa.get('ySplit', 0) or 0))
        m = re.search(r'<sheetView [^>]*>', head)
        sv = _attrs(m.group(0)) if m else {}
        self.gridlines = sv.get('showGridLines', '1') not in ('0', 'false')
        m = re.search(r'<autoFilter ref="([^"]+)"', tail)
        self.autofilter = parse_area(m.group(1)) if m else None
        self.has_drawing = '<drawing ' in tail
        self.dv = []
        for m in re.finditer(r'<dataValidation\b([^>]*)>(.*?)</dataValidation>', tail, re.S):
            a = _attrs(m.group(1))
            f1 = re.search(r'<formula1>(.*?)</formula1>', m.group(2), re.S)
            self.dv.append({'sqref': a.get('sqref', ''), 'type': a.get('type'), 'operator': a.get('operator'),
                            'prompt': html.unescape(a['prompt']) if a.get('prompt') else None,
                            'formula1': html.unescape(f1.group(1)) if f1 else None})
        self.cf = []
        for m in re.finditer(r'<conditionalFormatting\b.*?</conditionalFormatting>', tail, re.S):
            el = ET.fromstring(m.group(0))
            ranges = [parse_area(x) for x in el.get('sqref').split()]
            rules = []
            for r in el.findall('cfRule'):
                rd = {'type': r.get('type'), 'priority': int(r.get('priority')),
                      'dxf': int(r.get('dxfId')) if r.get('dxfId') is not None else None,
                      'stop': r.get('stopIfTrue') in ('1', 'true'),
                      'formulas': [f.text or '' for f in r.findall('formula')], 'ranges': ranges,
                      'sqref': el.get('sqref')}
                cs = r.find('colorScale')
                if cs is not None:
                    rd['cfvo'] = [(v.get('type'), v.get('val')) for v in cs.findall('cfvo')]
                    rd['colors'] = [normc(c.get('rgb')) for c in cs.findall('color')]
                db = r.find('dataBar')
                if db is not None:
                    rd['cfvo'] = [(v.get('type'), v.get('val')) for v in db.findall('cfvo')]
                    c = db.find('color')
                    rd['colors'] = [normc(c.get('rgb')) if c is not None else None]
                rules.append(rd)
            self.cf.extend(rules)

    def col_width(self, c):
        for d in self.cols:
            if d['min'] <= c <= d['max']:
                return d
        return None

    def hidden_cols(self):
        s = set()
        for d in self.cols:
            if d['hidden']:
                s.update(range(d['min'], d['max'] + 1))
        return s


class StyleInfo:
    __slots__ = ('nf', 'b', 'i', 'u', 'fc', 'bg', 'sz', 'font', 'ha', 'va', 'wrap', 'bd', 'strike', 'ind', 'rot', 'shrink')

    def __init__(self, **kw):
        for k, v in kw.items():
            setattr(self, k, v)


def _color_hex(col):
    if col is None:
        return None
    try:
        if getattr(col, 'type', None) == 'rgb' and isinstance(col.rgb, str):
            return normc(col.rgb)
    except Exception:
        pass
    return None


class StyleCache:
    def __init__(self, wb):
        self.wb = wb
        self.c = {}

    def get(self, sid):
        si = self.c.get(sid)
        if si is None:
            wb = self.wb
            st = wb._cell_styles[sid]
            nfid = st.numFmtId
            nf = BUILTIN_FORMATS.get(nfid, 'General') if nfid < 164 else wb._number_formats[nfid - 164]
            font = wb._fonts[st.fontId]
            fill = wb._fills[st.fillId]
            al = wb._alignments[st.alignmentId]
            bd = wb._borders[st.borderId]
            bg = None
            if getattr(fill, 'patternType', None) == 'solid':
                bg = _color_hex(fill.fgColor)
            sides = {}
            for side in ('left', 'right', 'top', 'bottom'):
                sd = getattr(bd, side, None)
                if sd is not None and sd.style:
                    sides[side[0]] = (sd.style, _color_hex(sd.color))
            si = StyleInfo(nf=nf, b=bool(font.b), i=bool(font.i), u=bool(font.u), fc=_color_hex(font.color),
                           bg=bg, sz=font.sz, font=font.name, ha=al.horizontal, va=al.vertical,
                           wrap=bool(al.wrap_text), bd=sides, strike=bool(font.strike),
                           ind=int(al.indent or 0), rot=int(al.textRotation or 0), shrink=bool(al.shrinkToFit))
            self.c[sid] = si
        return si


class NotLoaded(Exception):
    pass


class Unsupported(Exception):
    pass


class SheetData:
    def __init__(self, name, sid, idx, xml):
        self.name, self.id, self.idx, self.xml = name, sid, idx, xml
        self.rows = {}      # r -> list of raw values (index c-1); Formula for real formulas
        self.srows = {}     # r -> list of style ids
        self.max_row = 0
        self.max_col = 0
        self.rowdims = {}
        self.colstore = None   # big sheets: {c: [values from row 5..]}
        self.colstore_first = 5
        self.merge_cover = None

    def get(self, r, c):
        row = self.rows.get(r)
        if row is not None:
            return row[c - 1] if c - 1 < len(row) else None
        if self.colstore is not None and r >= self.colstore_first:
            cs = self.colstore.get(c)
            if cs is None:
                raise NotLoaded('%s col %s not kept in memory' % (self.name, num2col(c)))
            i = r - self.colstore_first
            return cs[i] if i < len(cs) else None
        return None

    def sid(self, r, c):
        row = self.srows.get(r)
        if row is not None and c - 1 < len(row):
            return row[c - 1]
        return 0

    def covered(self):
        """map (r,c) -> anchor (r1,c1) for merge-covered (non anchor) cells"""
        if self.merge_cover is None:
            d = {}
            for (r1, c1, r2, c2) in self.xml.merges:
                for r in range(r1, r2 + 1):
                    for c in range(c1, c2 + 1):
                        if (r, c) != (r1, c1):
                            d[(r, c)] = (r1, c1)
            self.merge_cover = d
        return self.merge_cover


_LAST_ROWDIMS = {}


def iter_sheet_cells(wb, ws, max_row=None, min_row=1):
    """yield (r, [cell dicts]) using openpyxl's internal streaming parser (fast, gives style ids)."""
    with ws._get_source() as src:
        parser = WorkSheetParser(src, ws._shared_strings, data_only=False, epoch=wb.epoch,
                                 date_formats=wb._date_formats, timedelta_formats=wb._timedelta_formats)
        for ridx, cells in parser.parse():
            if max_row is not None and ridx > max_row:
                break
            if ridx < min_row:
                continue
            yield ridx, cells
        _LAST_ROWDIMS[ws.title] = dict(parser.row_dimensions)


class Book:
    def __init__(self, path):
        self.path = path
        self.z = zipfile.ZipFile(path)
        self.names = []          # workbook order
        self.parts = {}          # name -> xml part
        self.sheets = {}         # name -> SheetData
        self.by_id = {}          # '03' -> name
        self.xml = {}            # name -> SheetXML
        self.wb = None
        self.styles = None
        self._read_workbook()
        self.dxfs = self._read_dxfs()

    def _read_workbook(self):
        wbx = ET.fromstring(self.z.read('xl/workbook.xml'))
        rels = ET.fromstring(self.z.read('xl/_rels/workbook.xml.rels'))
        rmap = {r.get('Id'): r.get('Target') for r in rels.findall('{%s}Relationship' % NS_PKG)}
        for s in wbx.find('{%s}sheets' % NS_MAIN).findall('{%s}sheet' % NS_MAIN):
            name = s.get('name')
            tgt = rmap[s.get('{%s}id' % NS_REL)].lstrip('/')
            if not tgt.startswith('xl/'):
                tgt = 'xl/' + tgt
            self.names.append(name)
            self.parts[name] = tgt
            self.by_id[name[:2]] = name
        for n in self.names:
            self.xml[n] = SheetXML(*xml_headtail(self.z, self.parts[n]))

    def _read_dxfs(self):
        st = ET.fromstring(self.z.read('xl/styles.xml'))
        out = []
        dx = st.find('{%s}dxfs' % NS_MAIN)
        if dx is None:
            return out
        for d in dx.findall('{%s}dxf' % NS_MAIN):
            p = {}
            f = d.find('{%s}font' % NS_MAIN)
            if f is not None:
                for tag in ('b', 'i', 'u'):
                    e = f.find('{%s}%s' % (NS_MAIN, tag))
                    if e is not None:
                        p[tag] = e.get('val', '1') not in ('0', 'false')
                c = f.find('{%s}color' % NS_MAIN)
                if c is not None and c.get('rgb'):
                    p['fc'] = normc(c.get('rgb'))
            fl = d.find('{%s}fill/{%s}patternFill' % (NS_MAIN, NS_MAIN))
            if fl is not None:
                c = fl.find('{%s}bgColor' % NS_MAIN)
                if c is None or not c.get('rgb'):
                    c = fl.find('{%s}fgColor' % NS_MAIN)
                if c is not None and c.get('rgb'):
                    p['bg'] = normc(c.get('rgb'))
            out.append(p)
        return out

    def open_wb(self):
        if self.wb is None:
            self.wb = openpyxl.load_workbook(self.path, read_only=True, data_only=False)
            self.styles = StyleCache(self.wb)
        return self.wb

    def load(self, name, max_row=None):
        wb = self.open_wb()
        ws = wb[name]
        sd = SheetData(name, name[:2], self.names.index(name), self.xml[name])
        for ridx, cells in iter_sheet_cells(wb, ws, max_row=max_row):
            if not cells:
                continue
            maxc = cells[-1]['column']
            vals = [None] * maxc
            sids = [0] * maxc
            for cell in cells:
                c = cell['column']
                v = cell['value']
                if cell['data_type'] == 'f':
                    v = Formula(v)
                vals[c - 1] = v
                sids[c - 1] = cell['style_id'] or 0
            sd.rows[ridx] = vals
            sd.srows[ridx] = sids
            if ridx > sd.max_row:
                sd.max_row = ridx
            if maxc > sd.max_col:
                sd.max_col = maxc
        sd.rowdims = {int(k): v for k, v in _LAST_ROWDIMS.get(ws.title, {}).items()}
        self.sheets[name] = sd
        return sd

    def sheet(self, sid):
        return self.sheets.get(self.by_id.get(sid))

    def nf(self, sd, r, c):
        return self.styles.get(sd.sid(r, c)).nf


# ----------------------------------------------------------------------------------------------
# Mini Excel formula engine (own implementation; covers every pattern in formula_patterns.json
# plus the conditional-format formulas)
# ----------------------------------------------------------------------------------------------
class XErr:
    __slots__ = ('code',)

    def __init__(self, code):
        self.code = code

    def __repr__(self):
        return self.code

    def __eq__(self, o):
        return isinstance(o, XErr) and o.code == self.code

    def __hash__(self):
        return hash(self.code)


NA, VALUE, DIV0, REF = XErr('#N/A'), XErr('#VALUE!'), XErr('#DIV/0!'), XErr('#REF!')


class LinkVal:
    __slots__ = ('label', 'loc')

    def __init__(self, label, loc):
        self.label, self.loc = label, loc

    def __repr__(self):
        return 'Link(%r -> %r)' % (self.label, self.loc)


class RangeVal:
    __slots__ = ('sheet', 'r1', 'c1', 'r2', 'c2')

    def __init__(self, sheet, r1, c1, r2, c2):
        self.sheet, self.r1, self.c1, self.r2, self.c2 = sheet, r1, c1, r2, c2

    def single(self):
        return self.r1 is not None and self.r1 == self.r2 and self.c1 == self.c2


_TOK = re.compile(r'''
 (?P<ws>\s+)
|(?P<str>"(?:[^"]|"")*")
|(?P<sref>(?:'(?:[^']|'')+'|[A-Za-z_][\w.]*)!(?:\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?|\$?[A-Z]{1,3}:\$?[A-Z]{1,3}))
|(?P<func>[A-Z][A-Z0-9.]*(?=\())
|(?P<bool>(?:TRUE|FALSE)(?![A-Za-z0-9_(]))
|(?P<ref>\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?(?![A-Za-z0-9_(])|\$?[A-Z]{1,3}:\$?[A-Z]{1,3}(?![A-Za-z0-9_(]))
|(?P<num>\d+\.?\d*(?:[eE][-+]?\d+)?|\.\d+)
|(?P<op><>|<=|>=|[-+*/^&=<>(),%])
|(?P<name>[A-Za-z_][\w.]*)
''', re.X)


def tokenize(text):
    out = []
    pos = 0
    n = len(text)
    while pos < n:
        m = _TOK.match(text, pos)
        if not m:
            raise Unsupported('cannot tokenize at %r' % text[pos:pos + 20])
        k = m.lastgroup
        if k != 'ws':
            out.append((k, m.group(k)))
        pos = m.end()
    return out


def _ref_node(sheet, txt):
    parts = txt.split(':')

    def one(p):
        m = re.match(r'^(\$?)([A-Z]{1,3})(\$?)(\d*)$', p)
        return (col2num(m.group(2)), m.group(1) == '$', int(m.group(4)) if m.group(4) else None, m.group(3) == '$')
    a = one(parts[0])
    b = one(parts[1]) if len(parts) > 1 else a
    return ('ref', sheet, a[0], a[1], a[2], a[3], b[0], b[1], b[2], b[3])


class Parser:
    CMP = ('=', '<>', '<', '>', '<=', '>=')

    def __init__(self, text):
        self.t = tokenize(text)
        self.i = 0

    def peek(self):
        return self.t[self.i] if self.i < len(self.t) else (None, None)

    def take(self):
        tok = self.peek()
        self.i += 1
        return tok

    def expect(self, v):
        k, x = self.take()
        if x != v:
            raise Unsupported('expected %r got %r' % (v, x))

    def parse(self):
        e = self.cmp()
        if self.i != len(self.t):
            raise Unsupported('trailing tokens %r' % (self.t[self.i:],))
        return e

    def cmp(self):
        a = self.concat()
        while self.peek()[0] == 'op' and self.peek()[1] in self.CMP:
            op = self.take()[1]
            a = ('bin', op, a, self.concat())
        return a

    def concat(self):
        a = self.add()
        while self.peek() == ('op', '&'):
            self.take()
            a = ('bin', '&', a, self.add())
        return a

    def add(self):
        a = self.mul()
        while self.peek()[0] == 'op' and self.peek()[1] in '+-' and self.peek()[1] in ('+', '-'):
            op = self.take()[1]
            a = ('bin', op, a, self.mul())
        return a

    def mul(self):
        a = self.pow()
        while self.peek()[0] == 'op' and self.peek()[1] in ('*', '/'):
            op = self.take()[1]
            a = ('bin', op, a, self.pow())
        return a

    def pow(self):
        a = self.unary()
        while self.peek() == ('op', '^'):
            self.take()
            a = ('bin', '^', a, self.unary())
        return a

    def unary(self):
        if self.peek() == ('op', '-'):
            self.take()
            return ('neg', self.unary())
        if self.peek() == ('op', '+'):
            self.take()
            return self.unary()
        a = self.primary()
        while self.peek() == ('op', '%'):
            self.take()
            a = ('bin', '/', a, ('num', 100))
        return a

    def primary(self):
        k, x = self.take()
        if k == 'num':
            v = float(x)
            return ('num', int(v) if re.fullmatch(r'\d+', x) else v)
        if k == 'str':
            return ('str', x[1:-1].replace('""', '"'))
        if k == 'bool':
            return ('bool', x == 'TRUE')
        if k == 'ref':
            return _ref_node(None, x)
        if k == 'sref':
            i = x.rindex('!')
            sh = x[:i]
            if sh.startswith("'"):
                sh = sh[1:-1].replace("''", "'")
            return _ref_node(sh, x[i + 1:])
        if k == 'func':
            self.expect('(')
            args = []
            if self.peek() != ('op', ')'):
                while True:
                    if self.peek() in (('op', ','), ('op', ')')):
                        args.append(('missing',))
                    else:
                        args.append(self.cmp())
                    if self.peek() == ('op', ','):
                        self.take()
                        continue
                    break
            self.expect(')')
            return ('func', x, args)
        if k == 'op' and x == '(':
            e = self.cmp()
            self.expect(')')
            return e
        raise Unsupported('unexpected token %r %r' % (k, x))


def to_text(v):
    if isinstance(v, LinkVal):
        v = v.label
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, (int, float)):
        return general_str(v)
    if isinstance(v, (dt.datetime, dt.date)):
        return general_str(to_serial(v))
    if isinstance(v, XErr):
        return v.code
    return str(v)


_NUMRE = re.compile(r'^\s*[-+]?(\d+\.?\d*|\.\d+)([eE][-+]?\d+)?\s*$')


def to_num(v):
    if isinstance(v, LinkVal):
        v = v.label
    if v is None:
        return 0
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, (dt.datetime, dt.date)):
        return to_serial(v)
    if isinstance(v, XErr):
        return v
    if isinstance(v, str) and _NUMRE.match(v):
        return float(v)
    return VALUE


def xl_cmp(a, b):
    if isinstance(a, LinkVal):
        a = a.label
    if isinstance(b, LinkVal):
        b = b.label
    if isinstance(a, (dt.datetime, dt.date)):
        a = to_serial(a)
    if isinstance(b, (dt.datetime, dt.date)):
        b = to_serial(b)
    if a is None and b is None:
        return 0
    if a is None:
        a = '' if isinstance(b, str) else (False if isinstance(b, bool) else 0)
    if b is None:
        b = '' if isinstance(a, str) else (False if isinstance(a, bool) else 0)

    def rank(x):
        return 2 if isinstance(x, bool) else (1 if isinstance(x, str) else 0)
    ra, rb = rank(a), rank(b)
    if ra != rb:
        return -1 if ra < rb else 1
    if ra == 1:
        a, b = a.casefold(), b.casefold()
    return (a > b) - (a < b)


def wildcard_re(s):
    out = []
    i = 0
    while i < len(s):
        ch = s[i]
        if ch == '~' and i + 1 < len(s) and s[i + 1] in '*?~':
            out.append(re.escape(s[i + 1]))
            i += 2
            continue
        out.append('.*' if ch == '*' else ('.' if ch == '?' else re.escape(ch)))
        i += 1
    return re.compile(''.join(out), re.S | re.I)


def crit_pred(crit):
    """Excel COUNTIF criterion -> predicate over evaluated cell values."""
    if isinstance(crit, LinkVal):
        crit = crit.label
    if isinstance(crit, XErr):
        return lambda v: isinstance(v, XErr) and v == crit
    if isinstance(crit, bool):
        return lambda v: isinstance(v, bool) and v == crit
    if is_num(crit) or isinstance(crit, (dt.datetime, dt.date)):
        n = to_serial(crit)
        return lambda v: (is_num(v) or isinstance(v, dt.datetime)) and num_eq(to_serial(v), n)
    s = to_text(crit)
    m = re.match(r'^(<=|>=|<>|=|<|>)?(.*)$', s, re.S)
    op, rest = m.group(1), m.group(2)
    n = float(rest) if _NUMRE.match(rest) else None

    def numval(v):
        return to_serial(v) if (is_num(v) or isinstance(v, dt.datetime)) else None
    if op in ('<', '>', '<=', '>='):
        f = {'<': lambda a, b: a < b, '>': lambda a, b: a > b, '<=': lambda a, b: a <= b, '>=': lambda a, b: a >= b}[op]
        if n is not None:
            return lambda v: numval(v) is not None and f(numval(v), n)
        rl = rest.casefold()
        return lambda v: isinstance(v, str) and f(v.casefold(), rl)
    if n is not None:
        rs = rest.strip()
        base = lambda v: (numval(v) is not None and num_eq(numval(v), n)) or (isinstance(v, str) and v.strip() == rs)
    elif rest == '':
        base = lambda v: v is None or v == ''
    else:
        rx = wildcard_re(rest)
        base = lambda v: isinstance(v, str) and rx.fullmatch(v) is not None
    if op == '<>':
        return lambda v: not base(v)
    return base


def round_half_away(x, n):
    q = Decimal(1).scaleb(-n)
    d = Decimal(repr(float(x))).quantize(q, rounding=ROUND_HALF_UP)
    return float(d)


class Engine:
    """Evaluates formulas stored in Book sheets (values memoised)."""

    def __init__(self, book):
        self.book = book
        self.ast_cache = {}
        self.memo = {}
        self.over = {}
        self.match_idx = {}
        self.range_cache = {}
        self.countif_cache = {}
        self.unsupported = collections.Counter()

    # --- plumbing
    def parse(self, text):
        a = self.ast_cache.get(text)
        if a is None:
            t = text[1:] if text.startswith('=') else text
            a = Parser(t).parse()
            self.ast_cache[text] = a
        return a

    def sheetdata(self, name):
        sd = self.book.sheets.get(name)
        if sd is None:
            raise NotLoaded('sheet %s not loaded' % name)
        return sd

    def raw(self, sname, r, c):
        return self.sheetdata(sname).get(r, c)

    def cellval(self, sname, r, c):
        key = (sname, r, c)
        if key in self.over:
            return self.over[key]
        v = self.sheetdata(sname).get(r, c)
        if isinstance(v, Formula):
            if key in self.memo:
                m = self.memo[key]
                if m is _EVALUATING:
                    return XErr('#CIRC!')
                return m
            self.memo[key] = _EVALUATING
            try:
                res = self.eval_text(sname, v.text)
            except Exception:
                self.memo.pop(key, None)
                raise
            self.memo[key] = res
            return res
        return v

    def eval_text(self, sname, text, dr=0, dc=0):
        ast = self.parse(text)
        return self.top(self.ev(ast, (sname, dr, dc)))

    def eval_cell(self, sname, r, c):
        return self.cellval(sname, r, c)

    def top(self, v):
        if isinstance(v, RangeVal):
            if v.single():
                return self.cellval(v.sheet, v.r1, v.c1)
            return VALUE
        if isinstance(v, list):
            return v[0] if v else None
        return v

    def sv(self, v):
        """scalar value, link unwrapped to its label"""
        if isinstance(v, RangeVal):
            if not v.single():
                return VALUE
            v = self.cellval(v.sheet, v.r1, v.c1)
        if isinstance(v, LinkVal):
            v = v.label
        return v

    def keep(self, v):
        """scalar value keeping LinkVal"""
        if isinstance(v, RangeVal):
            if not v.single():
                return VALUE
            return self.cellval(v.sheet, v.r1, v.c1)
        return v

    def rows_of(self, rng):
        sd = self.sheetdata(rng.sheet)
        r1 = rng.r1 if rng.r1 is not None else 1
        r2 = rng.r2 if rng.r2 is not None else sd.max_row
        return r1, r2

    def range_values(self, rng, raw=False):
        r1, r2 = self.rows_of(rng)
        key = (rng.sheet, r1, r2, rng.c1, rng.c2, raw)
        v = self.range_cache.get(key)
        if v is None:
            f = self.raw if raw else self.cellval
            v = []
            for r in range(r1, r2 + 1):
                for c in range(rng.c1, rng.c2 + 1):
                    x = f(rng.sheet, r, c)
                    if not raw and isinstance(x, LinkVal):
                        x = x.label
                    v.append(x)
            self.range_cache[key] = v
        return v

    def arr(self, v):
        if isinstance(v, RangeVal):
            return list(self.range_values(v)) if not v.single() else [self.sv(v)]
        if isinstance(v, list):
            return v
        if isinstance(v, LinkVal):
            return [v.label]
        return [v]

    # --- evaluation
    def ev(self, node, ctx):
        t = node[0]
        if t in ('num', 'str', 'bool'):
            return node[1]
        if t == 'ref':
            sheet = node[1] or ctx[0]
            dr, dc = ctx[1], ctx[2]
            c1 = node[2] + (0 if node[3] else dc)
            r1 = None if node[4] is None else node[4] + (0 if node[5] else dr)
            c2 = node[6] + (0 if node[7] else dc)
            r2 = None if node[8] is None else node[8] + (0 if node[9] else dr)
            return RangeVal(sheet, r1, c1, r2, c2)
        if t == 'neg':
            v = self.ev(node[1], ctx)
            if isinstance(v, list):
                return [x if isinstance(x, XErr) else -to_num(x) for x in v]
            n = to_num(self.sv(v))
            return n if isinstance(n, XErr) else -n
        if t == 'bin':
            return self.binop(node[1], self.ev(node[2], ctx), self.ev(node[3], ctx))
        if t == 'func':
            f = getattr(self, 'f_' + node[1].replace('.', '_'), None)
            if f is None:
                self.unsupported[node[1]] += 1
                raise Unsupported('function %s' % node[1])
            return f(node[2], ctx)
        if t == 'missing':
            return None
        raise Unsupported(str(t))

    def binop(self, op, a, b):
        a = self.arr(a) if (isinstance(a, RangeVal) and not a.single()) else a
        b = self.arr(b) if (isinstance(b, RangeVal) and not b.single()) else b
        if isinstance(a, list) or isinstance(b, list):
            la = a if isinstance(a, list) else None
            lb = b if isinstance(b, list) else None
            n = len(la if la is not None else lb)
            return [self.binop(op, la[i] if la is not None else a, lb[i] if lb is not None else b) for i in range(n)]
        a = self.sv(a)
        b = self.sv(b)
        if isinstance(a, XErr):
            return a
        if isinstance(b, XErr):
            return b
        if op == '&':
            return to_text(a) + to_text(b)
        if op in Parser.CMP:
            c = xl_cmp(a, b)
            return {'=': c == 0, '<>': c != 0, '<': c < 0, '>': c > 0, '<=': c <= 0, '>=': c >= 0}[op]
        x, y = to_num(a), to_num(b)
        if isinstance(x, XErr):
            return x
        if isinstance(y, XErr):
            return y
        if op == '+':
            return x + y
        if op == '-':
            return x - y
        if op == '*':
            return x * y
        if op == '/':
            return DIV0 if y == 0 else x / y
        if op == '^':
            return x ** y
        raise Unsupported(op)

    def truth(self, v):
        v = self.sv(v)
        if isinstance(v, XErr):
            return v
        if v is None:
            return False
        if isinstance(v, bool):
            return v
        if is_num(v):
            return v != 0
        if isinstance(v, str):
            if v.upper() == 'TRUE':
                return True
            if v.upper() == 'FALSE':
                return False
            return VALUE
        return VALUE

    # --- functions (args are AST nodes)
    def f_HYPERLINK(self, a, ctx):
        loc = self.sv(self.ev(a[0], ctx))
        if isinstance(loc, XErr):
            return loc
        lab = self.sv(self.ev(a[1], ctx)) if len(a) > 1 else loc
        if isinstance(lab, XErr):
            return lab
        return LinkVal(lab, to_text(loc))

    def f_IFERROR(self, a, ctx):
        v = self.ev(a[0], ctx)
        k = self.keep(v) if not isinstance(v, list) else v
        s = k.label if isinstance(k, LinkVal) else k
        if isinstance(s, XErr):
            return self.ev(a[1], ctx)
        return k

    def f_IF(self, a, ctx):
        c = self.truth(self.ev(a[0], ctx))
        if isinstance(c, XErr):
            return c
        if c:
            return self.keep(self.ev(a[1], ctx)) if len(a) > 1 else True
        return self.keep(self.ev(a[2], ctx)) if len(a) > 2 else False

    def _logic_args(self, a, ctx):
        out = []
        for n in a:
            v = self.ev(n, ctx)
            if isinstance(v, RangeVal) and not v.single():
                for x in self.range_values(v):
                    if isinstance(x, (bool, int, float)):
                        out.append(bool(x))
                    elif isinstance(x, XErr):
                        out.append(x)
                continue
            if isinstance(v, RangeVal):
                x = self.sv(v)
                if x is None:
                    continue
                if isinstance(x, str):
                    continue
                out.append(self.truth(x))
                continue
            out.append(self.truth(v))
        return out

    def f_AND(self, a, ctx):
        vs = self._logic_args(a, ctx)
        for v in vs:
            if isinstance(v, XErr):
                return v
        return all(vs)

    def f_OR(self, a, ctx):
        vs = self._logic_args(a, ctx)
        for v in vs:
            if isinstance(v, XErr):
                return v
        return any(vs)

    def f_NOT(self, a, ctx):
        v = self.truth(self.ev(a[0], ctx))
        return v if isinstance(v, XErr) else (not v)

    def f_ISNUMBER(self, a, ctx):
        v = self.sv(self.ev(a[0], ctx))
        return is_num(v) or isinstance(v, (dt.datetime, dt.date))

    def f_LEFT(self, a, ctx):
        v = self.sv(self.ev(a[0], ctx))
        if isinstance(v, XErr):
            return v
        n = to_num(self.sv(self.ev(a[1], ctx))) if len(a) > 1 else 1
        if isinstance(n, XErr):
            return n
        return to_text(v)[:int(n)]

    def f_UPPER(self, a, ctx):
        v = self.sv(self.ev(a[0], ctx))
        return v if isinstance(v, XErr) else to_text(v).upper()

    def f_TRIM(self, a, ctx):
        v = self.sv(self.ev(a[0], ctx))
        return v if isinstance(v, XErr) else re.sub(' +', ' ', to_text(v)).strip(' ')

    def f_SUBSTITUTE(self, a, ctx):
        v, o, n = (self.sv(self.ev(x, ctx)) for x in a[:3])
        for x in (v, o, n):
            if isinstance(x, XErr):
                return x
        v, o, n = to_text(v), to_text(o), to_text(n)
        if o == '':
            return v
        if len(a) > 3:
            k = int(to_num(self.sv(self.ev(a[3], ctx))))
            idx = -1
            for _ in range(k):
                idx = v.find(o, idx + 1)
                if idx < 0:
                    return v
            return v[:idx] + n + v[idx + len(o):]
        return v.replace(o, n)

    def f_SEARCH(self, a, ctx):
        f = self.sv(self.ev(a[0], ctx))
        w = self.sv(self.ev(a[1], ctx))
        if isinstance(f, XErr):
            return f
        if isinstance(w, XErr):
            return w
        st = int(to_num(self.sv(self.ev(a[2], ctx)))) if len(a) > 2 else 1
        rx = wildcard_re(to_text(f))
        m = rx.search(to_text(w), st - 1)
        return VALUE if m is None else m.start() + 1

    def f_YEAR(self, a, ctx):
        v = self.sv(self.ev(a[0], ctx))
        if isinstance(v, XErr):
            return v
        if isinstance(v, (dt.datetime, dt.date)):
            return v.year
        n = to_num(v)
        if isinstance(n, XErr):
            return n
        return (EPOCH + dt.timedelta(days=n)).year

    def f_DATE(self, a, ctx):
        y, m, d = (int(to_num(self.sv(self.ev(x, ctx)))) for x in a)
        return dt.datetime(y, m, d)

    def f_ROUND(self, a, ctx):
        x = to_num(self.sv(self.ev(a[0], ctx)))
        n = int(to_num(self.sv(self.ev(a[1], ctx))))
        if isinstance(x, XErr):
            return x
        return round_half_away(x, n)

    def f_AVERAGE(self, a, ctx):
        vals = []
        for n in a:
            v = self.ev(n, ctx)
            if isinstance(v, RangeVal):
                for x in self.range_values(v):
                    if isinstance(x, XErr):
                        return x
                    if is_num(x):
                        vals.append(x)
            else:
                x = to_num(self.sv(v))
                if isinstance(x, XErr):
                    return x
                vals.append(x)
        return DIV0 if not vals else sum(vals) / len(vals)

    def f_SUMPRODUCT(self, a, ctx):
        arrs = [self.arr(self.ev(n, ctx)) for n in a]
        n = len(arrs[0])
        tot = 0
        for i in range(n):
            p = 1
            for ar in arrs:
                x = ar[i]
                if isinstance(x, XErr):
                    return x
                p *= x if is_num(x) else 0
            tot += p
        return tot

    def f_COUNTA(self, a, ctx):
        n = 0
        for node in a:
            v = self.ev(node, ctx)
            if isinstance(v, RangeVal):
                n += sum(1 for x in self.range_values(v, raw=True) if x is not None)
            elif v is not None:
                n += 1
        return n

    def _countif_one(self, rng, crit):
        key = (rng.sheet, rng.r1, rng.r2, rng.c1, rng.c2, repr(crit) if not isinstance(crit, LinkVal) else repr(crit.label))
        c = self.countif_cache.get(key)
        if c is None:
            p = crit_pred(crit)
            c = sum(1 for x in self.range_values(rng) if p(x))
            self.countif_cache[key] = c
        return c

    def f_COUNTIF(self, a, ctx):
        rng = self.ev(a[0], ctx)
        crit = self.ev(a[1], ctx)
        if not isinstance(rng, RangeVal):
            return VALUE
        if isinstance(crit, RangeVal) and not crit.single():
            return [self._countif_one(rng, x) for x in self.range_values(crit)]
        return self._countif_one(rng, self.sv(crit))

    def f_COUNTIFS(self, a, ctx):
        pairs = []
        for i in range(0, len(a), 2):
            rng = self.ev(a[i], ctx)
            crit = self.sv(self.ev(a[i + 1], ctx))
            pairs.append((self.range_values(rng), crit_pred(crit)))
        n = len(pairs[0][0])
        return sum(1 for k in range(n) if all(p(vals[k]) for vals, p in pairs))

    def _match_index(self, rng):
        r1, r2 = self.rows_of(rng)
        key = (rng.sheet, r1, r2, rng.c1)
        d = self.match_idx.get(key)
        if d is None:
            d = {}
            for i, x in enumerate(self.range_values(RangeVal(rng.sheet, r1, rng.c1, r2, rng.c1))):
                if isinstance(x, str):
                    k = ('s', x.casefold())
                elif isinstance(x, bool):
                    k = ('b', x)
                elif is_num(x):
                    k = ('n', float(x))
                else:
                    continue
                if k not in d:
                    d[k] = i + 1
            self.match_idx[key] = d
        return d

    def f_MATCH(self, a, ctx):
        val = self.sv(self.ev(a[0], ctx))
        rng = self.ev(a[1], ctx)
        mt = to_num(self.sv(self.ev(a[2], ctx))) if len(a) > 2 else 1
        if isinstance(val, XErr):
            return val
        if mt != 0:
            raise Unsupported('MATCH type %r' % mt)
        if not isinstance(rng, RangeVal):
            return NA
        if isinstance(val, str) and any(ch in val for ch in '*?~'):
            rx = wildcard_re(val)
            for i, x in enumerate(self.range_values(rng)):
                if isinstance(x, str) and rx.fullmatch(x):
                    return i + 1
            return NA
        if isinstance(val, str):
            k = ('s', val.casefold())
        elif isinstance(val, bool):
            k = ('b', val)
        elif is_num(val):
            k = ('n', float(val))
        elif val is None:
            return NA
        else:
            k = ('n', float(to_serial(val)))
        if rng.c1 != rng.c2:
            raise Unsupported('MATCH over 2D range')
        return self._match_index(rng).get(k, NA)

    def f_INDEX(self, a, ctx):
        rng = self.ev(a[0], ctx)
        row = to_num(self.sv(self.ev(a[1], ctx)))
        if isinstance(row, XErr):
            return row
        col = to_num(self.sv(self.ev(a[2], ctx))) if len(a) > 2 else 1
        if isinstance(col, XErr):
            return col
        if not isinstance(rng, RangeVal):
            return VALUE
        r1, r2 = self.rows_of(rng)
        row, col = int(row), int(col)
        if row < 1 or r1 + row - 1 > r2 or col < 1 or rng.c1 + col - 1 > rng.c2:
            return REF
        return RangeVal(rng.sheet, r1 + row - 1, rng.c1 + col - 1, r1 + row - 1, rng.c1 + col - 1)


_EVALUATING = object()


# ----------------------------------------------------------------------------------------------
# Conditional formatting engine (Excel precedence: rules by priority; per property the first TRUE
# rule that sets it wins; stopIfTrue stops lower rules; colour scales / data bars computed)
# ----------------------------------------------------------------------------------------------
def _ref_deps(ast, out):
    if isinstance(ast, tuple):
        if ast and ast[0] == 'ref':
            if not ast[3] or not ast[7]:
                out['col'] = True
            if (ast[4] is not None and not ast[5]) or (ast[8] is not None and not ast[9]):
                out['row'] = True
            return
        for x in ast[1:]:
            if isinstance(x, (tuple, list)):
                _ref_deps(x, out)
    elif isinstance(ast, list):
        for x in ast:
            _ref_deps(x, out)


def percentile_inc(vals, p):
    s = sorted(vals)
    if not s:
        return None
    k = (len(s) - 1) * p
    f = math.floor(k)
    c = min(f + 1, len(s) - 1)
    return s[f] + (s[c] - s[f]) * (k - f)


def _lerp(c1, c2, t):
    a = [int(c1[i:i + 2], 16) for i in (1, 3, 5)]
    b = [int(c2[i:i + 2], 16) for i in (1, 3, 5)]
    return '#' + ''.join('%02X' % half_up(a[i] + (b[i] - a[i]) * t) for i in range(3))


SCALE_DIFF = collections.Counter()   # max channel difference seen on colour-scale cells -> count


def color_close(a, b, tol=1):
    a, b = normc(a), normc(b)
    if a == b:
        SCALE_DIFF[0] += 1
        return True
    if not a or not b or not a.startswith('#') or not b.startswith('#'):
        return False
    d = max(abs(int(a[i:i + 2], 16) - int(b[i:i + 2], 16)) for i in (1, 3, 5))
    SCALE_DIFF[d] += 1
    return d <= tol


class CFResult:
    def __init__(self):
        self.props = {}         # (r,c) -> {'b','i','u','bg','fc'}
        self.scale = set()      # (r,c) whose bg came from a colour scale
        self.bars = {}          # (r,c) -> (p, color)
        self.true_cells = collections.Counter()   # priority -> n cells TRUE
        self.true_rows = collections.defaultdict(set)
        self.errors = []


def compute_cf(eng, sd, dxfs, row_limit=None):
    res = CFResult()
    rules = sorted(sd.xml.cf, key=lambda r: r['priority'])
    maxr = sd.max_row if row_limit is None else min(sd.max_row, row_limit)
    stopped = set()
    for rule in rules:
        cells_ranges = []
        for (r1, c1, r2, c2) in rule['ranges']:
            cells_ranges.append((r1, c1, min(r2, maxr), c2))
        pr = rule['priority']
        if rule['type'] == 'expression':
            if not rule['formulas']:
                continue
            ftxt = rule['formulas'][0]
            try:
                ast = eng.parse('=' + ftxt)
            except Exception as e:
                res.errors.append('rule p%d unparsable: %s (%s)' % (pr, ftxt, e))
                continue
            deps = {}
            _ref_deps(ast, deps)
            ar, ac = rule['ranges'][0][0], rule['ranges'][0][1]
            dxf = dxfs[rule['dxf']] if rule['dxf'] is not None and rule['dxf'] < len(dxfs) else {}
            cache = {}

            def truth(r, c):
                key = (r if deps.get('row') else 0, c if deps.get('col') else 0)
                v = cache.get(key)
                if v is None:
                    try:
                        x = eng.ev(ast, (sd.name, r - ar, c - ac))
                        x = eng.sv(x)
                        v = (x is True) or (is_num(x) and x != 0)
                    except (NotLoaded, Unsupported) as e:
                        res.errors.append('rule p%d at %s: %s' % (pr, coord(r, c), e))
                        v = False
                    cache[key] = v
                return v
            for (r1, c1, r2, c2) in cells_ranges:
                for r in range(r1, r2 + 1):
                    if not deps.get('col'):
                        if not truth(r, c1):
                            continue
                        cols = range(c1, c2 + 1)
                    else:
                        cols = [c for c in range(c1, c2 + 1) if truth(r, c)]
                    for c in cols:
                        res.true_cells[pr] += 1
                        res.true_rows[pr].add(r)
                        if (r, c) in stopped:
                            continue
                        p = res.props.get((r, c))
                        if p is None:
                            p = res.props[(r, c)] = {}
                        for k, v in dxf.items():
                            if k not in p:
                                p[k] = v
                        if rule['stop']:
                            stopped.add((r, c))
        elif rule['type'] in ('colorScale', 'dataBar'):
            cells = []
            for (r1, c1, r2, c2) in cells_ranges:
                for r in range(r1, r2 + 1):
                    for c in range(c1, c2 + 1):
                        try:
                            v = eng.sv(eng.cellval(sd.name, r, c))
                        except (NotLoaded, Unsupported):
                            v = None
                        if is_num(v):
                            cells.append((r, c, v))
            vals = [v for _, _, v in cells]
            if not vals:
                continue
            lo_, hi_ = min(vals), max(vals)

            def thr(cfvo):
                t, val = cfvo
                if t == 'min':
                    return lo_
                if t == 'max':
                    return hi_
                if t == 'num':
                    return float(val)
                if t == 'percentile':
                    return percentile_inc(vals, float(val) / 100.0)
                if t == 'percent':
                    return lo_ + (hi_ - lo_) * float(val) / 100.0
                raise Unsupported('cfvo %s' % t)
            th = [thr(x) for x in rule['cfvo']]
            cols_ = rule['colors']
            res.true_cells[pr] += len(cells)
            for r, c, v in cells:
                res.true_rows[pr].add(r)
                if (r, c) in stopped:
                    continue
                if rule['type'] == 'colorScale':
                    if len(th) == 3:
                        lo, mid, hi = th
                        if v <= lo:
                            col = cols_[0]
                        elif v >= hi:
                            col = cols_[2]
                        elif v <= mid:
                            col = _lerp(cols_[0], cols_[1], (v - lo) / (mid - lo) if mid > lo else 1)
                        else:
                            col = _lerp(cols_[1], cols_[2], (v - mid) / (hi - mid) if hi > mid else 1)
                    else:
                        lo, hi = th
                        col = cols_[0] if v <= lo else (cols_[1] if v >= hi else _lerp(cols_[0], cols_[1], (v - lo) / (hi - lo)))
                    p = res.props.setdefault((r, c), {})
                    if 'bg' not in p:
                        p['bg'] = col
                        res.scale.add((r, c))
                else:
                    lo, hi = th[0], th[-1]
                    pp = 0.0 if hi <= lo else max(0.0, min(1.0, (v - lo) / (hi - lo)))
                    res.bars.setdefault((r, c), (pp, cols_[0]))
        else:
            res.errors.append('unsupported rule type %s p%d' % (rule['type'], pr))
    return res


def check_cf_counts(rep, sid, res):
    exp = EXPECTED_CF_COUNTS.get(sid)
    if not exp:
        return
    rep.stats['cf_rule_true_counts'][sid] = {'p%d' % p: [res.true_cells.get(p, 0), len(res.true_rows.get(p, ()))] for p in sorted(res.true_cells)}
    for pr, (kind, n) in exp.items():
        got = res.true_cells.get(pr, 0) if kind == 'cells' else len(res.true_rows.get(pr, ()))
        if got != n:
            rep.warn('selfcheck_cf_count', sid, 'verifier CF engine TRUE-%s count for rule p%d differs from spec' % (kind, pr),
                     where='p%d' % pr, exp=n, act=got)
    for e in res.errors[:20]:
        rep.warn('selfcheck_cf_eval', sid, e)


# ----------------------------------------------------------------------------------------------
# JSON side helpers
# ----------------------------------------------------------------------------------------------
def parse_style(s):
    """table style string 'flags|bg|fc[|flags]' or dict -> normalised props"""
    if s is None:
        return {}
    if isinstance(s, dict):
        return {'b': bool(s.get('b') or s.get('bold')), 'i': bool(s.get('i') or s.get('italic')),
                'u': bool(s.get('u') or s.get('underline')), 'bg': normc(s.get('bg')), 'fc': normc(s.get('fc'))}
    parts = str(s).split('|')
    flags = parts[0] + (parts[3] if len(parts) > 3 else '')
    return {'b': 'b' in flags, 'i': 'i' in flags, 'u': 'u' in flags,
            'bg': normc(parts[1]) if len(parts) > 1 else None, 'fc': normc(parts[2]) if len(parts) > 2 else None}


def decode_act(a):
    """JSON cell -> (display value, link dict, url, fmt)"""
    if isinstance(a, dict):
        if 't' in a:
            return a.get('t'), a.get('l'), a.get('u'), a.get('f')
        if 'v' in a:
            return a.get('v'), a.get('l'), a.get('u'), a.get('f')
        return ('<unrecognised object %s>' % short(a, 60)), None, None, None
    return a, None, None, None


def val_eq(e, a):
    if isinstance(e, LinkVal):
        e = e.label
    if e is None or (isinstance(e, str) and e == ''):
        return a is None or a == ''
    if isinstance(e, bool):
        return isinstance(a, bool) and a == e
    if is_num(e):
        return is_num(a) and num_eq(e, a)
    if isinstance(e, str):
        return isinstance(a, str) and a == e
    return e == a


def classify(e, a, kind):
    if kind == 'datetime':
        return 'datetime'
    if isinstance(e, str) and isinstance(a, str):
        if e.strip() == a.strip():
            return 'whitespace'
        if e.startswith('='):
            return 'literal_eq_string'
    if (isinstance(e, str) and is_num(a)) or (is_num(e) and isinstance(a, str)):
        return 'type'
    if kind == 'formula':
        return 'formula_value'
    if (e is None or e == '') != (a is None or a == ''):
        return 'missing_or_extra_value'
    return 'value'


class Ctx:
    def __init__(self, book, eng, rep, man, data_dir, quick):
        self.book, self.eng, self.rep, self.man, self.data_dir, self.quick = book, eng, rep, man, data_dir, quick
        self.first_row = {}     # table sid -> first data row
        self.last_row = {}
        self.ncols = {}
        for n in book.names:
            sid = n[:2]
            af = book.xml[n].autofilter
            if expected_mode(sid) == 'table' and af:
                self.first_row[sid] = af[0] + 1
                self.last_row[sid] = af[2]
                self.ncols[sid] = af[3]
        self.json = {}          # sid -> merged json (non chunked)
        self.semantic = []      # (src sid, where, target sid, r, [(col, key)])
        self.all_links = []     # (src sid, where, link dict)
        self.big = {}           # results of workers

    def name_to_id(self, name):
        return name[:2] if name in self.book.names else None

    def resolve_loc(self, loc):
        m = re.match(r"^#'?(.+?)'?!\$?([A-Z]{1,3})\$?(\d+)$", loc or '')
        if not m:
            return {'url': loc}
        name, row = m.group(1), int(m.group(3))
        sid = self.name_to_id(name)
        if sid is None:
            return {'bad': loc}
        mode = expected_mode(sid)
        if mode == 'grid':
            r = row
        elif mode == 'card':
            r = None
        else:
            fr = self.first_row.get(sid, 5)
            r = None if row < fr else row - fr
        return {'s': sid, 'r': r}


def expected_value(ctx, sd, r, c, raw=None, nf=None):
    """-> (value, kind) where kind in plain/datetime/formula/unchecked"""
    if raw is None:
        raw = sd.get(r, c)
    if isinstance(raw, Formula):
        try:
            v = ctx.eng.cellval(sd.name, r, c)
        except (NotLoaded, Unsupported) as e:
            return str(e), 'unchecked'
        if isinstance(v, XErr):
            v = v.code
        if isinstance(v, (dt.datetime, dt.date)):
            v = fmt_datetime(v, nf or ctx.book.nf(sd, r, c))
        return v, 'formula'
    if isinstance(raw, (dt.datetime, dt.date, dt.time)):
        return fmt_datetime(raw, nf or ctx.book.nf(sd, r, c)), 'datetime'
    return raw, 'plain'


def compare_cell(ctx, exp, kind, act):
    """-> list of (cat, msg, exp, act)"""
    out = []
    av, al, au, af = decode_act(act)
    if isinstance(exp, LinkVal):
        tgt = ctx.resolve_loc(exp.loc)
        if not val_eq(exp.label, av):
            out.append(('link_label', 'link label differs', exp.label, av))
        if 'url' in tgt:
            if au != tgt['url']:
                out.append(('link_missing', 'external url link differs', tgt['url'], au))
        elif 'bad' in tgt:
            out.append(('link_target', 'formula points to unknown sheet', tgt['bad'], al))
        elif not isinstance(al, dict):
            out.append(('link_missing', 'expected link, JSON has none', tgt, act))
        elif str(al.get('s')) != tgt['s'] or al.get('r') != tgt['r']:
            out.append(('link_target', 'link target differs', tgt, al))
        return out
    if al is not None or au is not None:
        out.append(('link_unexpected', 'JSON has a link where Excel has none', exp, act))
    if not val_eq(exp, av):
        out.append((classify(exp, av, kind), 'value differs (%s)' % kind, exp, av))
    return out


_MATCH_RE = re.compile(r"MATCH\(\"([^\"]+)\",'([^']+)'!\$([A-Z]+):\$[A-Z]+,0\)")


SEMANTIC_FIXED = {('03', 'DB'): [(1, 'B')], ('11', 'W'): [(1, 'B'), (3, 'F')], ('26', 'K'): [(0, 'A')]}


def record_semantic(ctx, src_sid, where, raw, act, sd=None, r=None, c=None):
    """remember MATCH-derived / fixed-row link targets for a landing-row check against target JSON"""
    if not isinstance(raw, Formula):
        return
    av, al, au, af = decode_act(act)
    if not isinstance(al, dict) or al.get('r') is None:
        return
    fx = SEMANTIC_FIXED.get((src_sid, num2col(c) if c else None))
    if fx and sd is not None:
        keys = [(tc, sd.get(r, col2num(scol))) for tc, scol in fx]
        ctx.semantic.append((src_sid, where, str(al.get('s')), al.get('r'), keys, 'fixed-row link, same key as source row'))
        return
    m = _MATCH_RE.search(raw.text)
    if m:
        tsid = ctx.name_to_id(m.group(2))
        ctx.semantic.append((src_sid, where, str(al.get('s')), al.get('r'), [(col2num(m.group(3)) - 1, m.group(1))],
                             'MATCH(%s, %s!%s)' % (m.group(1), m.group(2)[:2], m.group(3))))


# ----------------------------------------------------------------------------------------------
# TABLE sheets
# ----------------------------------------------------------------------------------------------
def load_json(path):
    with open(path, encoding='utf-8') as f:
        return json.load(f)


def excel_col_px(width):
    return max(40, min(420, half_up(width * 7.5)))


def table_colmap(sid, ncols, jcols, xlabels, rep):
    """Excel col (1-based) -> JSON column index (or None if dropped)."""
    n = len(jcols)
    if n == ncols:
        return {c: c - 1 for c in range(1, ncols + 1)}
    drop = DROPPABLE_COLS.get(sid, set())
    if drop and n == ncols - len(drop):
        m, j = {}, 0
        for c in range(1, ncols + 1):
            if c in drop:
                m[c] = None
            else:
                m[c] = j
                j += 1
        rep.warn('approximation', sid, 'hidden parity column(s) %s ("組") dropped from JSON (allowed by spec §5; zebra must come from ui)'
                 % ','.join(num2col(c) for c in sorted(drop)))
        return m
    # fall back to label matching
    jl = [str(c.get('label')) for c in jcols]
    m, j = {}, 0
    for c in range(1, ncols + 1):
        lab = str(xlabels.get(c))
        if j < n and jl[j] == lab:
            m[c] = j
            j += 1
        else:
            m[c] = None
            rep.err('column_missing', sid, 'Excel column missing in JSON columns', where=num2col(c), exp=lab)
    if j < n:
        rep.err('column_extra', sid, 'JSON has %d columns not mapped to Excel columns' % (n - j), act=jl[j:])
    return m


def note_cell(sd):
    for c in range(3, (len(sd.rows.get(1) or [])) + 1):
        v = sd.get(1, c)
        if v is not None and not isinstance(v, Formula):
            return c
    return None


def expected_bands(ctx, sd, ncols, colmap=None):
    out = []
    covered = set()
    for (r1, c1, r2, c2) in sd.xml.merges:
        if r1 == 2 and r2 == 2 and c1 <= ncols:
            v = sd.get(2, c1)
            covered.update(range(c1, c2 + 1))
            if v is None or (isinstance(v, str) and not v.strip()):
                continue
            si = ctx.book.styles.get(sd.sid(2, c1))
            out.append({'label': v, 'from': c1 - 1, 'to': min(c2, ncols) - 1, 'bg': si.bg, 'fc': si.fc})
    for c in range(1, ncols + 1):
        if c in covered:
            continue
        v = sd.get(2, c)
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        si = ctx.book.styles.get(sd.sid(2, c))
        out.append({'label': v, 'from': c - 1, 'to': c - 1, 'bg': si.bg, 'fc': si.fc})
    if colmap:
        for b in out:
            js = [colmap.get(c) for c in range(b['from'] + 1, b['to'] + 2) if colmap.get(c) is not None]
            if js:
                b['from'], b['to'] = js[0], js[-1]
    return sorted(out, key=lambda b: b['from'])


def check_table_meta(ctx, sd, J, sid, ncols, colmap, col_nf, rep):
    """title / notes / header_cells / bands / columns / freeze"""
    title = sd.get(1, 2)
    if J.get('title') != title:
        rep.err('meta', sid, 'title differs from B1', exp=title, act=J.get('title'))
    nc = note_cell(sd)
    if nc:
        ntext = sd.get(1, nc)
        notes = J.get('notes')
        if isinstance(notes, str):
            notes = [notes]
        ok = isinstance(notes, list) and (notes == [ntext] or '｜'.join(str(x) for x in notes) == ntext)
        if not ok:
            rep.err('meta', sid, 'notes do not reproduce %s1 (split on ｜ allowed)' % num2col(nc), exp=ntext, act=notes)
    # header area extra cells (e.g. 07 R1..Y1 KPI)
    hc = J.get('header_cells') or []
    for h in hc:
        al = decode_act(h.get('v'))[1]
        if isinstance(al, dict):
            ctx.all_links.append((sid, 'header_cells r%s c%s' % (h.get('r'), h.get('c')), al))
    extras = []
    for r in (1, 2, 3):
        for c in range(1, ncols + 1):
            raw = sd.get(r, c)
            if raw is None or (r == 1 and c in (1, 2, nc)) or r in (2, 3):
                continue
            if (r, c) in sd.covered():
                continue
            extras.append((r, c, raw))
    for (r, c, raw) in extras:
        exp, kind = expected_value(ctx, sd, r, c)
        found = None
        for h in hc:
            hv = h.get('v')
            av, al, au, af = decode_act(hv)
            if val_eq(exp, av) and (h.get('c') in (c, c - 1)) and h.get('r', r) == r:
                found = h
                break
        if found is None:
            rep.err('header_cells', sid, 'title-area cell not represented in header_cells', where=coord(r, c), exp=exp,
                    act=[h.get('v') for h in hc][:12])
            continue
        # style / number format / merge span of the header cell
        si = ctx.book.styles.get(sd.sid(r, c)) if sd.sid(r, c) else None
        e_ = static_expect(si, isinstance(exp, LinkVal))
        sts = J.get('styles') or []
        s = found.get('s')
        a_ = parse_style(sts[s] if isinstance(s, int) and 0 <= s < len(sts) else s) if s is not None else {}
        e2 = {'b': e_['b'], 'i': e_['i'], 'u': e_['u'], 'bg': _nb(e_['bg']), 'fc': _nf_(e_['fc'])}
        a2 = {'b': bool(a_.get('b')), 'i': bool(a_.get('i')), 'u': bool(a_.get('u')), 'bg': _nb(a_.get('bg')), 'fc': _nf_(a_.get('fc'))}
        if e2 != a2:
            rep.err('header_cells', sid, 'header cell style differs', where=coord(r, c), exp=e2, act=a2)
        num = exp.label if isinstance(exp, LinkVal) else exp
        if is_num(num) and norm_nf(si.nf if si else None) != norm_nf(found.get('f') or decode_act(found.get('v'))[3]):
            rep.err('header_cells', sid, 'header cell number format differs', where=coord(r, c), exp=si.nf if si else None, act=found.get('f'))
        m = [mm for mm in sd.xml.merges if (mm[0], mm[1]) == (r, c)]
        ecs, ers = (m[0][3] - m[0][1] + 1, m[0][2] - m[0][0] + 1) if m else (1, 1)
        if (found.get('cs') or 1) != ecs or (found.get('rs') or 1) != ers:
            rep.err('header_cells', sid, 'header cell merge span differs', where=coord(r, c), exp=(ecs, ers), act=(found.get('cs'), found.get('rs')))
    if not any(isinstance(h.get('v'), dict) and (h['v'].get('l') or {}).get('s') == '00' for h in hc) \
            and not any(k in J for k in ('home', 'home_link')):
        rep.add('info', 'home_link_absent', sid, 'A1 "⌂ 目錄" link not in JSON (spec: optional, sidebar covers it)')
    # bands
    eb = expected_bands(ctx, sd, ncols, colmap)
    jb = J.get('bands') or []
    if len(eb) == 1 and eb[0]['from'] == 0 and eb[0]['label'] == sd.name and not jb:
        rep.add('info', 'bands', sid, 'row-2 full-width sheet-name bar not emitted as band (allowed by spec)')
    else:
        if [(b['label'], b['from'], b['to']) for b in eb] != [(b.get('label'), b.get('from'), b.get('to')) for b in jb]:
            rep.err('bands', sid, 'bands differ from row-2 merged ranges',
                    exp=[(b['label'], b['from'], b['to']) for b in eb],
                    act=[(b.get('label'), b.get('from'), b.get('to')) for b in jb])
        else:
            for e, a in zip(eb, jb):
                if normc(a.get('bg')) != e['bg'] or _nf_(a.get('fc')) != _nf_(e['fc']):
                    rep.err('bands', sid, 'band colour differs', where=e['label'], exp=(e['bg'], e['fc']), act=(a.get('bg'), a.get('fc')))
    # columns
    jcols = J.get('columns') or []
    hidden = sd.xml.hidden_cols()
    for c in range(1, ncols + 1):
        jc = colmap.get(c)
        if jc is None:
            continue
        col = jcols[jc]
        lab = sd.get(4, c)
        if col.get('label') != lab:
            rep.err('column_meta', sid, 'column label differs from row 4', where=num2col(c), exp=lab, act=col.get('label'))
        nt = sd.get(3, c)
        jn = col.get('note')
        if (nt or None) != (jn or None):
            rep.err('column_meta', sid, 'column note differs from row 3', where=num2col(c), exp=nt, act=jn)
        if (c in hidden) != bool(col.get('hidden')):
            rep.err('column_meta', sid, 'column hidden flag differs', where=num2col(c), exp=c in hidden, act=col.get('hidden'))
        cw = sd.xml.col_width(c)
        if cw and cw['custom']:
            w = excel_col_px(cw['width'])
            if col.get('w') is None or abs(col.get('w') - w) > 1:
                rep.err('layout', sid, 'column width px differs (contract: w×7.5 clamp 40..420)', where=num2col(c), exp=w, act=col.get('w'))
        cnt = col_nf.get(c)
        if cnt:
            dom = cnt.most_common(1)[0][0]
            ex = norm_nf(dom)
            ac = norm_nf(col.get('fmt'))
            if ex != ac:
                if is_date_nf(dom):
                    if ac is not None:
                        rep.warn('column_meta', sid, 'date column fmt differs', where=num2col(c), exp=dom, act=col.get('fmt'))
                else:
                    rep.err('column_meta', sid, 'column fmt differs from Excel number_format', where=num2col(c), exp=dom, act=col.get('fmt'))
    fz = J.get('freeze_cols')
    if fz != sd.xml.xsplit:
        rep.err('meta', sid, 'freeze_cols differs from pane xSplit', exp=sd.xml.xsplit, act=fz)


def json_table_style_layers(J, colmap):
    styles = J.get('styles') or []
    pst = [parse_style(s) for s in styles]

    def sty(x):
        if isinstance(x, int) and not isinstance(x, bool):
            return pst[x] if 0 <= x < len(pst) else {'bad': x}
        return parse_style(x)
    cs = {}
    for e in J.get('cell_styles') or []:
        cs[(e[0], e[1])] = sty(e[2])
    rs = {}
    for e in J.get('row_styles') or []:
        rs[e[0]] = sty(e[1])
    cp = []
    for col in J.get('columns') or []:
        cp.append({'b': bool(col.get('bold') or col.get('b')), 'i': bool(col.get('italic') or col.get('i')),
                   'u': bool(col.get('underline') or col.get('u')),
                   'bg': normc(col.get('bg')), 'fc': normc(col.get('fc'))})
    return cs, rs, cp


def _nb(c):
    c = normc(c)
    return None if c == '#FFFFFF' else c


def _nf_(c):
    c = normc(c)
    return None if c == '#000000' else c


LINK_BLUE = '#0563C1'
BAR_TOL = 0.00005 + 1e-9   # bar fraction may be rounded to 4 decimals (sub-pixel)


def static_expect(si, is_link):
    """Excel static style of a table data cell as the page must show it.
    Link cells: the page draws links in its own link style, so Excel's Hyperlink font (blue #0563C1, underline)
    is not carried; any other font colour on a link cell is."""
    if si is None:
        return {'b': False, 'i': False, 'u': False, 'bg': None, 'fc': None}
    fc, u = si.fc, bool(si.u)
    if is_link:
        u = False
        if normc(fc) == LINK_BLUE:
            fc = None
    return {'b': bool(si.b), 'i': bool(si.i), 'u': u, 'bg': si.bg, 'fc': fc}


def eff_json(cs, rs, cp, i, jc):
    layers = [x for x in (cs.get((i, jc)), rs.get(i), cp[jc] if jc < len(cp) else None) if x]
    out = {'b': any(x.get('b') for x in layers), 'i': any(x.get('i') for x in layers),
           'u': any(x.get('u') for x in layers), 'bg': None, 'fc': None}
    for k in ('bg', 'fc'):
        for x in layers:
            if x.get(k):
                out[k] = x[k]
                break
    return out


def check_table_styles(ctx, sd, J, sid, cfres, first, last, ncols, colmap, rows_iter, rep, row_offset=0, link_cells=(), blank_cells=()):
    cs, rs, cp = json_table_style_layers(J, colmap)
    st = ctx.book.styles
    nbad = collections.Counter()
    nst = len(J.get('styles') or [])
    nrows_j = len(J.get('rows') or [])
    ncols_j = len(J.get('columns') or [])
    for key, lst, shape in (('cell_styles', J.get('cell_styles') or [], 3), ('row_styles', J.get('row_styles') or [], 2)):
        for e in lst:
            bad = len(e) != shape or not (0 <= e[0] < nrows_j) or (shape == 3 and not (0 <= e[1] < ncols_j))                 or (isinstance(e[-1], int) and not (0 <= e[-1] < nst))
            if bad:
                rep.err('style_index', sid, '%s entry out of range' % key, act=e)
    ncf, ncf_ok = 0, 0
    for r in rows_iter:
        i = r - first - row_offset
        for c in range(1, ncols + 1):
            jc = colmap.get(c)
            if jc is None:
                continue
            sid_ = sd.sid(r, c)
            si = st.get(sid_) if sid_ else None
            exp = static_expect(si, (r, c) in link_cells)
            cf = cfres.props.get((r, c)) if cfres else None
            if cf:
                for k in ('b', 'i', 'u', 'bg', 'fc'):
                    if k in cf:
                        exp[k] = cf[k]
            act = eff_json(cs, rs, cp, i, jc)
            probs = []
            empty = sd.get(r, c) is None or (r, c) in blank_cells
            if exp['b'] != act['b'] and not empty:
                probs.append('b')
            if exp['i'] != act['i'] and not empty:
                probs.append('i')
            if exp['u'] != act['u'] and not empty:
                probs.append('u')
            eb, ab = _nb(exp['bg']), _nb(act['bg'])
            if cfres and (r, c) in cfres.scale:
                scale_ok = color_close(eb or '#FFFFFF', ab or '#FFFFFF')
            if eb != ab and not (cfres and (r, c) in cfres.scale and scale_ok):
                probs.append('bg')
            if _nf_(exp['fc']) != _nf_(act['fc']) and not empty:
                probs.append('fc')
            if cf:
                ncf += 1
                ncf_ok += not probs
            if probs:
                cat = 'cf_style' if cf else 'static_style'
                for p in probs:
                    nbad[(num2col(c), p)] += 1
                rep.err(cat, sid, 'effective cell style differs (%s)' % ','.join(probs), where=coord(r, c),
                        exp=exp, act=act)
    rep.stats['cf_cells'][sid] = {'cells_with_cf_style': ncf, 'matching': ncf_ok}
    if nbad:
        rep.stats['style_mismatch_by_col'][sid] = {'%s:%s' % k: v for k, v in nbad.most_common(30)}
    # data bars
    exp_bars = {k: v for k, v in (cfres.bars.items() if cfres else []) if first <= k[0] <= last}
    jbars = {}
    for e in J.get('bars') or []:
        jbars[(e[0], e[1])] = (e[2], normc(e[3]) if len(e) > 3 else None)
    colbar = {jc: col.get('bar') for jc, col in enumerate(J.get('columns') or []) if col.get('bar')}
    if exp_bars:
        vals = {}
        for (r, c), (p, colr) in exp_bars.items():
            jc = colmap.get(c)
            if jc is None:
                continue
            got = jbars.get((r - first - row_offset, jc))
            if got is None:
                if jc in colbar:
                    continue
                if p > 0:
                    rep.err('cf_databar', sid, 'data bar missing', where=coord(r, c), exp=(round(p, 4), colr))
                continue
            gp, gc = got
            if abs(gp - p) > BAR_TOL:
                rep.err('cf_databar', sid, 'data bar length differs (p = (v-min)/(max-min), 4-decimal rounding allowed)', where=coord(r, c), exp=round(p, 6), act=gp)
            if gc and gc != colr:
                rep.err('cf_databar', sid, 'data bar colour differs', where=coord(r, c), exp=colr, act=gc)
        for jc, cb in colbar.items():
            c = [k for k, v in colmap.items() if v == jc][0]
            ev = [sd.get(r, c) for r in range(first, last + 1)]
            ev = [x for x in ev if is_num(x)]
            if normc(cb.get('color') or cb.get('c')) not in (None, list(exp_bars.values())[0][1]):
                rep.err('cf_databar', sid, 'column bar colour differs', where=num2col(c), exp=list(exp_bars.values())[0][1], act=cb)
            if ev and ('min' in cb and not num_eq(cb['min'], min(ev)) or 'max' in cb and not num_eq(cb['max'], max(ev))):
                rep.err('cf_databar', sid, 'column bar min/max differs', where=num2col(c), exp=(min(ev), max(ev)), act=cb)
        rep.add('info', 'cf_databar', sid, '%d data-bar cells expected; JSON has %d sparse bars%s' %
                (len(exp_bars), len(jbars), ' + column-level bar' if colbar else ''))
    elif jbars:
        rep.err('cf_databar', sid, 'JSON has bars but Excel has no data bar rule', act=len(jbars))


def check_subgrid(ctx, sd, J, sid, last, ncols, rep):
    exp = []
    for r in sorted(sd.rows):
        row = sd.rows[r]
        for ci, v in enumerate(row):
            c = ci + 1
            if v is None:
                continue
            if r > last or c > ncols:
                if (r, c) in sd.covered():
                    continue
                if r == 1:
                    continue      # row-1 merged note overflow
                exp.append((r, c))
    if not exp:
        return
    sg = J.get('subgrid') or J.get('side_grid') or J.get('side_tables') or J.get('subtables')
    if not sg:
        rep.err('subgrid_missing', sid, '%d Excel cells outside the table range are not emitted (no subgrid)' % len(exp),
                where='%s..%s' % (coord(*exp[0]), coord(*exp[-1])))
        return
    if not isinstance(sg, dict) or 'rows' not in sg:
        rep.warn('subgrid_unchecked', sid, 'sub-table present in non-grid shape; values not verified cell by cell', act=type(sg).__name__)
        rep.uncheck(sid, 'sub-table outside autofilter range present as %s; not verified cell-by-cell' % type(sg).__name__)
        return
    cells = {}
    for row in sg.get('rows') or []:
        for cell in row.get('cells') or []:
            cells[(row.get('r'), cell.get('c'))] = cell
    # find coordinate offset
    offs = [(0, 0)]
    og = sg.get('origin')
    if isinstance(og, dict) and og.get('r') and og.get('c'):
        offs.append((og['r'] - 1, og['c'] - 1))
    anchors = [x for x in exp if not (isinstance(sd.get(*x), str) and not sd.get(*x).strip())]
    r0, c0 = (anchors or exp)[0]
    v0, _ = expected_value(ctx, sd, r0, c0)
    for (rr, cc), cell in cells.items():
        av, al, au, af = decode_act(cell.get('v'))
        if val_eq(v0, av):
            offs.append((r0 - rr, c0 - cc))
    best = None
    for (dr, dc) in offs:
        ok = 0
        for (r, c) in exp[:50]:
            cell = cells.get((r - dr, c - dc))
            if cell is not None:
                ev, _ = expected_value(ctx, sd, r, c)
                av, al, au, af = decode_act(cell.get('v'))
                if val_eq(ev, av):
                    ok += 1
        if best is None or ok > best[0]:
            best = (ok, dr, dc)
    _, dr, dc = best
    if (dr, dc) != (0, 0):
        rep.add('info', 'subgrid', sid, 'subgrid uses offset coordinates (Excel r,c) = (r+%d, c+%d)' % (dr, dc))
    for (r, c) in exp:
        ev, kind = expected_value(ctx, sd, r, c)
        cell = cells.get((r - dr, c - dc))
        if cell is None:
            if isinstance(ev, str) and not ev.strip():
                rep.warn('subgrid', sid, 'whitespace-only cell outside table not emitted', where=coord(r, c), exp=ev)
            else:
                rep.err('subgrid', sid, 'cell outside table range missing from subgrid', where=coord(r, c), exp=ev)
            continue
        v = cell.get('v')
        act = {'t': v, 'l': cell.get('l')} if cell.get('l') else v
        for cat, msg, e, a in compare_cell(ctx, ev, kind, act):
            rep.err('subgrid_' + cat, sid, msg, where=coord(r, c), exp=e, act=a)
        si = ctx.book.styles.get(sd.sid(r, c))
        sty = grid_style_of(sg, cell)
        e_ = {'b': si.b, 'i': si.i, 'bg': _nb(si.bg), 'fc': _nf_(si.fc)}
        a_ = {'b': bool(sty.get('b')), 'i': bool(sty.get('i')), 'bg': _nb(sty.get('bg')), 'fc': _nf_(sty.get('fc'))}
        if e_ != a_:
            rep.err('subgrid_style', sid, 'sub-table cell style differs', where=coord(r, c), exp=e_, act=a_)
        lp = grid_layout_probs(si, sty, False)
        if lp:
            rep.err('subgrid_layout_style', sid, 'sub-table cell layout style differs (%s)' % ','.join(p[0] for p in lp),
                    where=coord(r, c), exp={p[0]: p[1] for p in lp}, act={p[0]: p[2] for p in lp})
        m = [mm for mm in sd.xml.merges if (mm[0], mm[1]) == (r, c)]
        ecs, ers = (m[0][3] - m[0][1] + 1, m[0][2] - m[0][0] + 1) if m else (1, 1)
        if (cell.get('cs') or 1) != ecs or (cell.get('rs') or 1) != ers:
            rep.err('subgrid_merge', sid, 'merge span differs', where=coord(r, c), exp=(ecs, ers), act=(cell.get('cs'), cell.get('rs')))
        if is_num(ev) and norm_nf(si.nf) and (cell.get('f') or sty.get('f') or sty.get('fmt')) != si.nf:
            rep.err('subgrid_numfmt', sid, 'sub-table number format not carried', where=coord(r, c), exp=si.nf, act=cell.get('f') or sty.get('f'))
    # styled empty cells inside the sub-table box (fills / borders) must be emitted too
    rr = [r - dr for (r, c) in exp]
    cc = [c - dc for (r, c) in exp]
    box = (min(rr) + dr, min(cc) + dc, max(rr) + dr, max(cc) + dc)
    for r in range(box[0], box[2] + 1):
        for c in range(box[1], box[3] + 1):
            if sd.get(r, c) is not None or (r, c) in sd.covered():
                continue
            si = ctx.book.styles.get(sd.sid(r, c)) if sd.sid(r, c) else None
            if si is not None and (_nb(si.bg) or si.bd):
                cell = cells.get((r - dr, c - dc))
                if cell is None:
                    rep.err('subgrid_styled_empty_missing', sid, 'empty filled/bordered cell of the sub-table not emitted',
                            where=coord(r, c), exp={'bg': si.bg, 'bd': si.bd})
                    continue
                sty = grid_style_of(sg, cell)
                if _nb(sty.get('bg')) != _nb(si.bg):
                    rep.err('subgrid_style', sid, 'empty cell fill differs', where=coord(r, c), exp=si.bg, act=sty.get('bg'))
                lp = grid_layout_probs(si, sty, True)
                if lp:
                    rep.err('subgrid_layout_style', sid, 'sub-table cell layout style differs (%s)' % ','.join(p[0] for p in lp),
                            where=coord(r, c), exp={p[0]: p[1] for p in lp}, act={p[0]: p[2] for p in lp})
    check_grid_overflow(sd, sid, sg, cells, sg.get('col_widths') or [], rep, ctx.book.styles, where_off=(dr, dc))


class ColAgg:
    """per Excel column aggregate over non-empty data cells (value classes, alignment, wrap, font)"""
    __slots__ = ('n', 'cls', 'ha', 'wrap', 'mono', 'gen_cls')

    def __init__(self):
        self.n = 0
        self.cls = collections.Counter()
        self.ha = collections.Counter()
        self.wrap = 0
        self.mono = 0
        self.gen_cls = collections.Counter()   # value classes of cells with general alignment

    def add(self, exp, kind, si):
        self.n += 1
        if isinstance(exp, LinkVal):
            k = 'link_num' if is_num(exp.label) else 'link'
        elif kind == 'datetime':
            k = 'date'
        elif isinstance(exp, bool):
            k = 'bool'
        elif is_num(exp):
            k = 'num'
        else:
            k = 'text'
        self.cls[k] += 1
        ha = XL_HA.get(si.ha if si else None)
        self.ha[ha] += 1
        if ha is None:
            self.gen_cls[k] += 1
        if si is not None and si.wrap:
            self.wrap += 1
        if si is not None and (si.font or '').lower() in MONO_FONTS:
            self.mono += 1

    def to_state(self):
        return [self.n, dict(self.cls), [[k, v] for k, v in self.ha.items()], self.wrap, self.mono, dict(self.gen_cls)]

    @classmethod
    def from_state(cls, s):
        a = cls()
        a.n, a.cls, a.ha, a.wrap, a.mono, a.gen_cls = s[0], collections.Counter(s[1]), \
            collections.Counter({k: v for k, v in s[2]}), s[3], s[4], collections.Counter(s[5])
        return a


def check_table_columns(ctx, sd, J, sid, ncols, colmap, agg, rep):
    """columns[].type / align / wrap / mono and header colours hbg/hfc (row 4) against Excel"""
    jcols = J.get('columns') or []
    hr = ctx.first_row[sid] - 1
    st = ctx.book.styles
    for c in range(1, ncols + 1):
        jc = colmap.get(c)
        if jc is None or jc >= len(jcols):
            continue
        col = jcols[jc]
        w = num2col(c)
        hs = st.get(sd.sid(hr, c)) if sd.sid(hr, c) else None
        if _nb(col.get('hbg')) != _nb(hs.bg if hs else None):
            rep.err('column_meta', sid, 'header fill (hbg) differs from row %d' % hr, where=w, exp=hs.bg if hs else None, act=col.get('hbg'))
        if _nf_(col.get('hfc')) != _nf_(hs.fc if hs else None):
            rep.err('column_meta', sid, 'header font colour (hfc) differs from row %d' % hr, where=w, exp=hs.fc if hs else None, act=col.get('hfc'))
        a = agg.get(c)
        if a is None or a.n == 0:
            continue
        n = float(a.n)
        fr = {'num': (a.cls['num'] + a.cls['link_num']) / n, 'date': a.cls['date'] / n,
              'link': a.cls['link'] / n, 'bool': a.cls['bool'] / n}
        t = col.get('type') or 'text'
        ok = {'num': fr['num'] >= 0.95, 'date': fr['date'] >= 0.95, 'link': fr['link'] >= 0.5,
              'bool': fr['bool'] == 1, 'text': True}.get(t, False)
        forced = next((k for k in ('bool', 'date', 'link', 'num') if fr[k] == 1), None)
        if not ok or (forced and t != forced):
            rep.err('column_type', sid, 'columns[].type inconsistent with the Excel values', where=w,
                    exp=forced or {k: round(v, 3) for k, v in fr.items()}, act=t)
        # alignment: majority of the non-empty cells; general -> Excel's value-type rule
        top, cnt = a.ha.most_common(1)[0]
        if top is None:
            # Excel General: numbers / dates right, booleans centred, text (incl. text-label links) left, per cell.
            # A homogeneous column gets that side; a mixed one must leave align unset (null) so the page applies
            # the same per-value rule (table.js: no align -> number right, boolean centre, text left).
            # Dates are pre-formatted strings in JSON, so with align unset the page would put them left:
            # accept whichever setting renders the most cells on the same side as Excel.
            g = a.gen_cls
            score = {'left': g['text'] + g['link'], 'right': g['num'] + g['date'] + g['link_num'], 'center': g['bool'],
                     None: g['num'] + g['link_num'] + g['bool'] + g['text'] + g['link']}
            best = max(score.values())
            ok_al = {k for k, v in score.items() if v == best}
            ea = None if None in ok_al else sorted(ok_al)[0]
        else:
            ok_al = {top}
            ea = top
        if (col.get('align') or None) not in ok_al:
            rep.err('column_align', sid, 'columns[].align differs from Excel alignment (%s)' % ('general -> by value type' if top is None else 'explicit'),
                    where=w, exp=ea, act=col.get('align'))
        if bool(col.get('wrap')) != (a.wrap * 2 > a.n):
            rep.err('column_meta', sid, 'columns[].wrap differs from the majority wrapText of the data cells', where=w,
                    exp=a.wrap * 2 > a.n, act=col.get('wrap'))
        if bool(col.get('mono')) != (a.mono * 2 > a.n):
            rep.err('column_meta', sid, 'columns[].mono differs from the majority font (Consolas…) of the data cells', where=w,
                    exp=a.mono * 2 > a.n, act=col.get('mono'))


UI_INDEX_LISTS = ('facets', 'facets_extra', 'search_cols', 'pre_cols')
UI_INDEX_ONE = ('key_col', 'alias_col', 'tag_col', 'link_col', 'count_col', 'prio_col', 'cmx_col', 'group_col', 'summary_col')


def _dv_list_for(sd, excel_col, first):
    for d in sd.xml.dv:
        if d['type'] != 'list' or not d['formula1']:
            continue
        for part in d['sqref'].split():
            r1, c1, r2, c2 = parse_area(part)
            if c1 <= excel_col <= c2 and r2 >= first:
                f = d['formula1']
                if f.startswith('"'):
                    return f.strip('"').split(',')
                return ('ref', f)
    return None


def _sort_key(v):
    if v is None or v == '':
        return (2, '')
    if is_num(v):
        return (0, v)
    return (1, str(v))


def check_ui(ctx, sd, J, sid, first, last, ncols, colmap, rep, colvals=None):
    """ui.* hints: indexes in range, and every value-dependent hint consistent with the (verified) data.
    colvals(jc) -> list of display values of JSON column jc over all rows."""
    ui = J.get('ui') or {}
    jcols = J.get('columns') or []
    n = len(jcols)
    rows = J.get('rows') or []
    if colvals is None:
        cache = {}

        def colvals(jc):
            if jc not in cache:
                cache[jc] = [decode_act(r[jc])[0] if jc < len(r) else None for r in rows]
            return cache[jc]

    def inr(i):
        return isinstance(i, int) and not isinstance(i, bool) and 0 <= i < n
    for k in UI_INDEX_LISTS:
        v = ui.get(k)
        if v is not None and not (isinstance(v, list) and all(inr(i) for i in v)):
            rep.err('ui', sid, 'ui.%s has an out-of-range column index' % k, act=v)
    for k in UI_INDEX_ONE:
        if k in ui and ui[k] is not None and not inr(ui[k]):
            rep.err('ui', sid, 'ui.%s out of range' % k, act=ui[k])
    for k in ('order', 'facet_options', 'sort_key', 'facet_values'):
        for key in (ui.get(k) or {}):
            if not inr(int(key)):
                rep.err('ui', sid, 'ui.%s key out of range' % k, act=key)
    # alias column really holds device aliases (03 column B)
    al = ui.get('alias_col')
    if inr(al):
        s03 = ctx.book.sheet('03')
        aliases = {s03.get(r, 2) for r in range(ctx.first_row['03'], ctx.last_row['03'] + 1)} if s03 else set()
        vals = [v for v in colvals(al) if v not in (None, '')]
        # spec 19 F: 52 rows hold several aliases joined by ';'
        hit = sum(1 for v in vals if v in aliases or (isinstance(v, str) and all(x in aliases for x in v.split(';'))))
        if aliases and vals and hit < len(vals):
            (rep.err if hit < 0.9 * len(vals) else rep.warn)('ui', sid, 'ui.alias_col values not all 03 device aliases', where=jcols[al].get('label'),
                                                           exp=len(vals), act=hit)
        rep.stats['ui_alias_hits'][sid] = [hit, len(vals)]
    lc = ui.get('link_col')
    if inr(lc) and (jcols[lc].get('type') != 'link'):
        rep.err('ui', sid, 'ui.link_col is not a link column', where=jcols[lc].get('label'), act=jcols[lc].get('type'))
    # custom sort order must rank every value of the column
    for key, order in (ui.get('order') or {}).items():
        vals = {v for v in colvals(int(key)) if v not in (None, '')}
        miss = vals - set(order)
        if miss:
            rep.err('ui', sid, 'ui.order does not cover every value', where=jcols[int(key)].get('label'), exp=sorted(map(str, vals)), act=order)
    # validation lists (input columns) straight from Excel <dataValidation>
    inv = {v: k for k, v in colmap.items() if v is not None}
    for key, opts in (ui.get('facet_options') or {}).items():
        exp = _dv_list_for(sd, inv.get(int(key)), first)
        if exp != opts:
            rep.err('ui', sid, 'ui.facet_options differ from the Excel data-validation list', where=jcols[int(key)].get('label'), exp=exp, act=opts)
    for jc in range(n):
        dvl = _dv_list_for(sd, inv.get(jc), first) if inv.get(jc) else None
        if isinstance(dvl, list) and (ui.get('facet_options') or {}).get(str(jc)) is None and not any(
                jc in (v.get('cols') or []) and v.get('options') == dvl for v in (J.get('validations') or [])):
            rep.err('ui', sid, 'Excel data-validation list not carried (ui.facet_options / validations)', where=jcols[jc].get('label'), exp=dvl)
    # every Excel data validation on the data rows is carried as validations[] (informational; page is read-only)
    exp_v = []
    for d in sd.xml.dv:
        for part in d['sqref'].split():
            r1, c1, r2, c2 = parse_area(part)
            if r2 < first:
                continue
            cols = [colmap[c] for c in range(c1, c2 + 1) if colmap.get(c) is not None]
            if not cols:
                continue
            e = {'cols': cols, 'type': d['type']}
            f = d['formula1']
            if d['type'] == 'list' and f:
                if f.startswith('"'):
                    e['options'] = f.strip('"').split(',')
                else:
                    e['source'] = f
            elif f:
                e['formula'] = f
                if d['operator']:
                    e['operator'] = d['operator']
            if d['prompt']:
                e['prompt'] = d['prompt']
            exp_v.append(e)
    if exp_v != (J.get('validations') or []):
        rep.err('ui', sid, 'validations differ from the Excel <dataValidation> list', exp=exp_v, act=J.get('validations'))
    # summary rows (15 subtotal rows)
    if 'summary_rows' in ui:
        sc, sv = ui.get('summary_col'), set(ui.get('summary_values') or [])
        exp = [i for i, v in enumerate(colvals(sc)) if v in sv] if inr(sc) else None
        if exp != ui['summary_rows'] or not exp:
            rep.err('ui', sid, 'ui.summary_rows differ from the rows whose summary column is %s' % sorted(sv), exp=exp, act=ui['summary_rows'])
    # sort_key: sorting by the key column must order the display column
    for key, kc in (ui.get('sort_key') or {}).items():
        pairs = sorted(zip(colvals(kc), colvals(int(key))), key=lambda p: _sort_key(p[0]))
        # only real timestamps must come out ordered; text placeholders (e.g. 10 F '（占位 1970）') are the
        # reason the key exists and may sit anywhere
        seq = [p[1] for p in pairs if isinstance(p[1], str) and re.match(r'^\d{4}-\d\d-\d\d', p[1])]
        nokey = sum(1 for k_, v_ in pairs if v_ not in (None, '') and k_ in (None, ''))
        if nokey:
            rep.err('ui', sid, 'ui.sort_key: rows with a display value but no sort key', act=nokey)
        bad = sum(1 for x, y in zip(seq, seq[1:]) if _sort_key(y) < _sort_key(x))
        if bad:
            rep.err('ui', sid, 'ui.sort_key: ordering by %s does not order %s' % (jcols[kc].get('label'), jcols[int(key)].get('label')), act=bad)
    # precomputed facet lists (chunked tables): every distinct value over ALL parts, first-appearance order
    if 'facet_values' in ui or J.get('part'):
        fv = ui.get('facet_values') or {}
        for fc in list(ui.get('facets') or []) + list(ui.get('facets_extra') or []):
            if not inr(fc):
                continue
            seen = {}
            for v in colvals(fc):
                if v is not None and v not in seen:
                    seen[v] = 1
            exp = list(seen) if len(seen) <= 500 else None
            act = fv.get(str(fc))
            if exp != act and not (exp is not None and act is not None and len(exp) == len(act) and all(val_eq(x, y) for x, y in zip(exp, act))):
                rep.err('ui', sid, 'ui.facet_values differ from the distinct values of the whole table', where=jcols[fc].get('label'),
                        exp=exp, act=act)
    for p in ui.get('presets') or []:
        if not inr(p.get('col')):
            rep.err('ui', sid, 'ui.presets column out of range', act=p)
        elif 'eq' in p and p['eq'] not in set(colvals(p['col'])):
            rep.err('ui', sid, 'ui.presets eq value does not occur in the column', act=p)


def verify_table(ctx, sd, sid, J, rep, cfres):
    first, last, ncols = ctx.first_row[sid], ctx.last_row[sid], ctx.ncols[sid]
    xlabels = {c: sd.get(4, c) for c in range(1, ncols + 1)}
    jcols = J.get('columns') or []
    colmap = table_colmap(sid, ncols, jcols, xlabels, rep)
    rows = J.get('rows') or []
    nx = last - first + 1
    if len(rows) != nx:
        rep.err('row_count', sid, 'JSON row count differs from Excel data rows (autofilter %d..%d)' % (first, last), exp=nx, act=len(rows))
    # key column sanity: rows after autofilter must not hold table data
    col_nf = collections.defaultdict(collections.Counter)
    st = ctx.book.styles
    rng = range(first, last + 1)
    if ctx.quick:
        rng = list(range(first, min(last, first + 300) + 1)) + list(range(max(first, last - 50), last + 1))
    ncmp = 0
    link_cells = set()
    blank_cells = set()
    agg = collections.defaultdict(ColAgg)
    for r in rng:
        i = r - first
        jr = rows[i] if i < len(rows) else None
        if jr is None:
            continue
        if not isinstance(jr, list):
            rep.err('row_shape', sid, 'row is not an array', where='row %d' % r, act=jr)
            continue
        for c in range(1, ncols + 1):
            jc = colmap.get(c)
            raw = sd.get(r, c)
            if raw is not None:
                nf = st.get(sd.sid(r, c)).nf
                col_nf[c][nf] += 1
            else:
                nf = None
            if jc is None:
                continue
            act = jr[jc] if jc < len(jr) else None
            if raw is None and act is None:
                continue
            ncmp += 1
            exp, kind = expected_value(ctx, sd, r, c, raw, nf)
            if kind == 'unchecked':
                rep.uncheck(sid, '%s: %s' % (coord(r, c), exp))
                continue
            for cat, msg, e, a in compare_cell(ctx, exp, kind, act):
                rep.err(cat, sid, msg, where=coord(r, c), exp=e, act=a)
            if isinstance(exp, LinkVal):
                link_cells.add((r, c))
            if raw is not None and (exp is None or exp == ''):
                blank_cells.add((r, c))       # formula showing nothing: fonts are invisible there
            if not (exp is None or exp == ''):
                agg[c].add(exp, kind, st.get(sd.sid(r, c)) if sd.sid(r, c) else None)
            # effective number format of numeric cells (cell {"v","f"} override, else column fmt)
            num = exp.label if isinstance(exp, LinkVal) else exp
            if is_num(num):
                av_, al_, au_, af_ = decode_act(act)
                jf = af_ if af_ is not None else (jcols[jc].get('fmt') if jc < len(jcols) else None)
                if norm_nf(jf) != norm_nf(nf or st.get(sd.sid(r, c)).nf):
                    rep.err('cell_numfmt', sid, 'effective number format of numeric cell differs', where=coord(r, c),
                            exp=nf or st.get(sd.sid(r, c)).nf, act=jf)
            av, al, au, af = decode_act(act)
            if al is not None:
                ctx.all_links.append((sid, coord(r, c), al))
                record_semantic(ctx, sid, coord(r, c), raw, act, sd, r, c)
    rep.stats['cells_compared'][sid] = ncmp
    check_table_columns(ctx, sd, J, sid, ncols, colmap, agg, rep)
    check_ui(ctx, sd, J, sid, first, last, ncols, colmap, rep)
    # (per-cell number formats are checked cell by cell above: effective fmt = {"v","f"} override or column fmt)
    check_table_meta(ctx, sd, J, sid, ncols, colmap, col_nf, rep)
    check_table_styles(ctx, sd, J, sid, cfres, first, last, ncols, colmap, rng, rep, link_cells=link_cells, blank_cells=blank_cells)
    check_subgrid(ctx, sd, J, sid, last, ncols, rep)
    return colmap


# ----------------------------------------------------------------------------------------------
# GRID sheets
# ----------------------------------------------------------------------------------------------
def grid_style_of(J, cell):
    s = cell.get('s')
    if s is None:
        return {}
    if isinstance(s, int) and not isinstance(s, bool):
        st = J.get('styles') or []
        return st[s] if 0 <= s < len(st) else {'__bad__': s}
    if isinstance(s, dict):
        return s
    if isinstance(s, str):
        return parse_style(s)
    return {}


MONO_FONTS = {'consolas', 'courier new', 'courier', 'lucida console', 'cascadia mono', 'cascadia code', 'menlo'}
# Excel alignment -> CSS-ish value the page must reproduce (None = general / not set)
XL_HA = {None: None, 'general': None, 'left': 'left', 'center': 'center', 'right': 'right', 'justify': 'justify',
         'fill': 'left', 'centerContinuous': 'center', 'distributed': 'center'}
# Excel default vertical alignment is bottom; the page must default to bottom too (grid.js: VA[..] || 'bottom')
XL_VA = {None: 'bottom', 'bottom': 'bottom', 'top': 'top', 'center': 'middle', 'justify': 'middle', 'distributed': 'middle'}


def _bd_norm(bd):
    """border spec -> {side: (style, colour)}; accepts {'t': 'thin #BFBFBF'} or 'thin' / 'l:thick:#BF8F00' strings"""
    out = {}
    if not bd:
        return out
    if isinstance(bd, str):
        for part in bd.split(';'):
            p = [x for x in re.split(r'[:\s]+', part.strip()) if x]
            if not p:
                continue
            if p[0] in ('t', 'r', 'b', 'l'):
                out[p[0]] = (p[1] if len(p) > 1 else 'thin', normc(p[2]) if len(p) > 2 else None)
            else:
                for k in 'trbl':
                    out[k] = (p[0], normc(p[1]) if len(p) > 1 else None)
        return out
    for k, v in bd.items():
        k = k[0]
        if isinstance(v, str):
            p = v.split()
            out[k] = (p[0], normc(p[1]) if len(p) > 1 else None)
        elif isinstance(v, dict):
            out[k] = (v.get('style') or v.get('s'), normc(v.get('color') or v.get('c')))
    return out


def grid_layout_probs(si, sty, empty):
    """font size / alignment / wrap / borders / monospace / indent / rotation / strike of a grid cell.
    -> list of (prop, expected, actual).  Font-only props are ignored on empty cells."""
    probs = []
    if si is None:
        e = {'sz': 11.0, 'ha': None, 'va': 'bottom', 'wrap': False, 'mono': False, 'ind': 0, 'rot': 0, 's': False, 'bd': {},
             'shrink': False}
    else:
        e = {'sz': float(si.sz or 11), 'ha': XL_HA.get(si.ha, si.ha), 'va': XL_VA.get(si.va, si.va), 'wrap': bool(si.wrap),
             'mono': (si.font or '').lower() in MONO_FONTS, 'ind': si.ind or 0, 'rot': si.rot or 0, 's': bool(si.strike),
             'bd': {k: v for k, v in si.bd.items()}, 'shrink': bool(si.shrink)}
    a = {'sz': float(sty.get('sz') or 11), 'ha': sty.get('ha') or None, 'va': sty.get('va') or 'bottom',
         'wrap': bool(sty.get('wrap')), 'mono': bool(sty.get('mono')), 'ind': sty.get('ind') or 0,
         'rot': sty.get('rot') or 0, 's': bool(sty.get('s') or sty.get('strike')), 'bd': _bd_norm(sty.get('bd')),
         'shrink': bool(sty.get('shrink'))}
    for k in ('ha', 'va', 'wrap', 'ind', 'rot', 'bd', 'shrink'):
        if e[k] != a[k]:
            probs.append((k, e[k], a[k]))
    if not empty:
        for k in ('sz', 'mono', 's'):
            if e[k] != a[k]:
                probs.append((k, e[k], a[k]))
    return probs


def text_px_estimate(s, sz_pt):
    """rough rendered width of the first line: CJK full-width = 1 em, others ~0.5 em"""
    import unicodedata
    em = (sz_pt or 11) * 4.0 / 3.0
    w = 0.0
    for ch in s.split('\n', 1)[0]:
        w += em if unicodedata.east_asian_width(ch) in ('W', 'F') else em * 0.5
    return w


def check_grid_overflow(sd, sid, J, jcells, colw, rep, st, where_off=(0, 0)):
    """ov (text overflow into empty right neighbours, CONTRACT v2 grid):
       hard invariants -> errors; width heuristic -> warnings; Excel overflow the page cannot show -> counted."""
    dr, dc = where_off
    merges_cov = sd.covered()
    anchors = {(m[0], m[1]) for m in sd.xml.merges}
    maxc = len(colw)
    clipped = []
    n_ov = 0
    for (r, c), cell in jcells.items():
        R, C = r + dr, c + dc                      # Excel coordinates
        ov = cell.get('ov')
        v = cell.get('v')
        sty = grid_style_of(J, cell)
        si = st.get(sd.sid(R, C)) if sd.sid(R, C) else None
        if ov:
            n_ov += 1
            bad = []
            if not isinstance(v, str):
                bad.append('value is not text')
            if sty.get('wrap') or (si is not None and si.wrap):
                bad.append('wrapped text cannot overflow')
            if XL_HA.get(si.ha if si else None) not in (None, 'left'):
                bad.append('only general/left text overflows to the right')
            if cell.get('cs') or (R, C) in anchors:
                bad.append('merged cell cannot overflow')
            for k in range(1, ov + 1):
                if sd.get(R, C + k) is not None:
                    bad.append('neighbour %s is not empty in Excel' % coord(R, C + k))
                if (R, C + k) in merges_cov or (R, C + k) in anchors:
                    bad.append('neighbour %s is merged' % coord(R, C + k))
                if (r, c + k) in jcells:
                    bad.append('neighbour %s is an emitted (styled) cell' % coord(R, C + k))
                if c + k > maxc:
                    bad.append('overflow beyond col_widths')
            if bad:
                rep.err('grid_overflow', sid, 'invalid ov=%d: %s' % (ov, '; '.join(sorted(set(bad)))), where=coord(R, C), act=v)
        # heuristic: text that clearly needs more room than its column
        if not isinstance(v, str) or cell.get('cs') or (R, C) in anchors or (si is not None and si.wrap) \
                or XL_HA.get(si.ha if si else None) not in (None, 'left'):
            continue
        need = text_px_estimate(v, (si.sz if si else 11)) + 6
        have = colw[c - 1] if c - 1 < maxc else 64
        if need <= have * 1.15:
            if ov and need < have * 0.85:
                rep.warn('grid_overflow', sid, 'ov given although the text seems to fit its column', where=coord(R, C), exp=0, act=ov)
            continue
        nb_empty = sd.get(R, C + 1) is None and (R, C + 1) not in merges_cov and (R, C + 1) not in anchors and c < maxc
        if not nb_empty:
            continue       # Excel clips too (neighbour holds a value or is merged)
        if (r, c + 1) in jcells:
            clipped.append(coord(R, C))   # Excel spills over an empty *styled* neighbour; the page clips here
            continue
        if not ov:
            rep.warn('grid_overflow', sid, 'long unwrapped text next to an empty unstyled cell has no ov (clipped on the page)',
                     where=coord(R, C), exp='ov>=1', act=v[:40])
    rep.stats['grid_overflow'][sid] = {'cells_with_ov': n_ov, 'excel_spills_over_styled_empty_neighbour': len(clipped),
                                       'examples': clipped[:12]}


def render_sectioned(v, nf):
    """independent renderer for sectioned number formats like '+0.0;-0.0;0.0' (pos;neg;zero)"""
    secs = nf.split(';')
    if v > 0 or len(secs) == 1:
        sec, x = secs[0], v
    elif v < 0:
        sec, x = (secs[1], -v) if len(secs) > 1 else (secs[0], v)
    else:
        sec, x = (secs[2] if len(secs) > 2 else secs[0]), v
    m = re.match(r'^(.*?)([#0,]*0)(\.0+)?(%?)(.*)$', sec)
    if not m:
        return None
    pre, ip, dp, pct, post = m.groups()
    if pct:
        x = x * 100
    nd = len(dp) - 1 if dp else 0
    q = Decimal(repr(x)).quantize(Decimal(1).scaleb(-nd), rounding=ROUND_HALF_UP)
    s = ('{:,.%df}' if ',' in ip else '{:.%df}') % nd
    body = s.format(q)
    strip = lambda t: t.replace('"', '').replace('\\', '')
    return strip(pre) + body + pct + strip(post)


def check_grid_meta(ctx, sd, sid, J, rep):
    """grid-level keys: title / notes / subtitle / freeze / autofilter / gridlines / 00 toc+legend / 14 named_tables"""
    t = sd.get(1, 2)
    if isinstance(t, str) and J.get('title') != t:
        rep.err('grid_meta', sid, 'title differs from B1', exp=t, act=J.get('title'))
    note = next((sd.get(1, c) for c in range(3, (len(sd.rows.get(1) or [])) + 1) if isinstance(sd.get(1, c), str)), None)
    if note is not None and '｜'.join(J.get('notes') or []) != note:
        rep.err('grid_meta', sid, 'notes do not reproduce the row-1 note (split on ｜)', exp=note, act=J.get('notes'))
    sub = sd.get(2, 2)
    if isinstance(sub, str) and sd.get(2, 1) is None and J.get('subtitle') != sub:
        rep.err('grid_meta', sid, 'subtitle differs from B2', exp=sub, act=J.get('subtitle'))
    if J.get('freeze_rows', 0) != sd.xml.ysplit or J.get('freeze_cols', 0) != sd.xml.xsplit:
        rep.err('grid_meta', sid, 'freeze_rows/cols differ from pane ySplit/xSplit', exp=(sd.xml.ysplit, sd.xml.xsplit),
                act=(J.get('freeze_rows'), J.get('freeze_cols')))
    af = sd.xml.autofilter
    ja = J.get('autofilter')
    if af and (not ja or (ja.get('r1'), ja.get('c1'), ja.get('r2'), ja.get('c2')) != af):
        rep.err('grid_meta', sid, 'autofilter range differs', exp=af, act=ja)
    if not af and ja:
        rep.err('grid_meta', sid, 'JSON autofilter but Excel has none', act=ja)
    if bool(J.get('gridlines', True)) != sd.xml.gridlines:
        rep.err('grid_meta', sid, 'gridlines flag differs from sheetView showGridLines', exp=sd.xml.gridlines, act=J.get('gridlines'))
    if sid == '00':
        # table of contents: every A-column HYPERLINK to a sheet
        exp_toc = []
        for r in sorted(sd.rows):
            if isinstance(sd.get(r, 1), Formula):
                v = ctx.eng.cellval(sd.name, r, 1)
                if isinstance(v, LinkVal) and v.label in ctx.book.names:
                    exp_toc.append({'id': v.label[:2], 'name': v.label, 'purpose': sd.get(r, 2), 'row_unit': sd.get(r, 3),
                                    'rows': sd.get(r, 4), 'sources': sd.get(r, 5)})
        if exp_toc != J.get('toc'):
            rep.err('grid_meta', sid, '00 toc differs from the A-column sheet links (rows 108..)', exp=exp_toc[:3], act=(J.get('toc') or [])[:3])
        if [x['name'] for x in exp_toc] != ctx.book.names:
            rep.warn('selfcheck_toc', sid, '00 TOC does not list every sheet in order', act=[x['name'] for x in exp_toc])
        # legend: section '6 圖例' = group rows (A only) + swatch rows (A styled sample, B description)
        r0 = next((r for r in sorted(sd.rows) if isinstance(sd.get(r, 1), str) and '圖例' in sd.get(r, 1)
                   and (r, 1) in {(m[0], m[1]) for m in sd.xml.merges}), None)
        exp_lg = []
        if r0:
            group = None
            r = r0 + 1
            while r <= sd.max_row:
                a, b = sd.get(r, 1), sd.get(r, 2)
                if isinstance(a, str) and (r, 1) in {(m[0], m[1]) for m in sd.xml.merges} and b is None and r > r0 + 1 \
                        and _nb(ctx.book.styles.get(sd.sid(r, 1)).bg) == '#1F4E79':
                    break                                         # next banner section
                if a is not None:
                    si = ctx.book.styles.get(sd.sid(r, 1))
                    if b is None:
                        group = a
                    else:
                        exp_lg.append({'group': group, 'label': a, 'desc': b, 'bg': _nb(si.bg), 'fc': _nf_(si.fc),
                                       'b': si.b, 'i': si.i, 'u': si.u, 'bd': bool(si.bd)})
                r += 1
        act_lg = [{'group': x.get('group'), 'label': x.get('label'), 'desc': x.get('desc'), 'bg': _nb(x.get('bg')), 'fc': _nf_(x.get('fc')),
                   'b': bool(x.get('b')), 'i': bool(x.get('i')), 'u': bool(x.get('u')), 'bd': bool(x.get('bd'))} for x in J.get('legend') or []]
        if exp_lg != act_lg or not exp_lg:
            bad = [(e, a) for e, a in zip(exp_lg, act_lg) if e != a][:3]
            rep.err('grid_meta', sid, '00 legend differs from the 圖例 section (%d vs %d items)' % (len(exp_lg), len(act_lg)), exp=bad or exp_lg[:2], act=act_lg[:2])
    if sid == '14':
        exp_nt = {}
        wbx = ctx.book.z.read('xl/workbook.xml').decode('utf-8')
        for m in re.finditer(r'<definedName name="([^"]+)"(?![^>]*localSheetId)[^>]*>([^<]+)</definedName>', wbx):
            nm, ref = m.group(1), html.unescape(m.group(2))
            if ref.startswith("'%s'!" % sd.name):
                exp_nt[nm] = parse_area(ref.split('!')[1].replace('$', ''))
        act_nt = {x.get('name'): parse_area(x.get('ref')) for x in J.get('named_tables') or []}
        if exp_nt != act_nt:
            rep.err('grid_meta', sid, 'named_tables differ from workbook definedNames', exp=exp_nt, act=act_nt)
        for x in J.get('named_tables') or []:
            r1, c1, r2, c2 = parse_area(x.get('ref'))
            if x.get('header_row') != r1 or x.get('last') != r2 or x.get('cols') != c2 - c1 + 1:
                rep.err('grid_meta', sid, 'named table header/last/cols inconsistent with its ref', where=x.get('name'), act=x)


def verify_grid(ctx, sd, sid, J, rep, cfres):
    st = ctx.book.styles
    jcells, jrowh = {}, {}
    for row in J.get('rows') or []:
        r = row.get('r')
        jrowh[r] = row.get('h')
        for cell in row.get('cells') or []:
            k = (r, cell.get('c'))
            if k in jcells:
                rep.err('grid_dup_cell', sid, 'duplicate cell', where=coord(*k))
            jcells[k] = cell
    cov = sd.covered()
    merges = {(m[0], m[1]): m for m in sd.xml.merges}
    expected = set()
    ncmp = 0
    for r in sorted(sd.rows):
        row = sd.rows[r]
        for ci in range(len(row)):
            c = ci + 1
            raw = row[ci]
            sidn = sd.sid(r, c)
            si = st.get(sidn) if sidn else None
            styled = si is not None and (si.bg is not None)
            bordered = si is not None and bool(si.bd)
            if raw is None and not styled and not bordered and (r, c) not in merges:
                continue
            if (r, c) in cov:
                if raw is not None:
                    rep.warn('grid_covered_value', sid, 'value inside a merged range (hidden in Excel)', where=coord(r, c), exp=raw)
                continue
            expected.add((r, c))
            cell = jcells.get((r, c))
            if cell is None:
                if raw is not None:
                    ev, _ = expected_value(ctx, sd, r, c)
                    rep.err('grid_cell_missing', sid, 'non-empty Excel cell missing in grid JSON', where=coord(r, c), exp=ev)
                elif styled or (r, c) in merges:
                    rep.err('grid_styled_empty_missing', sid, 'empty filled/merged cell not emitted (spec: emit styled empties)', where=coord(r, c), exp=si.bg if si else None)
                else:
                    rep.warn('grid_bordered_empty_missing', sid, 'empty bordered cell not emitted', where=coord(r, c))
                continue
            ncmp += 1
            # value
            if raw is not None:
                ev, kind = expected_value(ctx, sd, r, c, raw)
                if kind == 'unchecked':
                    rep.uncheck(sid, '%s: %s' % (coord(r, c), ev))
                else:
                    v = cell.get('v')
                    act = {'t': v, 'l': cell.get('l')} if cell.get('l') is not None else v
                    if isinstance(v, dict):
                        act = v
                    for cat, msg, e, a in compare_cell(ctx, ev, kind, act):
                        rep.err('grid_' + cat, sid, msg, where=coord(r, c), exp=e, act=a)
                    av, al, au, af = decode_act(act)
                    if al is not None:
                        ctx.all_links.append((sid, coord(r, c), al))
                    # number format
                    num = ev.label if isinstance(ev, LinkVal) else ev
                    if is_num(num) and si is not None and norm_nf(si.nf):
                        sty = grid_style_of(J, cell)
                        f = cell.get('f') or af or sty.get('f') or sty.get('fmt') or sty.get('nf')
                        if f != si.nf:
                            rep.err('grid_numfmt', sid, 'number format of numeric cell not carried', where=coord(r, c), exp=si.nf, act=f)
                        if ';' in si.nf or cell.get('d') is not None:
                            dexp = render_sectioned(num, si.nf)
                            if cell.get('d') != dexp:
                                rep.err('grid_numfmt', sid, 'pre-rendered display text d differs', where=coord(r, c), exp=dexp, act=cell.get('d'))
            elif cell.get('v') not in (None, ''):
                rep.err('grid_extra_value', sid, 'JSON has a value where Excel cell is empty', where=coord(r, c), act=cell.get('v'))
            # merge
            m = merges.get((r, c))
            ecs, ers = (m[3] - m[1] + 1, m[2] - m[0] + 1) if m else (1, 1)
            if (cell.get('cs') or 1) != ecs or (cell.get('rs') or 1) != ers:
                rep.err('grid_merge', sid, 'merge span differs', where=coord(r, c), exp=(ecs, ers), act=(cell.get('cs'), cell.get('rs')))
            # style
            sty = grid_style_of(J, cell)
            exp = {'b': bool(si and si.b), 'i': bool(si and si.i), 'u': bool(si and si.u),
                   'bg': si.bg if si else None, 'fc': si.fc if si else None}
            cf = cfres.props.get((r, c)) if cfres else None
            if cf:
                for k in ('b', 'i', 'u', 'bg', 'fc'):
                    if k in cf:
                        exp[k] = cf[k]
            act = {'b': bool(sty.get('b')), 'i': bool(sty.get('i')), 'u': bool(sty.get('u')),
                   'bg': normc(sty.get('bg')), 'fc': normc(sty.get('fc'))}
            probs = [k for k in ('b', 'i', 'u') if exp[k] != act[k] and raw is not None]
            eb, ab = _nb(exp['bg']), _nb(act['bg'])
            if cfres and (r, c) in cfres.scale:
                scale_ok = color_close(eb or '#FFFFFF', ab or '#FFFFFF')
            if eb != ab and not (cfres and (r, c) in cfres.scale and scale_ok):
                probs.append('bg')
            if _nf_(exp['fc']) != _nf_(act['fc']) and raw is not None:
                probs.append('fc')
            if probs:
                rep.err('cf_style' if cf else 'grid_style', sid, 'cell style differs (%s)' % ','.join(probs),
                        where=coord(r, c), exp={k: exp[k] for k in probs}, act={k: act[k] for k in probs})
            lp = grid_layout_probs(si, sty, raw is None)
            if lp:
                rep.err('grid_layout_style', sid, 'cell layout style differs (%s)' % ','.join(p[0] for p in lp),
                        where=coord(r, c), exp={p[0]: p[1] for p in lp}, act={p[0]: p[2] for p in lp})
            if si is not None and si.shrink and raw is not None:
                rep.stats['grid_shrink_to_fit'].setdefault(sid, []).append(coord(r, c))
            # data bar
            eb_ = cfres.bars.get((r, c)) if cfres else None
            jb = cell.get('bar')
            if eb_:
                p, colr = eb_
                if not jb:
                    if p > 0:
                        rep.err('cf_databar', sid, 'data bar missing', where=coord(r, c), exp=(round(p, 4), colr))
                else:
                    gp = jb.get('p')
                    if gp is None or abs(gp - p) > BAR_TOL:
                        rep.err('cf_databar', sid, 'data bar length differs (p = (v-min)/(max-min), 4-decimal rounding allowed)', where=coord(r, c), exp=round(p, 6), act=gp)
                    if normc(jb.get('c')) not in (None, colr):
                        rep.err('cf_databar', sid, 'data bar colour differs', where=coord(r, c), exp=colr, act=jb.get('c'))
            elif jb:
                rep.err('cf_databar', sid, 'unexpected data bar', where=coord(r, c), act=jb)
    rep.stats['cells_compared'][sid] = ncmp
    check_grid_overflow(sd, sid, J, jcells, J.get('col_widths') or [], rep, st)
    check_grid_meta(ctx, sd, sid, J, rep)
    for k, cell in jcells.items():
        if k not in expected and k not in cov:
            if cell.get('v') not in (None, ''):
                rep.err('grid_extra_cell', sid, 'JSON cell not present in Excel', where=coord(*k), act=cell.get('v'))
        elif k in cov:
            rep.warn('grid_covered_cell', sid, 'JSON emits a merge-covered cell', where=coord(*k))
    # rows heights
    for r in range(1, sd.max_row + 1):
        d = sd.rowdims.get(r) or {}
        custom = d.get('customHeight') in ('1', 'true') and d.get('ht')
        h = jrowh.get(r)
        if custom:
            e = half_up(float(d['ht']) * 4 / 3)
            if h is None or abs(h - e) > 1:
                rep.err('layout', sid, 'row height px differs (round(pt*4/3))', where='row %d' % r, exp=e, act=h)
        elif h not in (None,):
            rep.err('layout', sid, 'row without customHeight has explicit h (contract: null = auto)', where='row %d' % r, act=h)
    cw = J.get('col_widths') or []
    for d in sd.xml.cols:
        if not d['custom']:
            continue
        for c in range(d['min'], min(d['max'], sd.max_col + 3) + 1):
            e = half_up(d['width'] * 7.5)
            a = cw[c - 1] if c - 1 < len(cw) else None
            if a is None or abs(a - e) > 1:
                rep.err('layout', sid, 'col width px differs (w*7.5, no clamp for grid)', where=num2col(c), exp=e, act=a)
    hc = sd.xml.hidden_cols()
    jh = set()
    for x in J.get('hidden_cols') or []:
        jh.add(col2num(x) if isinstance(x, str) else x)
    if hc != jh and hc != {x + 1 for x in jh}:
        rep.err('grid_hidden', sid, 'hidden columns differ', exp=sorted(hc), act=sorted(jh))
    hr = {r for r, d in sd.rowdims.items() if d.get('hidden') in ('1', 'true')}
    if hr != set(J.get('hidden_rows') or []):
        rep.err('grid_hidden', sid, 'hidden rows differ', exp=sorted(hr), act=J.get('hidden_rows'))
    if not sd.xml.gridlines and J.get('gridlines') is not False:
        rep.err('layout', sid, 'showGridLines=0 in Excel but JSON gridlines is not false', act=J.get('gridlines'))
    secs = {s.get('r') for s in J.get('sections') or []}
    for r in EXPECTED_SECTIONS.get(sid, []):
        if r not in secs:
            rep.err('grid_sections', sid, 'expected section row missing', where='row %d' % r, exp=sd.get(r, 1), act=sorted(x for x in secs if x is not None))
    for s in J.get('sections') or []:
        if s.get('r') is None or not (1 <= s['r'] <= sd.max_row):
            rep.err('grid_sections', sid, 'section row out of range', act=s)


# ----------------------------------------------------------------------------------------------
# CHARTS
# ----------------------------------------------------------------------------------------------
NS_C = '{http://schemas.openxmlformats.org/drawingml/2006/chart}'
NS_A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
NS_XDR = '{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}'


def _zpath(base_dir, target):
    import posixpath
    if target.startswith('/'):
        return target.lstrip('/')
    return posixpath.normpath(posixpath.join(base_dir, target))


def _rels(z, path):
    import posixpath
    d, f = posixpath.split(path)
    rp = posixpath.join(d, '_rels', f + '.rels')
    if rp not in z.namelist():
        return {}
    root = ET.fromstring(z.read(rp))
    return {r.get('Id'): (r.get('Type'), _zpath(d, r.get('Target'))) for r in root}


def _srgb(el):
    if el is None:
        return None
    for path in ('%ssolidFill/%ssrgbClr' % (NS_A, NS_A), '%sln/%ssolidFill/%ssrgbClr' % (NS_A, NS_A, NS_A)):
        x = el.find(path)
        if x is not None:
            return normc(x.get('val'))
    return None


def parse_charts(ctx):
    book = ctx.book
    z = book.z
    out = collections.defaultdict(list)

    def ref_vals(f):
        if not f:
            return None
        i = f.rindex('!')
        sh = f[:i].strip("'").replace("''", "'")
        r1, c1, r2, c2 = parse_area(f[i + 1:].replace('$', ''))
        vals = []
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                v = ctx.eng.cellval(sh, r, c)
                if isinstance(v, LinkVal):
                    v = v.label
                vals.append(v)
        return vals
    for name in book.names:
        for rid, (typ, dpath) in _rels(z, book.parts[name]).items():
            if not typ.endswith('/drawing'):
                continue
            drels = _rels(z, dpath)
            droot = ET.fromstring(z.read(dpath))
            for anc in droot:
                tag = anc.tag.split('}')[1]
                if tag not in ('oneCellAnchor', 'twoCellAnchor', 'absoluteAnchor'):
                    continue
                fr = anc.find(NS_XDR + 'from')
                ch = anc.find('.//%schart' % NS_C)
                if ch is None:
                    continue
                cpath = drels[ch.get('{%s}id' % NS_REL)][1]
                ext = anc.find(NS_XDR + 'ext')
                croot = ET.fromstring(z.read(cpath))
                chart = croot.find(NS_C + 'chart')
                title_el = chart.find(NS_C + 'title')
                title = ''.join(t.text or '' for t in title_el.iter(NS_A + 't')) if title_el is not None else None
                pa = chart.find(NS_C + 'plotArea')
                axes = {}
                for ax in list(pa):
                    tg = ax.tag.split('}')[1]
                    if tg in ('valAx', 'catAx', 'dateAx'):
                        aid = ax.find(NS_C + 'axId').get('val')
                        t = ax.find(NS_C + 'title')
                        axes[aid] = (tg, ''.join(x.text or '' for x in t.iter(NS_A + 't')) if t is not None else None)
                groups = [g for g in pa if g.tag.endswith('Chart')]
                kinds = [g.tag.split('}')[1] for g in groups]
                first_val_ax = None
                series = []
                bar_dir, grouping = None, None
                for g in groups:
                    k = g.tag.split('}')[1]
                    axids = [a.get('val') for a in g.findall(NS_C + 'axId')]
                    vax = axids[1] if len(axids) > 1 else None
                    if first_val_ax is None:
                        first_val_ax = vax
                    if k == 'barChart':
                        bd = g.find(NS_C + 'barDir')
                        bar_dir = bd.get('val') if bd is not None else 'col'
                        gp = g.find(NS_C + 'grouping')
                        grouping = gp.get('val') if gp is not None else 'clustered'
                    stype = {'barChart': 'bar', 'lineChart': 'line', 'doughnutChart': 'doughnut', 'pieChart': 'pie',
                             'areaChart': 'area'}.get(k, k)
                    for s in g.findall(NS_C + 'ser'):
                        fx = lambda p: (s.find(p) is not None and s.find(p).text) or None
                        tx = fx('%stx/%sstrRef/%sf' % (NS_C, NS_C, NS_C))
                        cat = fx('%scat/%snumRef/%sf' % (NS_C, NS_C, NS_C)) or fx('%scat/%sstrRef/%sf' % (NS_C, NS_C, NS_C))
                        val = fx('%sval/%snumRef/%sf' % (NS_C, NS_C, NS_C))
                        sp = s.find(NS_C + 'spPr')
                        dpts = []
                        for dp in s.findall(NS_C + 'dPt'):
                            dpts.append(_srgb(dp.find(NS_C + 'spPr')))
                        nm = ref_vals(tx)
                        series.append({'name': nm[0] if nm else None, 'values': ref_vals(val), 'cats': ref_vals(cat),
                                       'color': _srgb(sp), 'colors': dpts or None, 'type': stype,
                                       'axis': 'y2' if (vax is not None and first_val_ax is not None and vax != first_val_ax) else 'y',
                                       'refs': (tx, cat, val)})
                if 'barChart' in kinds and 'lineChart' in kinds:
                    ctype = 'combo'
                else:
                    ctype = {'barChart': 'bar', 'lineChart': 'line', 'doughnutChart': 'doughnut', 'pieChart': 'pie',
                             'areaChart': 'area'}.get(kinds[0], kinds[0]) if kinds else None
                cats = next((s['cats'] for s in series if s['cats']), None)
                lg = chart.find(NS_C + 'legend')
                lp = lg.find(NS_C + 'legendPos').get('val') if lg is not None and lg.find(NS_C + 'legendPos') is not None else None
                vtitles = [axes.get(first_val_ax, (None, None))[1]]
                other = [a for a, (tg, _) in axes.items() if tg == 'valAx' and a != first_val_ax]
                out[name[:2]].append({
                    'title': title, 'type': ctype, 'horizontal': bar_dir == 'bar', 'stacked': grouping in ('stacked', 'percentStacked'),
                    'from': {'r': int(fr.find(NS_XDR + 'row').text) + 1, 'c': int(fr.find(NS_XDR + 'col').text) + 1} if fr is not None else None,
                    'size_px': {'w': half_up(int(ext.get('cx')) / 9525), 'h': half_up(int(ext.get('cy')) / 9525)} if ext is not None else None,
                    'categories': cats, 'series': series, 'legend': {'b': 'bottom', 'r': 'right', 't': 'top', 'l': 'left'}.get(lp, lp) if lg is not None else None,
                    'y_title': vtitles[0], 'y2_title': axes[other[0]][1] if other else None, 'part': cpath})
    return out


def _list_eq(e, a):
    if not isinstance(a, list) or len(a) != len(e):
        return False
    return all(val_eq(x, y) for x, y in zip(e, a))


def verify_charts(ctx, sid, J, xcharts, contract, rep):
    jc = J.get('charts') or []
    exp_n = EXPECTED_CHART_COUNT.get(sid, 0)
    if len(xcharts) != exp_n:
        rep.warn('selfcheck_charts', sid, 'chart count parsed from xlsx differs from spec', exp=exp_n, act=len(xcharts))
    if len(jc) != len(xcharts):
        rep.err('chart_count', sid, 'number of charts differs', exp=len(xcharts), act=len(jc))
    for i, xc in enumerate(xcharts):
        if i >= len(jc):
            break
        c = jc[i]
        w = 'chart %d (%s)' % (i + 1, c.get('id'))
        for k in ('title', 'type'):
            if c.get(k) != xc[k]:
                rep.err('chart', sid, '%s differs' % k, where=w, exp=xc[k], act=c.get(k))
        for k in ('horizontal', 'stacked'):
            if bool(c.get(k)) != xc[k]:
                rep.err('chart', sid, '%s differs' % k, where=w, exp=xc[k], act=c.get(k))
        af = (c.get('anchor') or {}).get('from')
        if af != xc['from']:
            rep.err('chart', sid, 'anchor.from differs', where=w, exp=xc['from'], act=af)
        if c.get('size_px') and c.get('size_px') != xc['size_px']:
            rep.err('chart', sid, 'size_px differs', where=w, exp=xc['size_px'], act=c.get('size_px'))
        if not c.get('size_px'):
            rep.err('chart', sid, 'size_px missing', where=w, exp=xc['size_px'])
        if not _list_eq(xc['categories'] or [], c.get('categories')):
            rep.err('chart_data', sid, 'categories differ from referenced cells', where=w, exp=xc['categories'], act=c.get('categories'))
        js = c.get('series') or []
        if len(js) != len(xc['series']):
            rep.err('chart_data', sid, 'series count differs', where=w, exp=len(xc['series']), act=len(js))
        for k, xs in enumerate(xc['series']):
            if k >= len(js):
                break
            s = js[k]
            ws = '%s series %d' % (w, k)
            if s.get('name') != xs['name']:
                rep.err('chart_data', sid, 'series name differs', where=ws, exp=xs['name'], act=s.get('name'))
            if not _list_eq(xs['values'] or [], s.get('values')):
                rep.err('chart_data', sid, 'series values differ from referenced cells %s' % xs['refs'][2], where=ws, exp=xs['values'], act=s.get('values'))
            if xs['colors']:
                if [normc(x) for x in (s.get('colors') or [])] != xs['colors']:
                    rep.err('chart', sid, 'per-point colours differ', where=ws, exp=xs['colors'], act=s.get('colors'))
            elif normc(s.get('color')) != xs['color']:
                rep.err('chart', sid, 'series colour differs', where=ws, exp=xs['color'], act=s.get('color'))
            if ctx_type(s.get('type'), c.get('type')) != xs['type']:
                rep.err('chart', sid, 'series type differs', where=ws, exp=xs['type'], act=s.get('type'))
            if (s.get('axis') or 'y') != xs['axis']:
                rep.err('chart', sid, 'series axis differs', where=ws, exp=xs['axis'], act=s.get('axis'))
        for k in ('y_title', 'y2_title'):
            if (c.get(k) or None) != (xc[k] or None):
                rep.err('chart', sid, '%s differs' % k, where=w, exp=xc[k], act=c.get(k))
        if 'legend' in c and c.get('legend') != xc['legend']:
            rep.err('chart', sid, 'legend position differs', where=w, exp=xc['legend'], act=c.get('legend'))
        # contract expectation (charts_contract.json by the chart analyst)
        cc = contract.get(c.get('id')) if contract else None
        if contract is not None and cc is None:
            rep.err('chart_contract', sid, 'chart id not in charts_contract.json', where=w, act=c.get('id'))
        if cc:
            for k in ('title', 'type', 'horizontal', 'stacked', 'categories', 'y_title', 'y2_title'):
                if k in cc and not (_list_eq(cc[k], c.get(k)) if isinstance(cc[k], list) else cc[k] == c.get(k)):
                    rep.err('chart_contract', sid, '%s differs from charts_contract.json' % k, where=w, exp=cc[k], act=c.get(k))
            if cc.get('anchor') != c.get('anchor'):
                rep.err('chart_contract', sid, 'anchor differs from charts_contract.json (to-cell is derived)', where=w, exp=cc.get('anchor'), act=c.get('anchor'))
            for k, cs_ in enumerate(cc.get('series') or []):
                s = js[k] if k < len(js) else {}
                for kk in ('name', 'color', 'type', 'axis', 'colors'):
                    if kk in cs_ and cs_[kk] is not None and cs_[kk] != s.get(kk):
                        rep.err('chart_contract', sid, 'series %s differs from charts_contract.json' % kk, where='%s series %d' % (w, k), exp=cs_[kk], act=s.get(kk))
                if not _list_eq(cs_.get('values') or [], s.get('values')):
                    rep.err('chart_contract', sid, 'series values differ from charts_contract.json', where='%s series %d' % (w, k), exp=cs_.get('values'), act=s.get('values'))
                for kk in ('line_px', 'point_radius', 'border'):
                    if kk in cs_ and cs_[kk] != s.get(kk):
                        rep.err('chart_contract_ext', sid, 'series %s differs from charts_contract.json' % kk, where='%s series %d' % (w, k), exp=cs_[kk], act=s.get(kk))
            for k in ('size_px', 'legend', 'cat_font_pt', 'y', 'y2', 'data_labels', 'hole', 'src'):
                act = c.get(k)
                if k in ('y', 'y2') and isinstance(act, dict) and 'title' in act:
                    # axis title inside y/y2 is an extra the contract file did not list: verify it against the
                    # chart XML <c:valAx><c:title> instead, and compare the remaining keys with the contract
                    xt = xc['y_title' if k == 'y' else 'y2_title']
                    if act.get('title') != xt:
                        rep.err('chart', sid, '%s.title differs from the axis title in the chart XML' % k, where=w, exp=xt, act=act.get('title'))
                    act = {kk: vv for kk, vv in act.items() if kk != 'title'}
                if k in cc and cc[k] != act:
                    rep.err('chart_contract_ext', sid, 'extension key %s differs from charts_contract.json' % k, where=w, exp=cc[k], act=c.get(k))


def ctx_type(t, ctype):
    return t if t else ('bar' if ctype in ('bar', 'combo') else ctype)


# ----------------------------------------------------------------------------------------------
# CARD (02_設備查詢卡)
# ----------------------------------------------------------------------------------------------
CARD_SHEET = '02_設備查詢卡'
_IDX03 = re.compile(r"INDEX\('03_設備總表'!\$([A-Z]+):\$[A-Z]+,\$J\$4\)")
_IDX08 = re.compile(r"INDEX\('08_變更歷程'!\$([A-Z]+):\$[A-Z]+,MATCH")


def card_expected_fields(ctx):
    sd = ctx.book.sheet('02')
    heads = {}
    for (r1, c1, r2, c2) in sd.xml.merges:
        if r1 == r2 and c1 == 2 and c2 == 8 and sd.get(r1, 2) and r1 >= 6:
            heads[r1] = sd.get(r1, 2)
    fields = []
    sec = None
    for r in range(5, 47):
        if r in heads:
            sec = heads[r]
            continue
        for c in (3, 6):
            raw = sd.get(r, c)
            if isinstance(raw, Formula) and '03_設備總表' in raw.text:
                cols = []
                for x in _IDX03.findall(raw.text):
                    i = col2num(x) - 1
                    if i not in cols:
                        cols.append(i)
                fields.append({'section': sec, 'label': sd.get(r, c - 1), 'cols': cols, 'cell': coord(r, c),
                               'raw': 'IF(INDEX(' in raw.text, 'nf': ctx.book.nf(sd, r, c), 'rc': (r, c)})
    return fields, heads


def card_expected_links(ctx):
    sd = ctx.book.sheet('02')
    out = []
    for r in (48, 49, 50, 51):
        for c in (2, 5):
            raw = sd.get(r, c)
            if not isinstance(raw, Formula):
                continue
            t = raw.text
            d = {'cell': coord(r, c), 'rc': (r, c)}
            m = re.search(r"MATCH\(\$C\$4,'([^']+)'!\$([A-Z]+):", t)
            if m:
                d['sheet'] = m.group(1)[:2]
                d['match_col'] = col2num(m.group(2)) - 1
                m2 = re.search(r"COUNTIF\('([^']+)'!\$([A-Z]+):\$[A-Z]+,\$C\$4(&\"#\*\")?\)", t)
                d['count_col'] = col2num(m2.group(2)) - 1
                d['count_mode'] = 'prefix#' if m2.group(3) else 'eq'
            elif 'DF:$DF' in t:
                d['param'] = {'sheet_col': col2num('DF') - 1, 'row_col': col2num('DG') - 1, 'count_col': col2num('DH') - 1}
            else:
                d['self'] = True
                d['sheet'] = '03'
            out.append(d)
    return out


def excel_link_index(ctx, spec_links):
    """alias -> {sid: [r0, n]} computed from the Excel sheets with COUNTIF / MATCH semantics"""
    res = collections.defaultdict(dict)
    for d in spec_links:
        if 'match_col' not in d:
            continue
        sd = ctx.book.sheet(d['sheet'])
        fr = ctx.first_row[d['sheet']]
        first = {}
        cnt = collections.Counter()
        mc, cc = d['match_col'] + 1, d['count_col'] + 1
        for r in range(1, sd.max_row + 1):
            v = sd.get(r, mc)
            if isinstance(v, str):
                first.setdefault(v.upper(), r)
            w = sd.get(r, cc)
            if isinstance(w, str):
                if d['count_mode'] == 'prefix#':
                    if '#' in w:
                        cnt[w.split('#', 1)[0].upper()] += 1
                else:
                    cnt[w.upper()] += 1
        for alias in set(first) | set(cnt):
            r = first.get(alias)
            res[alias][d['sheet']] = [None if r is None else r - fr, cnt.get(alias, 0)]
    return res


def card_engine_eval(ctx, q):
    eng = ctx.eng
    eng.over[(CARD_SHEET, 3, 3)] = None if q == '' else q
    for k in [k for k in eng.memo if k[0] == CARD_SHEET]:
        del eng.memo[k]
    eng.range_cache = {k: v for k, v in eng.range_cache.items() if k[0] != CARD_SHEET}
    out = {}
    sd = ctx.book.sheet('02')
    for r in range(1, 65):
        for c in range(1, 11):
            if isinstance(sd.get(r, c), Formula):
                try:
                    out[(r, c)] = eng.cellval(CARD_SHEET, r, c)
                except (NotLoaded, Unsupported) as e:
                    out[(r, c)] = 'UNCHECKED:%s' % e
    return out


def _jrows(ctx, sid):
    J = ctx.json.get(sid)
    return (J or {}).get('rows') or []


def _jv(row, i):
    if row is None or i >= len(row):
        return None
    return decode_act(row[i])[0]


def card_sim_json(ctx, J02, q):
    """simulate the front-end recipe (spec 02 §3-6) on the JSON tables"""
    out = {'status': None, 'alias': '', 'count': '', 'i3': None, 'links': {}, 'recent': []}
    if q == '':
        out['status'] = '請輸入位號…'
        return out
    key = re.sub(' +', ' ', q.replace('=', '')).strip().upper()
    r05 = _jrows(ctx, '05')
    i5 = next((i for i, row in enumerate(r05) if str(_jv(row, 0)) == key), None)
    if i5 is None:
        out['status'] = '查無此字串（萬用字元 * ? 不適用）：請到『05_位號索引』以 Ctrl+F 搜尋部分字串'
        return out
    row5 = r05[i5]
    alias = _jv(row5, 4) or ''
    out['alias'] = str(alias)
    out['count'] = _jv(row5, 7)
    r03 = _jrows(ctx, '03')
    i3 = next((i for i, row in enumerate(r03) if alias and _jv(row, 1) == alias), None)
    out['i3'] = i3
    st = '來源：%s［鍵 %s］' % (_jv(row5, 2), _jv(row5, 0))
    st += '（測試定義清單位號，無對應設備）' if i3 is None else ' → 目前位號 %s' % _jv(r03[i3], 0)
    out['status'] = st
    if i3 is None:
        return out
    for L in J02.get('links') or []:
        sheet = str(L.get('sheet') or '')
        if L.get('param'):
            p = L['param']
            row = r03[i3]
            df, dg, dh = _jv(row, p.get('sheet_col', 109)), _jv(row, p.get('row_col', 110)), _jv(row, p.get('count_col', 111))
            out['links']['param'] = (None if not df else (str(df)[:2], dg - 5 if is_num(dg) else None, dh))
            continue
        if L.get('self'):
            out['links']['03self'] = ('03', i3, None)
            continue
        rows = _jrows(ctx, sheet)
        mc, cc, mode = L.get('match_col'), L.get('count_col'), L.get('count_mode')
        if mc is None:
            continue
        r0 = next((i for i, row in enumerate(rows) if str(_jv(row, mc) or '').upper() == alias.upper()), None)
        if mode == 'prefix#':
            n = sum(1 for row in rows if str(_jv(row, cc) or '').upper().startswith(alias.upper() + '#'))
        else:
            n = sum(1 for row in rows if str(_jv(row, cc) or '').upper() == alias.upper())
        out['links'][sheet] = (sheet, r0, n)
    rc = J02.get('recent_changes') or {}
    r08 = _jrows(ctx, str(rc.get('sheet') or '08'))
    kc = rc.get('key_col', 0)
    idx = {}
    for i, row in enumerate(r08):
        v = _jv(row, kc)
        if isinstance(v, str) and v.upper().startswith(alias.upper() + '#'):
            idx[v.upper()] = i
    for k in range(1, (rc.get('k_max') or 10) + 1):
        i = idx.get('%s#%d' % (alias.upper(), k))
        vals = []
        for col in rc.get('columns') or []:
            if i is None:
                vals.append('')
                continue
            if 'cols' in col:
                vals.append((col.get('join') or ' → ').join('' if _jv(r08[i], x) is None else str(_jv(r08[i], x)) for x in col['cols']))
            else:
                v = _jv(r08[i], col.get('col'))
                vals.append('' if v is None else v)
        out['recent'].append(vals)
    return out


def verify_card(ctx, J, rep):
    sid = '02'
    sd = ctx.book.sheet('02')
    if J.get('mode') != 'card':
        rep.err('card', sid, 'mode is not card', act=J.get('mode'))
    if J.get('default_query') != sd.get(3, 3):
        rep.err('card', sid, 'default_query differs from C3', exp=sd.get(3, 3), act=J.get('default_query'))
    if J.get('title') != sd.get(1, 2):
        rep.err('card', sid, 'title differs from B1', exp=sd.get(1, 2), act=J.get('title'))
    hl = J.get('home_link')
    if isinstance(hl, dict) and isinstance(hl.get('l'), dict):
        ctx.all_links.append((sid, 'home_link', hl['l']))
        if hl['l'] != ctx.resolve_loc("#'00_說明'!A1"):
            rep.err('card', sid, 'home_link target differs from A1 HYPERLINK', exp=ctx.resolve_loc("#'00_說明'!A1"), act=hl['l'])
    notes = J.get('notes') or []
    if sd.get(2, 2) not in notes and '｜'.join(map(str, notes)) != sd.get(2, 2):
        rep.err('card', sid, 'notes do not contain B2', exp=sd.get(2, 2), act=notes)
    # --- fields
    ef, heads = card_expected_fields(ctx)
    jf = []
    for sec in J.get('sections') or []:
        for f in sec.get('fields') or []:
            jf.append((sec.get('label'), f))
    if len(jf) != len(ef):
        rep.err('card_fields', sid, 'field count differs', exp=len(ef), act=len(jf))
    for i, e in enumerate(ef):
        if i >= len(jf):
            break
        sl, f = jf[i]
        col = f.get('col')
        cols = col if isinstance(col, list) else [col]
        w = '%s %s' % (e['cell'], e['label'])
        if sl != e['section']:
            rep.err('card_fields', sid, 'section label differs', where=w, exp=e['section'], act=sl)
        if f.get('label') != e['label']:
            rep.err('card_fields', sid, 'field label differs', where=w, exp=e['label'], act=f.get('label'))
        if cols != e['cols']:
            rep.err('card_fields', sid, 'field 03 column index differs from INDEX formula', where=w, exp=e['cols'], act=col)
        if 'raw' in f and bool(f.get('raw')) != e['raw']:
            rep.err('card_fields', sid, 'raw/text shape differs', where=w, exp=e['raw'], act=f.get('raw'))
        if e['raw'] and norm_nf(f.get('fmt')) != norm_nf(e['nf']):
            rep.warn('card_fields', sid, 'field number format differs', where=w, exp=e['nf'], act=f.get('fmt'))
        if f.get('cell') is not None and str(f.get('cell')).split(':')[0] != e['cell']:
            rep.err('card_fields', sid, 'field cell differs', where=w, exp=e['cell'], act=f.get('cell'))
    # --- links
    el = card_expected_links(ctx)
    jl = J.get('links') or []
    if len(jl) != len(el):
        rep.err('card_links', sid, 'quick-link count differs', exp=len(el), act=len(jl))
    for e in el:
        cand = None
        for L in jl:
            if e.get('param') and L.get('param'):
                cand = L
            elif e.get('self') and (L.get('self') or (str(L.get('sheet')) == '03' and L.get('match_col') is None)):
                cand = L
            elif 'match_col' in e and str(L.get('sheet')) == e['sheet'] and not L.get('param'):
                cand = L
            if cand:
                break
        if cand is None:
            rep.err('card_links', sid, 'quick link missing', where=e['cell'], exp=e)
            continue
        for k in ('match_col', 'count_col', 'count_mode'):
            if k in e and cand.get(k) != e[k]:
                rep.err('card_links', sid, 'link %s differs' % k, where=e['cell'], exp=e[k], act=cand.get(k))
        if e.get('param'):
            for k, v in e['param'].items():
                if (cand.get('param') or {}).get(k) != v:
                    rep.err('card_links', sid, 'param link %s differs' % k, where=e['cell'], exp=v, act=cand.get('param'))
    # --- recent changes
    rc = J.get('recent_changes') or {}
    ecols = []
    for c in (2, 3, 4, 6, 7, 8):
        raw = sd.get(55, c)
        if isinstance(raw, Formula):
            xs = [col2num(x) - 1 for x in _IDX08.findall(raw.text)]
            ecols.append(xs if len(xs) > 1 else xs[0])
    jcols = [(x.get('cols') if 'cols' in x else x.get('col')) for x in rc.get('columns') or []]
    if jcols != ecols:
        rep.err('card_recent', sid, 'recent_changes columns differ from 08 INDEX formulas', exp=ecols, act=jcols)
    if str(rc.get('sheet')) != '08' or rc.get('key_col') != 0 or rc.get('k_max') != 10:
        rep.err('card_recent', sid, 'recent_changes sheet/key_col/k_max differ', exp=('08', 0, 10), act=(rc.get('sheet'), rc.get('key_col'), rc.get('k_max')))
    # --- rules (CF of sheet3.xml)
    xr = sd.xml.cf
    jr = J.get('rules') or []
    if len(jr) != len(xr):
        rep.err('card_rules', sid, 'CF rule count differs', exp=len(xr), act=len(jr))
    exp_targets = [{28, 29, 30, 31, 32, 33}, {48}, {98}, {'count'}]
    for i, et in enumerate(exp_targets):
        if i < len(jr) and set(jr[i].get('targets') or []) != et:
            rep.err('card_rules', sid, 'rule %d targets differ' % (i + 1), exp=sorted(map(str, et)), act=jr[i].get('targets'))
        if i < len(jr):
            dxf = ctx.book.dxfs[xr[i]['dxf']]
            s = jr[i].get('style') or {}
            for k in ('bg', 'fc'):
                if normc(s.get(k)) != dxf.get(k):
                    rep.err('card_rules', sid, 'rule %d %s differs from dxf' % (i + 1, k), exp=dxf.get(k), act=s.get(k))
            if bool(s.get('b')) != bool(dxf.get('b')):
                rep.err('card_rules', sid, 'rule %d bold differs from dxf' % (i + 1), exp=dxf.get('b'), act=s.get('b'))
    # --- link_index vs COUNTIF / MATCH semantics
    li = J.get('link_index')
    if li is None:
        rep.warn('card_link_index', sid, 'link_index absent (contract v2 asks for it)')
    else:
        ex = excel_link_index(ctx, el)
        aliases = [ctx.book.sheet('03').get(r, 2) for r in range(ctx.first_row['03'], ctx.last_row['03'] + 1)]
        nbad = 0
        for a in aliases:
            got = li.get(a) or {}
            exp = ex.get(a.upper(), {})
            for d in el:
                if 'match_col' not in d:
                    continue
                s = d['sheet']
                e = exp.get(s) or [None, 0]
                g = got.get(s)
                if g is None or g == [] or g == [None, 0] or g == [-1, 0]:
                    g = [None, 0]
                g = list(g)[:2]
                if g != e:
                    nbad += 1
                    rep.err('card_link_index', sid, 'link_index [r0,n] differs from MATCH/COUNTIF', where='%s → %s' % (a, s), exp=e, act=got.get(s))
        extra = set(li) - set(aliases)
        if extra:
            rep.warn('card_link_index', sid, 'link_index has aliases not in 03', act=sorted(extra)[:10])
        rep.stats['card'] ['link_index_checked'] = len(aliases)
    # --- engine vectors (self-check of the verifier) + JSON simulation vs engine
    for (q, alias, cnt, i3, status) in CARD_VECTORS:
        ev = card_engine_eval(ctx, q)
        e_alias = ev.get((4, 3))
        e_cnt = ev.get((4, 7))
        e_st = ev.get((3, 6))
        if (e_alias or '') != alias or (e_cnt if e_cnt is not None else '') != cnt or (status and e_st != status):
            rep.warn('selfcheck_card', sid, 'engine result differs from spec vector', where=q, exp=(alias, cnt, status), act=(e_alias, e_cnt, e_st))
        if '03' not in ctx.json or '05' not in ctx.json:
            continue
        sim = card_sim_json(ctx, J, q)
        if sim['status'] != e_st:
            rep.err('card_sim', sid, 'status from JSON lookup differs from Excel F3', where=q, exp=e_st, act=sim['status'])
        if sim['alias'] != (e_alias or ''):
            rep.err('card_sim', sid, 'alias from JSON lookup differs from Excel C4', where=q, exp=e_alias, act=sim['alias'])
        if not val_eq(e_cnt if e_cnt != '' else None, sim['count'] if sim['count'] != '' else None):
            rep.err('card_sim', sid, 'count from JSON lookup differs from Excel G4', where=q, exp=e_cnt, act=sim['count'])
        if i3 is not None and sim['i3'] != i3:
            rep.err('card_sim', sid, '03 row index differs', where=q, exp=i3, act=sim['i3'])
        if sim['i3'] is None:
            continue
        r03 = _jrows(ctx, '03')[sim['i3']]
        sd03 = ctx.book.sheet('03')
        for k, (sl, f) in enumerate(jf):
            cell = str(f.get('cell') or '').split(':')[0] or (ef[k]['cell'] if k < len(ef) else None)
            if not cell:
                continue
            r, c = parse_coord(cell)
            e = ev.get((r, c))
            col = f.get('col')
            if isinstance(col, list):
                a = re.sub(' +', ' ', ' '.join('' if _jv(r03, x) is None else str(_jv(r03, x)) for x in col)).strip()
            else:
                a = _jv(r03, col)
                if isinstance(e, (dt.datetime, dt.date)):
                    e = fmt_datetime(e, ctx.book.nf(sd03, ctx.first_row['03'] + sim['i3'], col + 1))
            if isinstance(e, str) and is_num(a):
                a = general_str(a)
            if not val_eq(e, a):
                rep.err('card_sim', sid, 'field value via JSON mapping differs from Excel card cell', where='%s %s' % (q, cell), exp=e, act=a)
        for d in el:
            e = ev.get(d['rc'])
            if isinstance(e, str) and e.startswith('UNCHECKED'):
                rep.uncheck(sid, 'card link %s: %s' % (d['cell'], e))
                continue
            key = 'param' if d.get('param') else ('03self' if d.get('self') else d.get('sheet'))
            s = sim['links'].get(key)
            if isinstance(e, LinkVal):
                tgt = ctx.resolve_loc(e.loc)
                m = re.search(r'(\d+) 筆', str(e.label))
                if s is None:
                    rep.err('card_sim', sid, 'JSON cannot reproduce link', where='%s %s' % (q, d['cell']), exp=(e.label, tgt), act=None)
                    continue
                if (tgt.get('s'), tgt.get('r')) != (s[0], s[1]):
                    rep.err('card_sim', sid, 'link target via JSON differs from Excel', where='%s %s' % (q, d['cell']), exp=tgt, act=s[:2])
                if m and s[2] is not None and int(m.group(1)) != s[2]:
                    rep.err('card_sim', sid, 'link count via JSON differs from Excel', where='%s %s' % (q, d['cell']), exp=e.label, act=s[2])
            elif isinstance(e, str) and e.endswith('（無）'):
                if s is not None and s[1] is not None and (s[2] or 0) > 0:
                    rep.err('card_sim', sid, 'Excel shows （無） but JSON finds rows', where='%s %s' % (q, d['cell']), exp=e, act=s)
        for k in range(10):
            erow = []
            for c in (2, 3, 4, 6, 7, 8):
                v = ev.get((55 + k, c))
                if isinstance(v, (dt.datetime, dt.date)):
                    v = fmt_datetime(v, ctx.book.nf(sd, 55 + k, c))
                erow.append('' if v is None else v)
            arow = sim['recent'][k] if k < len(sim['recent']) else None
            if arow is None:
                continue
            arow = list(arow)
            if arow and isinstance(arow[0], str) and len(arow[0]) > 16:
                arow[0] = arow[0][:16]
            if [str(x) for x in erow] != [str(x) for x in arow]:
                rep.err('card_sim', sid, 'recent change row via JSON differs from Excel', where='%s k=%d' % (q, k + 1), exp=erow, act=arow)
    ctx.eng.over.pop((CARD_SHEET, 3, 3), None)
    for k in [k for k in ctx.eng.memo if k[0] == CARD_SHEET]:
        del ctx.eng.memo[k]


# ----------------------------------------------------------------------------------------------
# CHUNKED sheets 21 / 22 (run in worker processes, streamed; JSON parts loaded one at a time)
# ----------------------------------------------------------------------------------------------
def big_worker(args):
    xlsx, data_dir, sid, quick, nsamples, entry = args
    rep = Report(nsamples)
    t0 = time.time()
    book = Book(xlsx)
    wb = book.open_wb()
    name = book.by_id[sid]
    ctx = Ctx(book, Engine(book), rep, {}, data_dir, quick)
    first, last, ncols = ctx.first_row[sid], ctx.last_row[sid], ctx.ncols[sid]
    sd = SheetData(name, sid, book.names.index(name), book.xml[name])
    book.sheets[name] = sd
    st = book.styles
    out = {'sid': sid, 'first': first, 'last': last, 'ncols': ncols}
    # JSON parts
    files = (entry or {}).get('files') or []
    parts = []
    for f in files:
        p = os.path.join(data_dir, f)
        if not os.path.exists(p):
            rep.err('missing_file', sid, 'part file missing', where=f)
            parts.append(None)
        else:
            parts.append(p)
    have_json = bool(parts) and all(parts)
    cur = {'k': -1, 'J': None, 'off': 0, 'n': 0}
    offsets = []
    part_meta = []

    def load_part(k):
        J = load_json(parts[k])
        cur.update(k=k, J=J, off=(J.get('part') or {}).get('row_offset', 0), n=len(J.get('rows') or []))
        sts = J.get('styles') or []
        cur['rs'] = {e[0]: parse_style(sts[e[1]] if isinstance(e[1], int) and e[1] < len(sts) else e[1]) for e in (J.get('row_styles') or [])}
        cur['cs_'], _rs_unused, cur['cp'] = json_table_style_layers(J, None)
        part_meta.append((k, J.get('part'), len(J.get('rows') or []), os.path.basename(parts[k])))
        return J

    colB, colD, parity = [], [], []
    counta = collections.Counter()
    col_nf = collections.defaultdict(collections.Counter)
    colmap = None
    agg = collections.defaultdict(ColAgg)
    uicols, inv_colmap, jcols0 = {}, {}, []
    jalias, jblock = [], []
    ncmp = 0
    zebra_rule = None
    for rule in sd.xml.cf:
        m = re.match(r'^\$([A-Z]+)(\d+)=(\d+)$', (rule['formulas'] or [''])[0])
        if rule['type'] == 'expression' and m:
            zebra_rule = (col2num(m.group(1)), int(m.group(3)), book.dxfs[rule['dxf']].get('bg'))
        else:
            rep.uncheck(sid, 'CF rule %s not evaluated in streaming mode' % rule['formulas'])
    exp_zebra_rows = 0
    jz_state = {'prev': None, 'par': -1}
    zebra_bad = 0
    jr_all = 0
    for ridx, cells in iter_sheet_cells(wb, wb[name]):
        if ridx <= 4:
            if cells:
                maxc = cells[-1]['column']
                vals, sids = [None] * maxc, [0] * maxc
                for cell in cells:
                    v = cell['value']
                    vals[cell['column'] - 1] = Formula(v) if cell['data_type'] == 'f' else v
                    sids[cell['column'] - 1] = cell['style_id'] or 0
                sd.rows[ridx], sd.srows[ridx] = vals, sids
            if ridx == 4 and have_json:
                J0 = load_part(0)
                xl = {c: sd.get(4, c) for c in range(1, ncols + 1)}
                colmap = table_colmap(sid, ncols, J0.get('columns') or [], xl, rep)
                offsets.append(0)
                jcols0 = J0.get('columns') or []
                inv_colmap = {v: k for k, v in colmap.items() if v is not None}
                ui0 = J0.get('ui') or {}
                need = set(ui0.get('facets') or []) | set(ui0.get('facets_extra') or [])
                for k in ('alias_col', 'summary_col', 'link_col'):
                    if isinstance(ui0.get(k), int):
                        need.add(ui0[k])
                for k in ('order', 'sort_key'):
                    for a_, b_ in (ui0.get(k) or {}).items():
                        need.add(int(a_))
                        if isinstance(b_, int):
                            need.add(b_)
                for p in ui0.get('presets') or []:
                    if isinstance(p.get('col'), int):
                        need.add(p['col'])
                uicols = {jc: [] for jc in need if jc in inv_colmap}
            continue
        if ridx > last:
            for cell in cells:
                if cell['value'] is not None:
                    rep.err('subgrid_missing', sid, 'value below autofilter range', where=coord(ridx, cell['column']))
            continue
        g = ridx - first
        vals = [None] * (ncols + 1)
        sids = [0] * (ncols + 1)
        dtcols = set()
        for cell in cells:
            c = cell['column']
            if c > ncols:
                if cell['value'] is not None:
                    rep.err('subgrid_missing', sid, 'value right of autofilter range', where=coord(ridx, c))
                continue
            sids[c] = cell['style_id'] or 0
            v = cell['value']
            if v is not None:
                nf = st.get(cell['style_id'] or 0).nf
                col_nf[c][nf] += 1
                counta[c] += 1
                if cell['data_type'] == 'f':
                    v = Formula(v)
                elif isinstance(v, (dt.datetime, dt.date, dt.time)):
                    v = fmt_datetime(v, nf)
                    dtcols.add(c)
            vals[c] = v
        colB.append(vals[2])
        if sid == '22':
            colD.append(vals[4])
        if zebra_rule:
            zc, zv, zbg = zebra_rule
            zon = is_num(vals[zc]) and vals[zc] == zv
            exp_zebra_rows += zon
        if not have_json:
            continue
        # locate part
        while cur['J'] is not None and g >= cur['off'] + cur['n'] and cur['k'] + 1 < len(parts):
            load_part(cur['k'] + 1)
            offsets.append(cur['off'])
            # chunk alignment: alias must not continue across a part boundary
            J = cur['J']
            if jalias and J.get('rows'):
                a0 = decode_act(J['rows'][0][1])[0] if len(J['rows'][0]) > 1 else None
                if a0 == jalias[-1]:
                    rep.err('chunk_alignment', sid, 'device alias split across parts', where='part %d' % cur['k'], act=a0)
        J = cur['J']
        li = g - cur['off']
        rows = J.get('rows') or []
        if li < 0 or li >= len(rows):
            rep.err('row_count', sid, 'Excel row has no JSON row (part offsets / counts)', where='row %d' % ridx)
            continue
        jr = rows[li]
        jr_all += 1
        jb = colmap.get(2)
        jalias.append(decode_act(jr[jb])[0] if jb is not None and jb < len(jr) else None)
        if sid == '22':
            jd = colmap.get(4)
            jblock.append(decode_act(jr[jd])[0] if jd is not None and jd < len(jr) else None)
        # zebra (json side)
        if zebra_rule:
            ui = J.get('ui') or {}
            rs = cur['rs']
            gcol = ui.get('group_col')
            zcol = ui.get('group_zebra') or ui.get('zebra')
            if li in rs:
                jbg = rs[li].get('bg')
            elif gcol is not None and zcol:
                key = decode_act(jr[gcol])[0] if gcol < len(jr) else None
                if key != jz_state['prev']:
                    jz_state['par'] += 1
                    jz_state['prev'] = key
                jbg = normc(zcol) if jz_state['par'] % 2 == 1 else None
            else:
                jbg = None
            if (normc(zbg) if zon else None) != jbg:
                zebra_bad += 1
                rep.err('cf_style', sid, 'device zebra fill differs ($%s=1 -> %s)' % (num2col(zc), zbg), where='row %d' % ridx,
                        exp=normc(zbg) if zon else None, act=jbg)
        else:
            zon, jbg = False, None
        for jc_, lst in uicols.items():
            lst.append(vals[inv_colmap[jc_]])
        if quick and not (li < 500):
            continue
        for c in range(1, ncols + 1):
            jc = colmap.get(c)
            if jc is None:
                continue
            raw = vals[c]
            act = jr[jc] if jc < len(jr) else None
            si = st.get(sids[c]) if sids[c] else None
            # static style of every cell (bold / italic / underline / fill / font colour) incl. empty fills
            e_ = static_expect(si, False)
            a_ = eff_json(cur['cs_'], cur['rs'], cur['cp'], li, jc)
            probs = [k for k in ('b', 'i', 'u') if raw is not None and e_[k] != a_[k]]
            if raw is not None and _nf_(e_['fc']) != _nf_(a_['fc']):
                probs.append('fc')
            ebg = normc(zbg) if zon else _nb(e_['bg'])        # CF fill wins over the static fill in Excel
            abg = _nb(a_['bg']) or jbg                           # page: column/row/cell fill, else zebra
            if ebg != abg:
                probs.append('bg')
            if probs:
                rep.err('static_style', sid, 'effective cell style differs (%s)' % ','.join(probs), where=coord(ridx, c),
                        exp={k: e_.get(k) for k in probs}, act={k: a_.get(k) for k in probs})
            if raw is None and act is None:
                continue
            ncmp += 1
            if isinstance(raw, Formula):
                rep.uncheck(sid, 'formula in data row %s not evaluated in streaming mode' % coord(ridx, c))
                continue
            kind = 'datetime' if c in dtcols else 'plain'
            for cat, msg, e, a in compare_cell(ctx, raw, kind, act):
                rep.err(cat, sid, msg, where=coord(ridx, c), exp=e, act=a)
            if raw is not None:
                agg[c].add(raw, kind, si)
                if is_num(raw):
                    af_ = decode_act(act)[3]
                    jf = af_ if af_ is not None else jcols0[jc].get('fmt')
                    nf = si.nf if si else 'General'
                    if norm_nf(jf) != norm_nf(nf):
                        rep.err('cell_numfmt', sid, 'effective number format of numeric cell differs', where=coord(ridx, c), exp=nf, act=jf)
    # --- meta, counts, parts
    nx = last - first + 1
    if have_json:
        tot = 0
        for k, pm, n, fn in part_meta:
            pm = pm or {}
            if pm.get('index') != k or pm.get('of') != len(parts):
                rep.err('chunk_meta', sid, 'part index/of wrong', where=fn, exp=(k, len(parts)), act=(pm.get('index'), pm.get('of')))
            if pm.get('row_offset') != tot:
                rep.err('chunk_meta', sid, 'part row_offset not cumulative', where=fn, exp=tot, act=pm.get('row_offset'))
            tot += n
        if len(part_meta) != len(parts):
            rep.err('chunk_meta', sid, 'not all parts were consumed (Excel rows ended early?)', exp=len(parts), act=len(part_meta))
        if tot != nx and len(part_meta) == len(parts):
            rep.err('row_count', sid, 'sum of part rows differs from Excel data rows', exp=nx, act=tot)
        po = (entry or {}).get('part_offsets')
        if po != [pm.get('row_offset') for _, pm, _, _ in part_meta if pm] and len(part_meta) == len(parts):
            rep.err('manifest', sid, 'manifest part_offsets differ from parts row_offset', exp=[pm.get('row_offset') for _, pm, _, _ in part_meta if pm], act=po)
        J0 = load_json(parts[0])
        check_table_meta(ctx, sd, J0, sid, ncols, colmap, col_nf, rep)
        if not quick:
            check_table_columns(ctx, sd, J0, sid, ncols, colmap, agg, rep)
        out['uicols'] = uicols
        out['colmap'] = colmap
        for k in range(1, len(parts)):
            Jk = load_json(parts[k])
            for key in ('columns', 'title', 'bands'):
                if Jk.get(key) != J0.get(key):
                    rep.warn('chunk_meta', sid, 'part metadata %s differs from part 0' % key, where=os.path.basename(parts[k]))
        if zebra_rule and not ((J0.get('ui') or {}).get('group_zebra') or J0.get('row_styles')):
            rep.err('cf_style', sid, 'no zebra representation (ui.group_zebra / row_styles) for CF $%s=1' % num2col(zebra_rule[0]))
    if zebra_rule:
        exp_n = EXPECTED_CF_COUNTS.get(sid, {}).get(1, (None, None))[1]
        if exp_n is not None and exp_zebra_rows != exp_n:
            rep.warn('selfcheck_cf_count', sid, 'zebra rows counted by verifier differ from spec', exp=exp_n, act=exp_zebra_rows)
        rep.stats['zebra'][sid] = {'expected_shaded_rows': exp_zebra_rows, 'mismatched_rows': zebra_bad}
    rep.stats['cells_compared'][sid] = ncmp
    rep.stats['timing'][sid + '_worker_s'] = round(time.time() - t0, 1)
    # alias blocks must be contiguous (spec) -> check on Excel
    seen, prev = set(), None
    for a in colB:
        if a != prev:
            if a in seen:
                rep.warn('selfcheck_blocks', sid, 'alias block not contiguous in Excel', act=a)
            seen.add(a)
            prev = a
    out.update(report=rep.state(), colB=colB, colD=colD, counta=dict(counta), header={r: sd.rows.get(r) for r in (1, 2, 3, 4)},
               jalias=jalias, jblock=jblock, max_row=last, json_rows=jr_all, have_json=have_json)
    return out


# ----------------------------------------------------------------------------------------------
# manifest / links / expected numbers
# ----------------------------------------------------------------------------------------------
def check_manifest(ctx, man, rep):
    sheets = man.get('sheets') or []
    ids = [str(s.get('id')) for s in sheets]
    exp_ids = [n[:2] for n in ctx.book.names]
    if ids != exp_ids:
        rep.err('manifest', '-', 'sheet ids/order differ from workbook', exp=exp_ids, act=ids)
    for s in sheets:
        sid = str(s.get('id'))
        name = ctx.book.by_id.get(sid)
        if name is None:
            rep.err('manifest', sid, 'unknown sheet id')
            continue
        if s.get('name') != name:
            rep.err('manifest', sid, 'name differs', exp=name, act=s.get('name'))
        if s.get('mode') != expected_mode(sid):
            rep.err('manifest', sid, 'mode differs from expected mode map', exp=expected_mode(sid), act=s.get('mode'))
        files = s.get('files') or []
        if not files:
            rep.err('manifest', sid, 'no files listed')
        tot = 0
        for f in files:
            p = os.path.join(ctx.data_dir, f)
            if not os.path.exists(p):
                rep.err('missing_file', sid, 'file listed in manifest does not exist', where=f)
            else:
                tot += os.path.getsize(p)
        if s.get('bytes') is not None and s.get('bytes') != tot:
            rep.warn('manifest', sid, 'bytes differs from file sizes', exp=tot, act=s.get('bytes'))
        if expected_mode(sid) == 'table':
            nx = ctx.last_row[sid] - ctx.first_row[sid] + 1
            if s.get('rows') != nx:
                rep.err('manifest', sid, 'rows differs from Excel data rows', exp=nx, act=s.get('rows'))
            nc = ctx.ncols[sid]
            if s.get('cols') != nc:
                if sid in DROPPABLE_COLS and s.get('cols') == nc - len(DROPPABLE_COLS[sid]):
                    rep.warn('approximation', sid, 'manifest cols = Excel cols − dropped parity column', exp=nc, act=s.get('cols'))
                else:
                    rep.err('manifest', sid, 'cols differs from Excel', exp=nc, act=s.get('cols'))
            if EXPECTED_ROWS.get(sid) != nx:
                rep.warn('selfcheck_rows', sid, 'Excel data rows differ from spec', exp=EXPECTED_ROWS.get(sid), act=nx)
            if EXPECTED_COLS.get(sid) != nc:
                rep.warn('selfcheck_cols', sid, 'Excel columns differ from spec', exp=EXPECTED_COLS.get(sid), act=nc)
            if EXPECTED_FREEZE.get(sid) != ctx.book.xml[name].xsplit:
                rep.warn('selfcheck_freeze', sid, 'pane xSplit differs from spec', exp=EXPECTED_FREEZE.get(sid), act=ctx.book.xml[name].xsplit)
        if sid in CHUNKED:
            po = s.get('part_offsets')
            if not isinstance(po, list) or len(po) != len(files) or (po and po[0] != 0):
                rep.err('manifest', sid, 'part_offsets missing or inconsistent with files', act=po)
            elif any(b <= a for a, b in zip(po, po[1:])):
                rep.err('manifest', sid, 'part_offsets not increasing', act=po)
    # descriptive fields: 00 TOC row (C purpose, D row unit, E rows, F sources), tab colours, groups
    toc = {}
    s00 = ctx.book.sheet('00')
    if s00 is not None:
        for r in sorted(s00.rows):
            if isinstance(s00.get(r, 1), Formula):
                v = ctx.eng.cellval(s00.name, r, 1)
                if isinstance(v, LinkVal) and v.label in ctx.book.names:
                    toc[v.label[:2]] = (s00.get(r, 2), s00.get(r, 3), s00.get(r, 4), s00.get(r, 5))
    groups = man.get('groups') or []
    for s in sheets:
        sid = str(s.get('id'))
        name = ctx.book.by_id.get(sid)
        if name is None:
            continue
        t = toc.get(sid)
        if t is None:
            rep.err('manifest', sid, 'sheet not found in the 00 table of contents')
        else:
            for k, v in zip(('desc', 'row_unit', 'toc_rows', 'sources'), t):
                if s.get(k) != v:
                    rep.err('manifest', sid, 'manifest %s differs from the 00 TOC row' % k, exp=v, act=s.get(k))
        m = re.search(r'<tabColor [^>]*rgb="([0-9A-Fa-f]{6,8})"', ctx.book.xml[name].head)
        if normc(m.group(1) if m else None) != normc(s.get('tab_color')):
            rep.err('manifest', sid, 'tab_color differs from <sheetPr><tabColor>', exp=m.group(1) if m else None, act=s.get('tab_color'))
        if s.get('group') not in groups:
            rep.err('manifest', sid, 'group not in manifest groups', act=s.get('group'))
        if s.get('short') and s.get('short') != name[3:]:
            rep.warn('manifest', sid, 'short name is not the sheet name without its id', exp=name[3:], act=s.get('short'))
        if expected_mode(sid) == 'grid':
            J = ctx.json.get(sid)
            if J is not None and s.get('rows') != J.get('max_row'):
                rep.err('manifest', sid, 'grid rows differ from max_row', exp=J.get('max_row'), act=s.get('rows'))
    if [str(s.get('group')) for s in sheets] and len(set(s.get('group') for s in sheets)) != len(groups):
        rep.warn('manifest', '-', 'some manifest groups are empty', act=groups)
    wbm = man.get('workbook') or {}
    for k, (sidx, r, c) in {'title': ('00', 1, 2), 'subtitle': ('00', 2, 2), 'summary_title': ('01', 1, 2), 'summary_subtitle': ('01', 2, 2)}.items():
        sdx = ctx.book.sheet(sidx)
        if sdx is not None and k in wbm and wbm[k] != sdx.get(r, c):
            rep.err('manifest', '-', 'workbook.%s differs from %s %s' % (k, sidx, coord(r, c)), exp=sdx.get(r, c), act=wbm[k])
    se = man.get('search') or {}
    if str(se.get('index_sheet')) != '05' or se.get('key_col') != 0 or se.get('alias_col') != 4:
        rep.err('manifest', '-', 'search block differs from contract', exp={'index_sheet': '05', 'key_col': 0, 'alias_col': 4}, act=se)


def target_rowcount(ctx, sid):
    if sid in CHUNKED:
        b = ctx.big.get(sid)
        return (b['last'] - b['first'] + 1) if b else None
    J = ctx.json.get(sid)
    return len(J.get('rows') or []) if J else None


def target_value(ctx, sid, r, col):
    if sid in CHUNKED:
        b = ctx.big.get(sid)
        if not b or not b.get('have_json'):
            raise NotLoaded('chunked JSON %s not available' % sid)
        lst = b['jalias'] if col == 1 else (b['jblock'] if col == 3 and sid == '22' else None)
        if lst is None:
            raise NotLoaded('column %d of %s not kept' % (col, sid))
        return lst[r] if r < len(lst) else '<out of range>'
    rows = (ctx.json.get(sid) or {}).get('rows') or []
    if r >= len(rows):
        return '<out of range>'
    return _jv(rows[r], col)


def check_links(ctx, rep):
    n = 0
    for (src, where, l) in ctx.all_links:
        n += 1
        s = str(l.get('s'))
        r = l.get('r')
        name = ctx.book.by_id.get(s)
        if name is None:
            rep.err('link_invalid', src, 'link target sheet does not exist', where=where, act=l)
            continue
        mode = expected_mode(s)
        if mode == 'card':
            if r is not None:
                rep.err('link_invalid', src, 'card target must have r=null', where=where, act=l)
        elif mode == 'grid':
            mr = ctx.book.sheets[name].max_row if name in ctx.book.sheets else None
            if not isinstance(r, int) or (mr and not 1 <= r <= mr):
                rep.err('link_invalid', src, 'grid target row out of range', where=where, act=l)
        else:
            if r is not None:
                nr = target_rowcount(ctx, s)
                if not isinstance(r, int) or r < 0 or (nr is not None and r >= nr):
                    rep.err('link_invalid', src, 'table target row out of range', where=where, act=l, exp=nr)
    rep.stats['links']['checked'] = n
    ns, nbad = 0, 0
    for (src, where, tsid, r, keys, desc) in ctx.semantic:
        for col, key in keys:
            ns += 1
            try:
                v = target_value(ctx, tsid, r, col)
            except NotLoaded as e:
                rep.uncheck(src, 'semantic link %s: %s' % (where, e))
                continue
            if str(v).upper() != str(key).upper() and not (is_num(v) and is_num(key) and num_eq(v, key)):
                nbad += 1
                rep.err('link_semantic', src, 'link lands on a row whose key differs (%s)' % desc, where='%s → %s r=%s col %d' % (where, tsid, r, col),
                        exp=key, act=v)
    rep.stats['links']['semantic_checked'] = ns
    rep.stats['links']['semantic_bad'] = nbad


def _grid_cell(J, r, c):
    for row in J.get('rows') or []:
        if row.get('r') == r:
            for cell in row.get('cells') or []:
                if cell.get('c') == c:
                    return cell
    return None


def check_expected(ctx, man, rep):
    eng = ctx.eng
    # engine self-check against spec numbers
    for sid, cells in EXPECTED_CELLS.items():
        sd = ctx.book.sheet(sid)
        for co, ev in cells.items():
            r, c = parse_coord(co)
            try:
                v = eng.cellval(sd.name, r, c)
            except (NotLoaded, Unsupported) as e:
                rep.warn('selfcheck_expected', sid, 'cannot evaluate %s: %s' % (co, e))
                continue
            if not (is_num(v) and num_eq(v, ev)):
                rep.warn('selfcheck_expected', sid, 'verifier evaluation differs from spec value', where=co, exp=ev, act=v)
            J = ctx.json.get(sid)
            if not J:
                continue
            if expected_mode(sid) == 'grid':
                cell = _grid_cell(J, r, c)
                got = decode_act(cell.get('v'))[0] if cell else None
                if not (is_num(got) and num_eq(got, ev)):
                    rep.err('expected_value', sid, 'JSON value differs from spec value', where=co, exp=ev, act=got)
            else:
                hc = J.get('header_cells') or []
                if not any(is_num(decode_act(h.get('v'))[0]) and num_eq(decode_act(h.get('v'))[0], ev) for h in hc):
                    rep.err('expected_value', sid, 'header KPI value not found in header_cells', where=co, exp=ev)
    # all verification formulas evaluate to ✓ (14 G6:G22, L29:L37; 24 IF cells)
    for sid in ('14', '24'):
        sd = ctx.book.sheet(sid)
        J = ctx.json.get(sid)
        nf, nck = 0, 0
        for r, row in sd.rows.items():
            for ci, raw in enumerate(row):
                if isinstance(raw, Formula) and raw.text.startswith('=IF('):
                    nf += 1
                    try:
                        v = eng.cellval(sd.name, r, ci + 1)
                    except Exception:
                        v = None
                    if v != '✓':
                        rep.warn('selfcheck_expected', sid, 'IF verification formula not ✓ in verifier', where=coord(r, ci + 1), act=v)
                    if J:
                        cell = _grid_cell(J, r, ci + 1)
                        if not cell or cell.get('v') != v:
                            rep.err('expected_value', sid, 'IF ✓/✗ cell differs', where=coord(r, ci + 1), exp=v, act=cell and cell.get('v'))
                        else:
                            nck += 1
        rep.stats['expected']['%s_if_cells' % sid] = '%d/%d' % (nck, nf)
    # 01 KPIs
    J = ctx.json.get('01')
    if J:
        for co, (lab, s, rr, f) in EXPECTED_01_KPI.items():
            r, c = parse_coord(co)
            cell = _grid_cell(J, r, c)
            if not cell:
                rep.err('expected_value', '01', 'KPI cell missing', where=co)
                continue
            v = cell.get('v')
            l = cell.get('l')
            if isinstance(v, dict):
                v, l = v.get('t'), v.get('l')
            if not val_eq(lab, v):
                rep.err('expected_value', '01', 'KPI label differs', where=co, exp=lab, act=v)
            if not isinstance(l, dict) or str(l.get('s')) != s or l.get('r') != rr:
                rep.err('expected_value', '01', 'KPI link differs', where=co, exp={'s': s, 'r': rr}, act=l)
            if f:
                sty = grid_style_of(J, cell)
                ff = cell.get('f') or sty.get('f') or sty.get('fmt')
                if ff != f:
                    rep.err('expected_value', '01', 'KPI number format differs', where=co, exp=f, act=ff)
    # 00 TOC D108:D135 row counts vs manifest
    sd = ctx.book.sheet('00')
    ment = {str(s.get('id')): s for s in man.get('sheets') or []}
    for r in range(108, 136):
        a = sd.get(r, 1)
        lab = eng.cellval(sd.name, r, 1) if isinstance(a, Formula) else a
        lab = lab.label if isinstance(lab, LinkVal) else lab
        d = sd.get(r, 4)
        sid = str(lab)[:2]
        m = ment.get(sid)
        if not m:
            continue
        if expected_mode(sid) == 'table':
            if m.get('rows') != d:
                rep.err('expected_rowcount', sid, 'manifest rows differ from 00_說明 TOC D%d' % r, exp=d, act=m.get('rows'))
        else:
            rep.add('info', 'toc_rows', sid, '00 TOC D%d=%s (layout rows), manifest rows=%s' % (r, d, m.get('rows')))
    # card default
    if '03' in ctx.json and '05' in ctx.json and '02' in ctx.json:
        sim = card_sim_json(ctx, ctx.json['02'], CARD_DEFAULT[0])
        if (sim['alias'], sim['i3']) != CARD_DEFAULT[1:]:
            rep.err('expected_value', '02', 'default query does not resolve to D00759 / 03 idx 229', exp=CARD_DEFAULT, act=(sim['alias'], sim['i3']))


# ----------------------------------------------------------------------------------------------
# report output
# ----------------------------------------------------------------------------------------------
def write_reports(rep, out_json, out_md, meta):
    by_cat = collections.Counter()
    by_sheet = collections.Counter()
    for (lvl, cat, sh), n in rep.counts.items():
        if lvl == 'error':
            by_cat[cat] += n
            by_sheet[sh] += n
    entries = []
    for k, n in sorted(rep.counts.items(), key=lambda kv: ({'error': 0, 'warning': 1, 'info': 2}.get(kv[0][0], 3), kv[0][2], -kv[1])):
        entries.append({'level': k[0], 'category': k[1], 'sheet': k[2], 'count': n, 'samples': rep.samples.get(k, [])})
    summary = {'errors': rep.total('error'), 'warnings': rep.total('warning'), 'infos': rep.total('info'),
               'errors_by_category': dict(by_cat.most_common()), 'errors_by_sheet': dict(sorted(by_sheet.items())),
               'unchecked': len(rep.unchecked)}
    doc = {'summary': summary, 'meta': meta, 'entries': entries, 'unchecked': rep.unchecked[:500],
           'notes': rep.notes, 'stats': rep.stats}
    with open(out_json, 'w', encoding='utf-8') as f:
        json.dump(doc, f, ensure_ascii=False, indent=1, default=str)
    L = []
    L.append('# AMS data verification report\n')
    L.append('- xlsx: `%s`\n- data: `%s`\n- mode: %s\n- run: %s, %.1f s\n' % (meta['xlsx'], meta['data'], meta['mode'], meta['started'], meta['seconds']))
    L.append('\n## Summary\n\n**%d errors, %d warnings, %d info, %d unchecked items**\n' % (summary['errors'], summary['warnings'], summary['infos'], summary['unchecked']))
    if by_cat:
        L.append('\n| error category | count |\n|---|---|\n' + ''.join('| %s | %d |\n' % kv for kv in by_cat.most_common()))
        L.append('\n| sheet | errors |\n|---|---|\n' + ''.join('| %s | %d |\n' % kv for kv in sorted(by_sheet.items())))
    for lvl in ('error', 'warning', 'info'):
        es = [e for e in entries if e['level'] == lvl]
        if not es:
            continue
        L.append('\n## %ss\n' % lvl.capitalize())
        for e in es:
            L.append('\n### [%s] %s — %s × %d\n' % (e['sheet'], e['category'], lvl, e['count']))
            for s in e['samples']:
                w = (' @%s' % s['where']) if s.get('where') else ''
                x = ''
                if 'expected' in s:
                    x = ' — expected `%s` / actual `%s`' % (str(s['expected']).replace('`', "'"), str(s['actual']).replace('`', "'"))
                L.append('- %s%s%s\n' % (s['msg'], w, x))
    if rep.unchecked:
        L.append('\n## Unchecked / partially checked (explicit)\n')
        for u in rep.unchecked[:200]:
            L.append('- %s\n' % u)
    if rep.notes:
        L.append('\n## Notes\n')
        for u in rep.notes:
            L.append('- %s\n' % u)
    L.append('\n## Stats\n\n```\n%s\n```\n' % json.dumps(rep.stats, ensure_ascii=False, indent=1, default=str)[:20000])
    with open(out_md, 'w', encoding='utf-8') as f:
        f.write(''.join(L))
    return summary


STATIC_UNCHECKED = [
    # what is still only partially checked, after the data-integrator pass (everything else is compared exactly)
    'colour-scale fills compared with ±1/255 per channel for rounding (stats.colour_scale_channel_diff shows the real spread); '
    'Excel ties (mid==max) resolved as max colour',
    'data-bar fraction p compared to (v-min)/(max-min) with 4-decimal rounding allowed (±0.00005)',
    'grid overflow ov: hard invariants are errors (text, unwrapped, general/left, unmerged, empty unstyled neighbours); '
    'whether ov is needed at all uses a text-width estimate (CJK 1 em, other 0.5 em) and only warns',
    'ui.facets / facets_extra / search_cols / key_col choices are presentation hints from the specs: indexes are range-checked, '
    'alias_col contents, order, facet_options, facet_values, summary_rows, sort_key, presets and validations are checked against the data',
    'front-end rendering itself is out of scope here (see the browser tests); this report checks the JSON the page consumes',
]


def main():
    ap = argparse.ArgumentParser(description='Independent verifier: xlsx vs docs/data')
    ap.add_argument('xlsx')
    ap.add_argument('data')
    ap.add_argument('--quick', action='store_true', help='sample rows instead of exhaustive comparison')
    ap.add_argument('--samples', type=int, default=8)
    ap.add_argument('--charts-contract', default=DEFAULT_CHARTS_CONTRACT)
    ap.add_argument('--no-workers', action='store_true', help='skip 21/22 streaming')
    ap.add_argument('--out', default=HERE)
    a = ap.parse_args()
    t0 = time.time()
    started = time.strftime('%Y-%m-%d %H:%M:%S')
    rep = Report(a.samples)
    for u in STATIC_UNCHECKED:
        rep.uncheck('-', u)
    meta = {'xlsx': a.xlsx, 'data': a.data, 'mode': 'quick' if a.quick else 'exhaustive', 'started': started}
    out_json = os.path.join(a.out, 'verify_report.json')
    out_md = os.path.join(a.out, 'verify_report.md')
    man_path = os.path.join(a.data, 'manifest.json')
    if not os.path.exists(man_path):
        rep.err('missing_file', '-', 'manifest.json missing', where=man_path)
        meta['seconds'] = time.time() - t0
        s = write_reports(rep, out_json, out_md, meta)
        print('manifest.json missing -> %d error(s)' % s['errors'])
        return 1
    man = load_json(man_path)
    entries = {str(s.get('id')): s for s in man.get('sheets') or []}
    print('[verify] opening workbook …', flush=True)
    book = Book(a.xlsx)
    pool = None
    asyncs = {}
    if not a.no_workers:
        pool = mp.get_context('spawn').Pool(2)
        for sid in sorted(CHUNKED):
            asyncs[sid] = pool.apply_async(big_worker, ((a.xlsx, a.data, sid, a.quick, a.samples, entries.get(sid)),))
    book.open_wb()
    for n in book.names:
        if n[:2] not in CHUNKED:
            book.load(n)
    print('[verify] workbook loaded %.1fs' % (time.time() - t0), flush=True)
    eng = Engine(book)
    ctx = Ctx(book, eng, rep, man, a.data, a.quick)
    for sid, s in entries.items():
        if sid in CHUNKED:
            continue
        files = s.get('files') or []
        Js = []
        for f in files:
            p = os.path.join(a.data, f)
            if os.path.exists(p):
                try:
                    Js.append(load_json(p))
                except Exception as e:
                    rep.err('json_invalid', sid, 'cannot parse JSON: %s' % e, where=f)
        if not Js:
            continue
        J = Js[0]
        if len(Js) > 1 and 'rows' in J:
            J = dict(J)
            J['rows'] = [r for x in Js for r in (x.get('rows') or [])]
        ctx.json[sid] = J
    check_manifest(ctx, man, rep)
    # 21/22 worker results
    for sid, fut in asyncs.items():
        try:
            res = fut.get()
        except Exception as e:
            rep.err('verifier_crash', sid, 'worker failed: %r' % e)
            continue
        rep.merge(res['report'])
        ctx.big[sid] = res
        name = book.by_id[sid]
        sd = SheetData(name, sid, book.names.index(name), book.xml[name])
        for r, vals in res['header'].items():
            if vals is not None:
                sd.rows[int(r)] = vals
        sd.colstore = {2: res['colB']}
        if res['colD']:
            sd.colstore[4] = res['colD']
        sd.max_row = res['max_row']
        book.sheets[name] = sd
        if res.get('have_json'):
            rep.stats['rows'][sid] = res['json_rows']
            try:
                J0 = load_json(os.path.join(a.data, entries[sid]['files'][0]))
                uic = res.get('uicols') or {}

                def colvals(jc, uic=uic):
                    if jc not in uic:
                        raise KeyError('column %s not collected by worker' % jc)
                    return uic[jc]
                check_ui(ctx, sd, J0, sid, ctx.first_row[sid], ctx.last_row[sid], ctx.ncols[sid], res.get('colmap') or {},
                         rep, colvals=colvals)
            except Exception as e:
                import traceback
                rep.err('verifier_crash', sid, 'check_ui (chunked): %r' % e, where=traceback.format_exc()[-600:])
    if pool:
        pool.close()
    print('[verify] workers done %.1fs' % (time.time() - t0), flush=True)
    contract = None
    if a.charts_contract and os.path.exists(a.charts_contract):
        contract = {c['id']: c for c in load_json(a.charts_contract)}
    else:
        rep.uncheck('-', 'charts_contract.json not found; charts checked against xlsx only')
    xcharts = parse_charts(ctx)
    for n in book.names:
        sid = n[:2]
        if sid in CHUNKED:
            continue
        sd = book.sheets[n]
        J = ctx.json.get(sid)
        t1 = time.time()
        cfres = None
        if sd.xml.cf and expected_mode(sid) != 'card':
            cfres = compute_cf(eng, sd, book.dxfs)
            check_cf_counts(rep, sid, cfres)
        if J is None:
            rep.err('missing_file', sid, 'no JSON for sheet')
            continue
        if J.get('mode') not in (None, expected_mode(sid)):
            rep.err('meta', sid, 'JSON mode differs', exp=expected_mode(sid), act=J.get('mode'))
        try:
            m = expected_mode(sid)
            if m == 'table':
                verify_table(ctx, sd, sid, J, rep, cfres)
            elif m == 'grid':
                verify_grid(ctx, sd, sid, J, rep, cfres)
                if xcharts.get(sid) or J.get('charts') or sid in EXPECTED_CHART_COUNT:
                    verify_charts(ctx, sid, J, xcharts.get(sid, []), contract, rep)
            else:
                verify_card(ctx, J, rep)
        except Exception as e:
            import traceback
            rep.err('verifier_crash', sid, 'exception while verifying: %r' % e, where=traceback.format_exc()[-600:])
        rep.stats['timing'][sid + '_s'] = round(time.time() - t1, 2)
    check_links(ctx, rep)
    try:
        check_expected(ctx, man, rep)
    except Exception as e:
        import traceback
        rep.err('verifier_crash', '-', 'check_expected: %r' % e, where=traceback.format_exc()[-600:])
    if eng.unsupported:
        rep.warn('selfcheck_engine', '-', 'unsupported functions encountered', act=dict(eng.unsupported))
    rep.stats['colour_scale_channel_diff'] = {str(k): v for k, v in sorted(SCALE_DIFF.items())}
    meta['seconds'] = round(time.time() - t0, 1)
    s = write_reports(rep, out_json, out_md, meta)
    print('\n==== verify_data summary (%s, %.1fs) ====' % (meta['mode'], meta['seconds']))
    print('errors: %d   warnings: %d   info: %d   unchecked: %d' % (s['errors'], s['warnings'], s['infos'], s['unchecked']))
    for k, v in s['errors_by_category'].items():
        print('  %-32s %d' % (k, v))
    print('by sheet:', ', '.join('%s:%d' % kv for kv in s['errors_by_sheet'].items()))
    print('report: %s\n        %s' % (out_json, out_md))
    return 0 if s['errors'] == 0 else 1


if __name__ == '__main__':
    mp.freeze_support()
    sys.exit(main())
