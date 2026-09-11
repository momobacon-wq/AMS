# -*- coding: utf-8 -*-
"""mode = "grid" sheet builder (layout sheets 00, 01, 14, 16, 17, 18, 24 and table sub-grids)."""
import datetime as dt
import unicodedata

from amsx_common import col_px, row_px, fmt_number
from .formula import Formula, Link, XlError
from .values import plain_value
from .cf import evaluate_cf
from . import config

MONO_FONTS = ("consolas", "courier new", "courier", "lucida console", "cascadia mono", "cascadia code", "menlo")
VA_MAP = {"top": "top", "center": "middle", "bottom": "bottom", "justify": "middle", "distributed": "middle"}
HA_MAP = {"left": "left", "center": "center", "right": "right", "justify": "justify", "fill": "left",
          "centerContinuous": "center", "distributed": "center"}


def _norm_fc(c):
    return None if c in (None, "#000000") else c


class GridStyles(object):
    def __init__(self):
        self.list = []
        self.idx = {}

    def get(self, d):
        key = tuple(sorted((k, v if not isinstance(v, dict) else tuple(sorted(v.items()))) for k, v in d.items()))
        i = self.idx.get(key)
        if i is None:
            i = self.idx[key] = len(self.list)
            self.list.append(d)
        return i


def _border(sides):
    out = {}
    for k in ("t", "r", "b", "l"):
        if k in sides:
            st, col = sides[k]
            out[k] = st + (" " + col if col else "")
    return out


def cell_style_dict(xf, ov, is_link=False):
    f = xf.font
    d = {}
    b = bool(f.b)
    it = bool(f.i)
    u = bool(f.u)
    fc = _norm_fc(f.color)
    bg = xf.fill
    bd = dict(xf.border) if xf.border else {}
    if ov:
        if "b" in ov:
            b = bool(ov["b"])
        if "i" in ov:
            it = bool(ov["i"])
        if "u" in ov:
            u = bool(ov["u"])
        if "fc" in ov:
            fc = _norm_fc(ov["fc"])
        if "bg" in ov:
            bg = ov["bg"]
        elif "bg_white" in ov:
            bg = None
        if "bd" in ov:
            bd.update(ov["bd"])
    if b:
        d["b"] = 1
    if it:
        d["i"] = 1
    if u:
        d["u"] = 1
    if f.strike:
        d["s"] = 1
    if bg:
        d["bg"] = bg
    if fc:
        d["fc"] = fc
    if f.sz is not None and f.sz != 11:
        d["sz"] = f.sz if f.sz != int(f.sz) else int(f.sz)
    ha = HA_MAP.get(xf.ha)
    if ha:
        d["ha"] = ha
    va = VA_MAP.get(xf.va)
    if va:
        d["va"] = va
    if xf.wrap:
        d["wrap"] = 1
    if xf.shrink:
        d["shrink"] = 1   # shrinkToFit (01 KPI numbers); informational, text fits at full size there
    if xf.indent:
        d["ind"] = xf.indent
    if xf.rotation:
        d["rot"] = xf.rotation
    if bd:
        d["bd"] = _border(bd)
    if (f.name or "").lower() in MONO_FONTS:
        d["mono"] = 1
    return d


def _text_px(s, sz_pt):
    px = (sz_pt or 11) * 4.0 / 3.0
    w = 0.0
    for ch in s:
        if ch == "\n":
            break
        ea = unicodedata.east_asian_width(ch)
        w += px if ea in ("W", "F") else px * 0.55
    return w


def build_grid_region(ctx, sheet_name, region, relative=False, charts=None):
    """Build a grid object for ``region`` (r1, c1, r2, c2) of a sheet.

    relative=True renumbers rows/cols from 1 (used for table sub-grids) and
    records the Excel origin.
    """
    ev, wb, resolver = ctx.ev, ctx.wb, ctx.resolver
    sh = wb.load_sheet(sheet_name)
    sid = sh.id
    r1, c1, r2, c2 = region
    mm = sh.merge_map()
    cf = evaluate_cf(ev, sheet_name, region)
    if not relative:
        ctx.cf_hits[sid] = [(p, t, None, h) for p, t, rg, h in cf.rule_hits]
    ov_cells = cf.cells
    gst = GridStyles()
    dr = r1 - 1 if relative else 0
    dc = c1 - 1 if relative else 0
    ncols = c2 - c1 + 1
    max_c_needed = c2
    if charts:
        for ch in charts:
            max_c_needed = max(max_c_needed, ch["anchor"]["to"]["c"])
    if not relative:
        # explicitly sized margin columns just right of the used range (e.g. 01 Z, width 2) are part of the layout
        for mn, mx, w, hid, ol in sh.cols:
            if w is not None and c2 < mn <= c2 + 3:
                max_c_needed = max(max_c_needed, min(mx, c2 + 3))
    col_widths = [col_px(sh.col_width(c)) for c in range(c1, max_c_needed + 1)]
    rows_out = []
    cellinfo = {}  # (r, c) -> dict for overflow pass
    for r in range(r1, r2 + 1):
        vals = sh.rows.get(r) or []
        srow = sh.srows.get(r) or []
        cells = []
        for c in range(c1, min(c2, len(vals)) + 1):
            m = mm.get((r, c))
            if m and m[0] == "covered":
                continue
            raw = vals[c - 1]
            s_id = srow[c - 1] if c <= len(srow) else 0
            xf = wb.styles.xfs[s_id]
            ov = ov_cells.get((r, c))
            link = None
            if isinstance(raw, Formula):
                val = ev.value(sheet_name, r, c)
                if isinstance(val, Link):
                    link = resolver.resolve(val.target, sid, "%s%d" % (_cl(c), r))
                    val = val.label
                elif isinstance(val, XlError):
                    ctx.warnings.append("%s %s%d: formula error %s" % (sid, _cl(c), r, val.code))
                if val == "":
                    val = None
            else:
                val = raw
            sd = cell_style_dict(xf, ov, link is not None)
            has_bg = "bg" in sd or "bd" in sd
            if val is None and not has_bg and not (ov and "bar" in ov):
                continue
            cell = {"c": c - dc}
            nf = xf.numfmt
            if isinstance(val, (dt.datetime, dt.date, dt.time)):
                cell["v"] = plain_value(val, nf)
            else:
                cell["v"] = plain_value(val, nf) if val is not None else None
            if sd:
                cell["s"] = gst.get(sd)
            if m and m[0] == "anchor":
                if m[2] > 1:
                    cell["cs"] = m[2]
                if m[1] > 1:
                    cell["rs"] = m[1]
            if link is not None:
                if "u" in link:
                    cell["l"] = {"u": link["u"]}
                else:
                    cell["l"] = link
            if isinstance(val, (int, float)) and not isinstance(val, bool) and nf not in ("General", "@"):
                cell["f"] = nf
                if ";" in nf:  # sectioned format (e.g. +0.0;-0.0;0.0): pre-rendered display text
                    dtxt = fmt_number(val, nf)
                    if dtxt is not None:
                        cell["d"] = dtxt
            if ov and "bar" in ov:
                cell["bar"] = {"p": ov["bar"][0], "c": ov["bar"][1]}
            cells.append(cell)
            cellinfo[(r, c)] = (cell, sd)
        ht = sh.row_ht.get(r)
        h = row_px(ht[0]) if ht and ht[1] else None
        if cells or h is not None or r in sh.hidden_rows:
            ro = {"r": r - dr, "h": h, "cells": cells}
            rows_out.append(ro)
    # overflow: unwrapped text spilling into empty, unstyled right neighbours
    for (r, c), (cell, sd) in cellinfo.items():
        v = cell.get("v")
        if not isinstance(v, str) or sd.get("wrap") or cell.get("cs") or sd.get("ha") not in (None, "left"):
            continue
        need = _text_px(v, sd.get("sz", 11)) + 6
        have = col_widths[c - c1] if c - c1 < len(col_widths) else col_px(None)
        if need <= have:
            continue
        n = 0
        k = c + 1
        while k <= max_c_needed and need > have:
            if (r, k) in cellinfo or (mm.get((r, k)) is not None):
                break
            have += col_widths[k - c1] if k - c1 < len(col_widths) else col_px(None)
            n += 1
            k += 1
        if n:
            cell["ov"] = n
    g = {"col_widths": col_widths, "rows": rows_out, "styles": gst.list,
         "max_row": r2 - dr, "max_col": c2 - dc}
    if relative:
        g["origin"] = {"r": r1, "c": c1}
    return g


def _cl(c):
    from openpyxl.utils.cell import get_column_letter
    return get_column_letter(c)


def detect_sections(sh, ev):
    """Section headings: merged full-width banners (00) or bold >=12pt text in col A."""
    secs = []
    mm = sh.merge_map()
    for r in sorted(sh.rows):
        v = sh.raw(r, 1)
        if not isinstance(v, str) or not v.strip():
            continue
        xf = sh.xf(r, 1)
        m = mm.get((r, 1))
        f = xf.font
        banner = m and m[0] == "anchor" and m[2] >= max(3, sh.max_col - 1) and xf.fill and f.b
        heading = f.b and (f.sz or 11) >= 12 and r > 1 and _norm_fc(f.color) == "#1F4E79"
        if banner or heading:
            secs.append({"r": r, "label": v})
    return secs


def build_grid(ctx, sheet_name, charts=None):
    wb = ctx.wb
    sh = wb.load_sheet(sheet_name)
    sid = sh.id
    region = (1, 1, sh.max_row, sh.max_col)
    g = build_grid_region(ctx, sheet_name, region, charts=charts)
    title = sh.raw(1, 2) if isinstance(sh.raw(1, 2), str) else None
    notes = []
    for c in range(3, sh.max_col + 1):
        v = sh.raw(1, c)
        if isinstance(v, str):
            notes = v.split("｜")
            break
    if sid in config.GRID_SECTIONS:
        sections = [dict(s) for s in config.GRID_SECTIONS[sid]]
    else:
        sections = detect_sections(sh, ctx.ev)
        for r in config.GRID_EXTRA_SECTIONS.get(sid, []):
            v = sh.raw(r, 1)
            if isinstance(v, str) and all(s["r"] != r for s in sections):
                sections.append({"r": r, "label": v})
        sections.sort(key=lambda s: s["r"])
    obj = {"id": sid, "name": sheet_name, "mode": "grid", "title": title, "notes": notes}
    sub = sh.raw(2, 2)
    if isinstance(sub, str) and sh.raw(2, 1) is None:
        obj["subtitle"] = sub
    obj.update(g)
    obj["sections"] = sections
    obj["charts"] = charts or []
    obj["hidden_cols"] = [c for c in range(1, len(g["col_widths"]) + 1) if sh.col_hidden(c)]
    obj["hidden_rows"] = sorted(sh.hidden_rows)
    obj["gridlines"] = sh.show_gridlines
    obj["freeze_rows"] = sh.pane["y"] if sh.pane else 0
    obj["freeze_cols"] = sh.pane["x"] if sh.pane else 0
    if sh.autofilter:
        r1, c1, r2, c2 = sh.autofilter
        obj["autofilter"] = {"r1": r1, "c1": c1, "r2": r2, "c2": c2}
    return obj
