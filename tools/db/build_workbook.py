# -*- coding: utf-8 -*-
"""Assemble the AMS database workbook(s) from sheets_*.py modules.
Usage: python build_workbook.py <order.json> <out_main.xlsx> <out_detail.xlsx> [detail_threshold_rows]

order.json (from make_order.py): sheets order/numbers, findings, sources, notes.
Fixed numbers: 00_說明, 01_摘要, 02_設備查詢卡 (formula sheet built here), 03_設備總表, 04_位號索引.
Post-processing applied to every sheet (Excel and web share the result, see sheets_final.pkl):
  * 設備總表: columns reordered to [現行AMS位號, 設備別名 (alias), 設備鍵, ...] + 3 param-link columns
  * every sheet with a DeviceKey column gets a '設備別名 (alias)' column right after it
  * 參數現值 sorted by device (alias) with alias at column 1; 參數變更歷程 / 事件總帳 / 位號指派歷程 get '別名#序 (最新=1)'
Sheets with rows >= threshold go to the detail workbook (with a stub sheet in the main one).
"""
import sys, os, json, importlib.util, glob, datetime, math, re, pickle
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import sqlite3

HERE = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(os.path.dirname(HERE), 'AmsDb.sqlite')

order_path, out_main, out_detail = sys.argv[1:4]
THRESH = int(sys.argv[4]) if len(sys.argv) > 4 else 120000
cfg = json.load(open(order_path, encoding='utf-8'))
TITLE = cfg.get('title', '新複循環廠 AMS Device Manager 資料庫解析')
ALIAS = '設備別名 (alias)'
CARD_NUM = '02'
DEFAULT_QUERY = cfg.get('default_query', 'G12HAP70BT001')

# ---- load modules (cached in sheets_cache.pkl; delete the file to force a rebuild)
conn = sqlite3.connect(DB)
CACHE = os.path.join(HERE, 'sheets_cache.pkl')
raw_sheets = pickle.load(open(CACHE, 'rb')) if os.path.exists(CACHE) else {}
all_sheets = {}   # (module_key, name) -> sheet dict
for path in sorted(glob.glob(os.path.join(HERE, 'sheets_*.py'))):
    key = os.path.basename(path)[7:-3]
    if key in cfg.get('skip_modules', []):
        continue
    if key in raw_sheets:
        sheets = raw_sheets[key]
    else:
        spec = importlib.util.spec_from_file_location('sheets_' + key, path)
        mod = importlib.util.module_from_spec(spec)
        try:
            spec.loader.exec_module(mod)
            sheets = mod.build(conn)
        except Exception as e:
            print(f'!! module {key} failed: {e!r}', flush=True)
            continue
        raw_sheets[key] = sheets
        pickle.dump(raw_sheets, open(CACHE, 'wb'), protocol=5)
    for s in sheets:
        if [key, s['name']] in cfg.get('skip_sheets', []):
            print(f'skipped {key}/{s["name"]}', flush=True); continue
        s = dict(s); s['df'] = s['df'].copy(); s['module'] = key
        ren = {}
        for col in s['df'].columns:
            new = str(col)
            for pat, rep in cfg.get('header_renames', []):
                new = re.sub(pat, rep, new)
            if new != str(col):
                ren[col] = new
        if ren:
            s['df'] = s['df'].rename(columns=ren)
        all_sheets[(key, s['name'])] = s
        print(f'loaded {key}/{s["name"]}: {len(s["df"])} rows x {s["df"].shape[1]} cols', flush=True)


def find_col(df, *prefixes):
    cols = [str(c) for c in df.columns]
    for p in prefixes:
        if p in cols:
            return df.columns[cols.index(p)]
        for c in df.columns:
            if str(c).startswith(p):
                return c
    return None


def get(key, name):
    return all_sheets.get((key, name))


# ---- post-processing: 設備總表 order + alias map
dev = get('devices', '設備總表')
df = dev['df']
c_key = find_col(df, '設備鍵'); c_alias = find_col(df, '舊工作簿別名'); c_tag = find_col(df, '現行AMS位號')
df = df.rename(columns={c_alias: ALIAS})
cols = [c_tag, ALIAS, c_key] + [c for c in df.columns if c not in (c_tag, ALIAS, c_key)]
df = df[cols]
alias_of = {int(k): (a if isinstance(a, str) else '') for k, a in zip(df[c_key], df[ALIAS])}
dev['df'] = df

for (k, name), s in all_sheets.items():
    d = s['df']
    if ALIAS in d.columns:
        continue
    old = find_col(d, '舊工作簿別名')
    if old is not None:
        s['df'] = d.rename(columns={old: ALIAS}); continue
    kc = None
    for c in d.columns:
        if 'DeviceKey' in str(c):
            kc = c; break
    if kc is None:
        continue
    vals = pd.to_numeric(d[kc], errors='coerce')
    ali = vals.map(lambda v: alias_of.get(int(v), '') if pd.notna(v) else '')
    pos = list(d.columns).index(kc) + 1
    d.insert(pos, ALIAS, ali.values)
    s['df'] = d
    print(f'alias column added to {k}/{name} after {kc!r} ({(ali != "").sum()} matched)', flush=True)

# 參數現值: alias to column 1, sorted by device
pv = get('params', '參數現值')
d = pv['df']
c0 = d.columns[0]
sort_cols = [ALIAS] + [c for c in (find_col(d, '區塊索引'), find_col(d, '參數 (')) if c is not None]
d = d.sort_values(sort_cols, kind='stable').reset_index(drop=True)
order_cols = [c0, ALIAS] + [c for c in d.columns if c not in (c0, ALIAS)]
d = d[order_cols]
pv['df'] = d
# param-link columns on 設備總表 (Excel row = 0-based index + 4; the web extractor uses the same numbers)
grp = d.groupby(ALIAS, sort=False)
first = grp.indices
starts = {a: int(ix[0]) for a, ix in first.items() if a}
counts = {a: int(len(ix)) for a, ix in first.items() if a}
ddf = dev['df']
ddf['參數現值表'] = ddf[ALIAS].map(lambda a: '參數現值' if a in starts else '')
ddf['參數現值起始列 (Excel列號)'] = ddf[ALIAS].map(lambda a: starts[a] + 4 if a in starts else None).astype('Int64')
ddf['參數現值筆數'] = ddf[ALIAS].map(lambda a: counts.get(a, 0)).astype(int)
dev['df'] = ddf


def add_seq(sheet, time_prefixes, label='別名#序 (最新=1)'):
    d = sheet['df']
    tc = find_col(d, *time_prefixes)
    if tc is None or ALIAS not in d.columns:
        print('!! cannot add seq to', sheet['name']); return
    t = pd.to_datetime(d[tc], errors='coerce')
    tmp = pd.DataFrame({'a': d[ALIAS].fillna(''), 't': t, 'i': range(len(d))})
    tmp = tmp.sort_values(['a', 't', 'i'], ascending=[True, False, False], kind='stable')
    tmp['k'] = tmp.groupby('a').cumcount() + 1
    seq = tmp.sort_values('i')['k'].values
    vals = [f'{a}#{k}' if a else '' for a, k in zip(d[ALIAS].fillna(''), seq)]
    d[label] = vals
    sheet['df'] = d


add_seq(get('params', '參數變更歷程'), ['變更記錄時間(台灣)'])
add_seq(get('events', '事件總帳'), ['時間(台灣)'])
add_seq(get('devices', '位號指派歷程'), ['指派開始(台灣)'])

# ---- ordering (numbers come from order.json; 02 is the formula card)
ordered = []
used = set()
for o in cfg['sheets']:
    k = (o['module_key'], o['sheet_name'])
    if k in all_sheets and k not in used:
        ordered.append((o['number'], all_sheets[k])); used.add(k)
    else:
        print(f'!! ordered sheet not found: {o}', flush=True)
n = max(int(x[0]) for x in ordered) + 1
for kk, s in all_sheets.items():
    if kk not in used:
        ordered.append((f'{n:02d}', s)); used.add(kk); n += 1
        print(f'appended unordered sheet {kk}', flush=True)
ordered.sort(key=lambda x: x[0])
num_of = {s['name']: num for num, s in ordered}
sheet_title = {s['name']: f'{num}_{s["name"]}' for num, s in ordered}
# the param-link column must carry the full sheet name (the web front-end resolves it by name)
ddf = dev['df']
ddf['參數現值表'] = ddf['參數現值表'].map(lambda v: sheet_title['參數現值'] if v else '')
dev['df'] = ddf

# final sheet set for the web extractor
pickle.dump({'title': TITLE, 'ordered': [(num, {k: v for k, v in s.items()}) for num, s in ordered], 'cfg': cfg,
             'card_number': CARD_NUM, 'default_query': DEFAULT_QUERY},
            open(os.path.join(HERE, 'sheets_final.pkl'), 'wb'), protocol=5)

# ---- styles
HDR_FILL = PatternFill('solid', fgColor='DDEBF7')
TITLE_FONT = Font(bold=True, size=14)
HDR_FONT = Font(bold=True)
NOTE_FONT = Font(italic=True, color='555555')
LINK_FONT = Font(color='0563C1', underline='single')
BOLD = Font(bold=True)
LABEL_FONT = Font(color='595959', size=9)
SEC_FONT = Font(bold=True, color='FFFFFF')
SEC_FILL = PatternFill('solid', fgColor='1F4E79')
INPUT_FILL = PatternFill('solid', fgColor='FFF2CC')
HELPER_FONT = Font(color='A6A6A6', size=8)
SEV_FILL = {'高': 'F8CBAD', '中': 'FFE699', '低': 'E2EFDA', '資訊': 'F2F2F2'}
DT_FMT = 'yyyy-mm-dd hh:mm:ss'


def safe_title(s):
    s = re.sub(r'[\\/*?:\[\]]', '_', s)
    return s[:31]


SENTINEL_DATES = {(1899, 12, 30), (1900, 1, 1), (1970, 1, 1), (1, 1, 1), (2036, 2, 5), (2036, 2, 7)}
BLANK_STRINGS = {'nan', 'None', 'NaT', '<NA>', 'NULL'}


def to_cell(ws, v, fmt=None):
    if v is None or v is pd.NA or v is pd.NaT:
        return None
    if isinstance(v, float) and math.isnan(v):
        return None
    try:
        if not isinstance(v, (str, bytes, bytearray, memoryview, list, dict, tuple)) and pd.api.types.is_scalar(v) and pd.isna(v):
            return None
    except Exception:
        pass
    if isinstance(v, pd.Timestamp):
        v = v.to_pydatetime().replace(tzinfo=None)
    if isinstance(v, (datetime.datetime, datetime.date)):
        if (v.year, v.month, v.day) in SENTINEL_DATES or v.year < 1900:
            return None  # sentinel / unset dates are shown blank (see 00_說明)
        c = WriteOnlyCell(ws, value=v)
        c.number_format = DT_FMT
        return c
    if isinstance(v, (pd.Timedelta,)):
        return str(v)
    if isinstance(v, (bytes, bytearray, memoryview)):
        return bytes(v).hex()
    if hasattr(v, 'item') and not isinstance(v, (str, bytes)):
        try:
            v = v.item()
        except Exception:
            pass
    if isinstance(v, bool):
        return '是' if v else '否'
    if isinstance(v, str):
        if v in BLANK_STRINGS:
            return None
        v = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', v)
        if len(v) > 32000:
            v = v[:32000] + '…'
        if v.startswith(('=', '+', '-', '@')) and len(v) > 1 and not re.match(r'^-?\d', v):
            c = WriteOnlyCell(ws, value=v); c.data_type = 's'  # keep as text, never as a formula
            return c
        return v
    return v


def col_width(series, header):
    try:
        sample = series.dropna().astype(str).head(400)
        m = max([len(header)] + [min(len(x), 60) for x in sample]) if len(sample) else len(header)
    except Exception:
        m = len(header)
    return min(max(8, m * 1.15 + 2), 60)


def write_sheet(wb, number, s):
    df = s['df']
    name = safe_title(f'{number}_{s["name"]}')
    ws = wb.create_sheet(title=name)
    cols = list(df.columns)
    widths = [col_width(df[c], str(c)) for c in cols]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    c1 = WriteOnlyCell(ws, value='=HYPERLINK("#\'00_說明\'!A1","⌂ 目錄")'); c1.font = LINK_FONT
    c2 = WriteOnlyCell(ws, value=f'{number}_{s["name"]}'); c2.font = TITLE_FONT
    c3 = WriteOnlyCell(ws, value=f'一列＝{s.get("one_row","")}｜{len(df):,} 列｜來源表：{", ".join(s.get("source_tables", []))}'); c3.font = NOTE_FONT
    ws.append([c1, c2, None, c3])
    notes = s.get('notes') or []
    c = WriteOnlyCell(ws, value=s.get('purpose', '')); c.font = NOTE_FONT
    ws.append([None, c, None, WriteOnlyCell(ws, value='；'.join(str(x) for x in notes)[:2000]) if notes else None])
    hdr = []
    for col in cols:
        h = WriteOnlyCell(ws, value=str(col)); h.font = HDR_FONT; h.fill = HDR_FILL
        h.alignment = Alignment(wrap_text=True, vertical='center')
        hdr.append(h)
    ws.append(hdr)
    ws.freeze_panes = 'A4'
    if len(df):
        ws.auto_filter.ref = f'A3:{get_column_letter(len(cols))}{len(df) + 3}'
    for row in df.itertuples(index=False, name=None):
        ws.append([to_cell(ws, v) for v in row])
    return name


def write_stub(wb, number, s, detail_file):
    name = safe_title(f'{number}_{s["name"]}')
    ws = wb.create_sheet(title=name)
    ws.column_dimensions['B'].width = 90
    c1 = WriteOnlyCell(ws, value='=HYPERLINK("#\'00_說明\'!A1","⌂ 目錄")'); c1.font = LINK_FONT
    c2 = WriteOnlyCell(ws, value=f'{number}_{s["name"]}'); c2.font = TITLE_FONT
    ws.append([c1, c2])
    ws.append([None, WriteOnlyCell(ws, value=f'本表 {len(s["df"]):,} 列，過大，已放在明細活頁簿：{detail_file}（同名工作表）')])
    ws.append([None, WriteOnlyCell(ws, value=s.get('purpose', ''))])
    ws.append([None, WriteOnlyCell(ws, value='欄位：' + '、'.join(str(c) for c in s['df'].columns))])
    for n_ in (s.get('notes') or []):
        ws.append([None, WriteOnlyCell(ws, value=str(n_))])
    return name


def write_readme(wb, entries, detail_entries, detail_file, is_detail=False):
    ws = wb.create_sheet(title='00_說明')
    for col, w in zip('ABCDEFG', (14, 34, 60, 22, 12, 40, 30)):
        ws.column_dimensions[col].width = w
    t = WriteOnlyCell(ws, value=TITLE + ('（明細活頁簿）' if is_detail else '')); t.font = TITLE_FONT
    ws.append([t])
    ws.append([WriteOnlyCell(ws, value=f'建置 {datetime.datetime.now():%Y-%m-%d %H:%M}｜來源 20260912.ams_bckup（SQL Server 2014 原生備份，資料庫 AmsDb，伺服器 AMS1SVR\\EMERSON2014）｜還原於 WSL SQL Server 2025 Express → SQLite → 本簿')])
    ws.append([])
    h = WriteOnlyCell(ws, value='1  資料來源'); h.font = BOLD; ws.append([h])
    for row in cfg.get('sources', []):
        ws.append([None] + [WriteOnlyCell(ws, value=str(x)) for x in row])
    ws.append([])
    h = WriteOnlyCell(ws, value='2  閱讀須知'); h.font = BOLD; ws.append([h])
    for i, row in enumerate(cfg.get('reading_notes', []), 1):
        ws.append([None, WriteOnlyCell(ws, value=f'{i}'), WriteOnlyCell(ws, value=str(row))])
    ws.append([])
    h = WriteOnlyCell(ws, value='3  工作表目錄（點名稱跳轉）'); h.font = BOLD; ws.append([h])
    hdr = ['工作表', '用途', '一列代表', '列數', '欄數', '來源表', '備註']
    ws.append([None] + [WriteOnlyCell(ws, value=x) for x in hdr])
    if not is_detail:
        link = WriteOnlyCell(ws, value=f'=HYPERLINK("#\'{CARD_NUM}_設備查詢卡\'!C3","{CARD_NUM}_設備查詢卡")'); link.font = LINK_FONT
        ws.append([None, link, WriteOnlyCell(ws, value='輸入位號（目前/舊位號、識別時位號、HostTag、裝置ID、設備鍵、別名），即顯示該設備的識別、位置、關鍵組態、最近參數變更、事件與位號歷程'),
                   WriteOnlyCell(ws, value='版面（查詢卡）'), None, None, WriteOnlyCell(ws, value='04_位號索引、03_設備總表、設備參數統計、參數變更歷程、事件總帳、位號指派歷程（公式查表）')])
    for (st, s, where) in entries:
        link = WriteOnlyCell(ws, value=f'=HYPERLINK("#\'{st}\'!A1","{st}")'); link.font = LINK_FONT
        ws.append([None, link, WriteOnlyCell(ws, value=s.get('purpose', '')), WriteOnlyCell(ws, value=s.get('one_row', '')),
                   WriteOnlyCell(ws, value=len(s['df'])), WriteOnlyCell(ws, value=s['df'].shape[1]),
                   WriteOnlyCell(ws, value=', '.join(s.get('source_tables', []))), WriteOnlyCell(ws, value=where)])
    if detail_entries and not is_detail:
        ws.append([])
        h = WriteOnlyCell(ws, value=f'4  明細活頁簿 {detail_file} 內的大表'); h.font = BOLD; ws.append([h])
        for (st, s, where) in detail_entries:
            ws.append([None, WriteOnlyCell(ws, value=st), WriteOnlyCell(ws, value=s.get('purpose', '')), WriteOnlyCell(ws, value=s.get('one_row', '')),
                       WriteOnlyCell(ws, value=len(s['df'])), WriteOnlyCell(ws, value=s['df'].shape[1]), WriteOnlyCell(ws, value=', '.join(s.get('source_tables', [])))])
    ws.append([])
    h = WriteOnlyCell(ws, value='5  各表閱讀註記'); h.font = BOLD; ws.append([h])
    for (st, s, where) in entries + detail_entries:
        for n_ in (s.get('notes') or []):
            ws.append([None, WriteOnlyCell(ws, value=st), WriteOnlyCell(ws, value=str(n_))])


def write_summary(wb):
    ws = wb.create_sheet(title='01_摘要')
    for col, w in zip('ABCDE', (14, 8, 90, 60, 22)):
        ws.column_dimensions[col].width = w
    c1 = WriteOnlyCell(ws, value='=HYPERLINK("#\'00_說明\'!A1","⌂ 目錄")'); c1.font = LINK_FONT
    t = WriteOnlyCell(ws, value=TITLE + '：重點發現'); t.font = TITLE_FONT
    ws.append([c1, t])
    ws.append([])
    for k, v in cfg.get('kpis', []):
        ws.append([None, None, WriteOnlyCell(ws, value=str(k)), WriteOnlyCell(ws, value=str(v))])
    ws.append([])
    ws.append([None] + [WriteOnlyCell(ws, value=x) for x in ('嚴重度', '發現', '證據', '工作表')])
    for f in cfg.get('top_findings', []):
        sev = WriteOnlyCell(ws, value=f.get('severity', ''))
        if f.get('severity') in SEV_FILL:
            sev.fill = PatternFill('solid', fgColor=SEV_FILL[f['severity']])
        ws.append([None, sev, WriteOnlyCell(ws, value=f.get('text', '')), WriteOnlyCell(ws, value=f.get('evidence', '')), WriteOnlyCell(ws, value=f.get('sheet', ''))])
    ws.append([])
    ws.append([None, WriteOnlyCell(ws, value='待決事項')])
    for i, q in enumerate(cfg.get('open_questions', []), 1):
        ws.append([None, WriteOnlyCell(ws, value=str(i)), WriteOnlyCell(ws, value=str(q))])


# ---------------------------------------------------------------- 02_設備查詢卡 (formulas)
def q(sheet_title):
    return "'" + sheet_title.replace("'", "''") + "'"


def ref(sheet_name, prefix):
    """('設備總表', '製造商') -> ("'03_設備總表'", 'Y') column letter by header prefix."""
    s = next(s for num, s in ordered if s['name'] == sheet_name)
    c = find_col(s['df'], prefix)
    if c is None:
        raise KeyError(f'{sheet_name}: no column starting with {prefix!r}')
    return q(sheet_title[sheet_name]), get_column_letter(list(s['df'].columns).index(c) + 1)


def is_dt_col(sheet_name, prefix):
    s = next(s for num, s in ordered if s['name'] == sheet_name)
    c = find_col(s['df'], prefix)
    return pd.api.types.is_datetime64_any_dtype(s['df'][c])


def write_card(wb):
    ws = wb.create_sheet(title=f'{CARD_NUM}_設備查詢卡')
    widths = {'A': 3, 'B': 20, 'C': 26, 'D': 16, 'E': 22, 'F': 20, 'G': 26, 'H': 16, 'I': 30, 'J': 16, 'K': 18, 'L': 16, 'M': 20}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A5'
    S03, T03 = q(sheet_title['設備總表']), ref('設備總表', '現行AMS位號')[1]
    S04 = q(sheet_title['位號索引'])
    A03 = ref('設備總表', ALIAS)[1]
    K04, A04, N04, SRC04 = ref('位號索引', '查詢鍵')[1], ref('位號索引', '設備別名')[1], ref('位號索引', '此字串對應設備數')[1], ref('位號索引', '來源類型')[1]

    rowno = [0]
    _append = ws.append
    def append(cells):
        rowno[0] += 1; _append(cells)
    ws.append = append

    def C(v, font=None, fill=None, fmt=None, align=None):
        c = WriteOnlyCell(ws, value=v)
        if font: c.font = font
        if fill: c.fill = fill
        if fmt: c.number_format = fmt
        if align: c.alignment = align
        return c

    back = C('=HYPERLINK("#\'00_說明\'!A1","⌂ 目錄")', LINK_FONT)
    ws.append([back, C(f'{CARD_NUM}_設備查詢卡', TITLE_FONT), None, C('在 C3 輸入位號後按 Enter：可用目前/舊 AMS 位號、識別時位號、HostTag、裝置ID、設備鍵、別名（大小寫不拘）；查表來源 04_位號索引 → 03_設備總表 等', NOTE_FONT)])
    ws.append([None, C('公式查表（INDEX/MATCH/COUNTIF），不需巨集；找不到請到 04_位號索引 以 Ctrl+F 搜尋部分字串；「(台灣)」欄為 UTC+8', NOTE_FONT)])
    # row 3: input + status
    status = (f'=IFERROR(IF(C4="","請輸入位號…",IF(E4="","查無此字串：請到 04_位號索引 以 Ctrl+F 搜尋部分字串",'
              f'"找到（來源："&M4&"）→ 目前位號 "&INDEX({S03}!${T03}:${T03},I4)&IF(K4>1,"；此字串對應 "&K4&" 台，只顯示優先序最高者",""))),"請輸入位號（C3 為文字格式；若 Excel 把輸入當公式，請先鍵入單引號）")')
    ws.append([None, C('查詢位號 →', BOLD), C(DEFAULT_QUERY, Font(bold=True, size=12), INPUT_FILL, '@'), None, C(status, Font(color='1F4E79'))])
    # row 4: helpers (small grey)
    ws.append([None, C('查詢鍵', HELPER_FONT), C('=UPPER(TRIM(SUBSTITUTE(C3,"=","")))', HELPER_FONT),
               C('索引列', HELPER_FONT), C(f'=IFERROR(MATCH(C4,{S04}!${K04}:${K04},0),"")', HELPER_FONT),
               C('別名', HELPER_FONT), C(f'=IF(E4="","",INDEX({S04}!${A04}:${A04},E4))', HELPER_FONT),
               C('總表列', HELPER_FONT), C(f'=IF(G4="","",IFERROR(MATCH(G4,{S03}!${A03}:${A03},0),""))', HELPER_FONT),
               C('對應設備數', HELPER_FONT), C(f'=IF(E4="","",INDEX({S04}!${N04}:${N04},E4))', HELPER_FONT),
               C('來源類型', HELPER_FONT), C(f'=IF(E4="","",INDEX({S04}!${SRC04}:${SRC04},E4))', HELPER_FONT)])
    ws.append([])

    def field(sheet_name, prefix, rowref):
        S, L = ref(sheet_name, prefix)
        f = f'=IF({rowref}="","",IF(INDEX({S}!${L}:${L},{rowref})="","",INDEX({S}!${L}:${L},{rowref})))'
        return C(f, None, None, DT_FMT if is_dt_col(sheet_name, prefix) else None)

    def section(title):
        ws.append([None, C(title, SEC_FONT, SEC_FILL), C(None, None, SEC_FILL), C(None, None, SEC_FILL), C(None, None, SEC_FILL), C(None, None, SEC_FILL), C(None, None, SEC_FILL), C(None, None, SEC_FILL), C(None, None, SEC_FILL)])

    def pairs(sheet_name, rowref, items):
        # items: list of (label, prefix); 2 pairs per row: B/C and F/G, values may be long → E and I unused
        for i in range(0, len(items), 2):
            row = [None]
            for j in range(2):
                if i + j < len(items):
                    lab, pre = items[i + j]
                    row += [C(lab, LABEL_FONT), field(sheet_name, pre, rowref), None, None]
            ws.append(row)

    R03 = '$I$4'
    section('1  識別與位號（03_設備總表）')
    pairs('設備總表', R03, [('現行 AMS 位號', '現行AMS位號'), ('設備別名', ALIAS), ('設備鍵', '設備鍵'), ('識別時裝置位號', '識別時裝置位號'),
                          ('現行位號品質', '現行位號品質'), ('識別位號品質', '識別位號品質'), ('前一個 AMS 位號', '前一個AMS位號'), ('位號指派次數', '位號指派次數'),
                          ('現行位號指派時間(台灣)', '現行位號指派時間'), ('指派者', '指派者'), ('首次識別時間(台灣)', '首次識別時間'), ('機組', '機組 ('),
                          ('機組說明', '機組說明'), ('KKS 系統碼', 'KKS系統碼'), ('KKS 設備碼', 'KKS設備碼'), ('KKS 流水號', 'KKS流水號')])
    ws.append([])
    section('2  設備（03_設備總表）')
    pairs('設備總表', R03, [('節點種類', '節點種類'), ('協定＋版本', '協定+版本'), ('製造商', '製造商 ('), ('型號', '型號 ('), ('型號代碼', '型號代碼'), ('設備版本', '設備版本 ('),
                          ('主類別', '主類別'), ('次類別', '次類別'), ('裝置ID (HART Device ID)', '裝置ID'), ('處置', '處置 ('), ('設備 GUID', '設備GUID'), ('設備版本鍵', '設備版本鍵')])
    ws.append([])
    section('3  實體位置（03_設備總表：DeviceLocation / NetworkInfo）')
    pairs('設備總表', R03, [('網路名稱', '網路名稱'), ('網路種類', '網路種類'), ('MUX 位址', 'MUX位址'), ('通道', '通道 ('), ('COM 埠號', 'COM埠號'), ('HART 版本', 'HART版本'),
                          ('MUX HART UID', 'MUX HART UID'), ('裝置 HART UID', '裝置 HART UID'), ('FF 連結設備', 'FF連結設備名稱'), ('FF Link 編號', 'FF Link編號'),
                          ('主機位號 (HostTag)', '主機位號'), ('識別狀態', '識別狀態'), ('AMS 路徑', 'AMS路徑'), ('主機路徑', '主機路徑'),
                          ('FF 資源區塊數', 'FF資源區塊數'), ('FF 轉換區塊數', 'FF轉換區塊數'), ('FF 功能區塊數', 'FF功能區塊數'), ('伺服器', '伺服器')])
    ws.append([])
    # 4 關鍵組態 from 設備參數統計 by alias
    Sst, Ast = ref('設備參數統計', ALIAS)
    ws.append([None, C('總表列(參數統計)', HELPER_FONT), C(f'=IF($G$4="","",IFERROR(MATCH($G$4,{Sst}!${Ast}:${Ast},0),""))', HELPER_FONT)])
    RST = f'$C${rowno[0]}'
    section('4  關鍵組態現值（設備參數統計；量程/單位/阻尼/寫入保護…）')
    pairs('設備參數統計', RST, [('參數數', '參數數'), ('變更次數', '變更次數'), ('LRV 量程下限', 'LRV'), ('URV 量程上限', 'URV'), ('量程參數', '量程參數'), ('單位', '單位(解碼)'),
                            ('阻尼 (s)', '阻尼(s)'), ('寫入保護', '寫入保護'), ('轉換函數', '轉換函數'), ('輪詢位址', '輪詢位址'), ('描述 (descriptor)', '描述 ('), ('訊息 (message)', '訊息 ('),
                            ('HART 日期', 'HART 日期'), ('最終組裝號', '最終組裝號'), ('FF XD_SCALE EU0', 'FF XD_SCALE EU0'), ('FF XD_SCALE EU100', 'FF XD_SCALE EU100'),
                            ('FF XD_SCALE 單位', 'FF XD_SCALE 單位'), ('首次記錄(台灣)', '首次記錄'), ('最後記錄(台灣)', '最後記錄'), ('最後事件說明', '最後事件說明'), ('最後使用者', '最後使用者')])
    ws.append([])
    # 5 counts
    section('5  相關筆數（點數字跳到該表；篩選「設備別名」欄）')
    cnt_items = [('參數變更歷程', '參數變更'), ('事件總帳', '事件'), ('位號指派歷程', '位號指派'), ('FF區塊清單', 'FF 區塊'), ('同步失敗統計', '同步失敗'), ('刪除裝置事件', '刪除裝置事件'), ('SnapOn裝置檔案資訊', 'SnapOn 檔案'), ('測試定義指派', '測試定義指派')]
    row = [None]
    for i, (sn, lab) in enumerate(cnt_items):
        if sn not in sheet_title or find_col(next(s for _, s in ordered if s['name'] == sn)['df'], ALIAS) is None:
            continue
        S, L = ref(sn, ALIAS)
        f = f'=IF($G$4="","",HYPERLINK("#{S}!A"&IFERROR(MATCH($G$4,{S}!${L}:${L},0),1),COUNTIF({S}!${L}:${L},$G$4)&" 筆"))'
        row += [C(lab, LABEL_FONT), C(f, LINK_FONT)]
        if len(row) >= 9:
            ws.append(row); row = [None]
    S, L = ref('設備總表', '參數現值筆數')
    row += [C('參數現值（明細簿）', LABEL_FONT), C(f'=IF($I$4="","",INDEX({S}!${L}:${L},$I$4)&" 筆")')]
    ws.append(row)
    ws.append([])

    def listing(title, sheet_name, seq_prefix, columns, kmax):
        S, Lseq = ref(sheet_name, seq_prefix)
        section(title)
        ws.append([None, C('k', LABEL_FONT), C('↗', LABEL_FONT)] + [C(lab, LABEL_FONT) for lab, _ in columns])
        for k in range(1, kmax + 1):
            r = f'IFERROR(MATCH($G$4&"#{k}",{S}!${Lseq}:${Lseq},0),"")'
            cells = [None, C(k, HELPER_FONT), C(f'=IF({r}="","",HYPERLINK("#{S}!A"&{r},"↗"))', LINK_FONT)]
            for lab, pre in columns:
                _, L = ref(sheet_name, pre)
                fmt = DT_FMT if is_dt_col(sheet_name, pre) else None
                cells.append(C(f'=IF({r}="","",IF(INDEX({S}!${L}:${L},{r})="","",INDEX({S}!${L}:${L},{r})))', None, None, fmt))
            ws.append(cells)
        ws.append([])

    listing('6  最近 10 筆參數變更（參數變更歷程，k=1 最新）', '參數變更歷程', '別名#序',
            [('變更記錄時間(台灣)', '變更記錄時間(台灣)'), ('參數', '參數 ('), ('中文名稱', '參數中文名稱'), ('關鍵組態類別', '關鍵組態類別'), ('前值', '前值'), ('新值', '新值'), ('事件說明', '事件說明'), ('使用者', '使用者 ('), ('前值記錄時間(台灣)', '前值記錄時間(台灣)'), ('前值使用者', '前值使用者')], 10)
    listing('7  最近 15 筆事件（事件總帳，k=1 最新）', '事件總帳', '別名#序',
            [('時間(台灣)', '時間(台灣)'), ('類別', '類別中文'), ('描述', '描述 ('), ('使用者', '使用者 ('), ('電腦 IP', '電腦 IP'), ('事件當時位號', '事件當時位號'), ('其他 (Other)', '其他 ('), ('本事件寫入參數筆數', '本事件寫入參數筆數')], 15)
    listing('8  位號指派歷程（k=1 最新）', '位號指派歷程', '別名#序',
            [('位號', '位號 ('), ('狀態', '狀態'), ('指派開始(台灣)', '指派開始(台灣)'), ('指派結束(台灣)', '指派結束(台灣)'), ('開始事件', '開始事件 ('), ('開始事件使用者', '開始事件使用者'), ('結束事件', '結束事件 ('), ('結束事件使用者', '結束事件使用者')], 6)
    return ws.title


# ---- split
main_entries, detail_entries = [], []
for number, s in ordered:
    (detail_entries if len(s['df']) >= THRESH else main_entries).append((number, s))

wb = Workbook(write_only=True)
readme_entries = [(safe_title(f'{n}_{s["name"]}'), s, '') for n, s in main_entries]
stub_entries = [(safe_title(f'{n}_{s["name"]}'), s, '明細活頁簿') for n, s in detail_entries]
write_readme(wb, readme_entries, stub_entries, os.path.basename(out_detail))
write_summary(wb)
write_card(wb)
print('wrote 02_設備查詢卡', flush=True)
for number, s in ordered:
    if len(s['df']) >= THRESH:
        write_stub(wb, number, s, os.path.basename(out_detail))
    else:
        write_sheet(wb, number, s)
    print(f'wrote {number}_{s["name"]}', flush=True)
wb.save(out_main)
print('saved', out_main, os.path.getsize(out_main) // 1024, 'KB', flush=True)

if detail_entries:
    wb2 = Workbook(write_only=True)
    write_readme(wb2, [(safe_title(f'{n}_{s["name"]}'), s, '') for n, s in detail_entries], [], '', is_detail=True)
    for number, s in detail_entries:
        write_sheet(wb2, number, s)
        print(f'wrote detail {number}_{s["name"]}', flush=True)
    wb2.save(out_detail)
    print('saved', out_detail, os.path.getsize(out_detail) // 1024, 'KB', flush=True)
print('BUILD_DONE')
